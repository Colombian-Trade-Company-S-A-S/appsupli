"""
Pruebas de la sincronización con Odoo.

Odoo se reemplaza por un cliente falso con un organigrama chico: una dirección
con un área, un área sin dirección y empleados de cada tipo. Nunca se llama a
la red.
"""
import pytest
from rest_framework.test import APIClient

from apps.accounts.models import Area, Departamento, User
from apps.odoo.client import EstadoOdoo, OdooError
from apps.odoo.models import Sincronizacion
from apps.odoo.sync import separar_nombre, sincronizar

pytestmark = pytest.mark.django_db

EMPRESA = [1, 'COLOMBIAN TRADE COMPANY SAS']


def departamento(id_, nombre, padre=None, activo=True):
    return {'id': id_, 'name': nombre, 'parent_id': [padre, ''] if padre else False, 'active': activo}


def empleado(id_, nombre, *, correo=False, cedula=False, jefe=None, depto=None, cargo='ASESOR', activo=True):
    return {
        'id': id_, 'name': nombre, 'active': activo, 'work_email': correo, 'identification_id': cedula,
        'job_id': False, 'job_title': cargo, 'department_id': [depto, ''] if depto else False,
        'parent_id': [jefe, ''] if jefe else False, 'company_id': EMPRESA,
        'work_location_id': [3, 'Regional Centro Norte'], 'write_date': '2026-10-01 00:00:00',
    }


class OdooFalso:
    def __init__(self, departamentos=None, empleados=None, error=None):
        self._departamentos = departamentos if departamentos is not None else [
            departamento(28, 'Commercial Directorate'),
            departamento(22, 'SALES', padre=28),
            departamento(7, 'TECH'),
        ]
        self._empleados = empleados if empleados is not None else [
            empleado(1, 'ORDOÑEZ ARENAS JUAN CAMILO', correo='jordonez@supli.tech', cedula='100', depto=28),
            empleado(2, 'CALDERON MEJIA ANGIE LORENA', correo='acalderon@supli.tech', cedula='200', jefe=1, depto=22),
            empleado(3, 'ACUÑA RODRÍGUEZ EBERTO DE JESÚS', cedula='300', jefe=2, depto=22),
            empleado(4, 'AYALA PARRA LAURA CATALINA', correo='lauraayala@coltrade.com.co', cedula='400', depto=7),
        ]
        self._error = error

    def departamentos(self):
        if self._error:
            raise OdooError(self._error)
        return self._departamentos

    def empleados(self):
        return self._empleados


def todo(cliente, **kw):
    opciones = {'crear': True, 'actualizar': True, 'desactivar': True, **kw}
    return sincronizar(cliente=cliente, **opciones)


def test_separa_apellidos_y_nombres_con_particulas():
    assert separar_nombre('ACUÑA RODRÍGUEZ EBERTO DE JESÚS') == ('Eberto de Jesús', 'Acuña Rodríguez')
    assert separar_nombre('CORREA ARIAS DANIEL') == ('Daniel', 'Correa Arias')
    assert separar_nombre('PORRAS MARIA') == ('Maria', 'Porras')
    assert separar_nombre('katherin') == ('Katherin', '')


def test_crea_empleados_con_jerarquia_y_organigrama():
    registro = todo(OdooFalso())

    assert registro.estado == 'ok', registro.error
    assert registro.resumen['creados'] == 4
    juan = User.objects.get(odoo_id=1)
    angie = User.objects.get(odoo_id=2)
    asesor = User.objects.get(odoo_id=3)
    laura = User.objects.get(odoo_id=4)

    assert (angie.first_name, angie.last_name) == ('Angie Lorena', 'Calderon Mejia')
    assert angie.manager == juan and asesor.manager == angie
    assert juan.kind == angie.kind == User.Kind.LEADER
    assert asesor.kind == User.Kind.COLLABORATOR
    # Dirección y área salen del árbol, no del texto del nombre.
    assert (angie.direccion, angie.area.name) == ('Commercial Directorate', 'SALES')
    assert (juan.direccion, juan.area) == ('Commercial Directorate', None)
    assert (laura.direccion, laura.area.name) == ('', 'TECH')
    assert angie.organizacion == 'COLOMBIAN TRADE COMPANY SAS'
    assert angie.regional == 'Regional Centro Norte'
    # Sin clave: nadie entra hasta que tenga Microsoft o el admin le ponga una.
    assert not angie.has_usable_password()


def test_solo_copia_correos_del_dominio_de_ingreso():
    registro = todo(OdooFalso())

    assert User.objects.get(odoo_id=2).email == 'acalderon@supli.tech'
    assert User.objects.get(odoo_id=3).email is None
    assert User.objects.get(odoo_id=4).email is None
    tipos = {(e['odoo_id'], e['tipo']) for e in registro.excepciones}
    assert (4, 'correo_otro_dominio') in tipos


def test_empareja_por_correo_sin_duplicar():
    existente = User.objects.create_user(
        email='acalderon@supli.tech', username='angie', first_name='Angie', password='clave-123456'
    )
    todo(OdooFalso())

    existente.refresh_from_db()
    assert existente.odoo_id == 2
    assert User.objects.filter(email='acalderon@supli.tech').count() == 1
    # Su clave sigue sirviendo: emparejar no es recrear.
    assert existente.check_password('clave-123456')


def test_empareja_por_cedula():
    existente = User.objects.create(username='eberto', first_name='Eberto', cedula='300')
    todo(OdooFalso())
    existente.refresh_from_db()
    assert existente.odoo_id == 3


def test_es_idempotente():
    todo(OdooFalso())
    registro = todo(OdooFalso())
    assert registro.resumen.get('creados', 0) == 0
    assert registro.resumen['sin_cambios'] == 4
    assert User.objects.filter(odoo_id__isnull=False).count() == 4
    assert Departamento.objects.count() == 3


def test_actualiza_lo_que_cambia_en_odoo():
    todo(OdooFalso())
    cliente = OdooFalso()
    cliente._empleados[2]['job_title'] = 'PROMOTOR'
    registro = todo(cliente)
    assert registro.resumen['actualizados'] == 1
    assert User.objects.get(odoo_id=3).position == 'PROMOTOR'


def test_sin_actualizar_no_toca_a_los_existentes():
    todo(OdooFalso())
    cliente = OdooFalso()
    cliente._empleados[2]['job_title'] = 'PROMOTOR'
    registro = sincronizar(cliente=cliente, crear=True, actualizar=False, desactivar=False)
    assert registro.resumen['existentes_sin_actualizar'] == 4
    assert User.objects.get(odoo_id=3).position == 'ASESOR'


def test_un_vacio_en_odoo_no_borra_lo_que_hay():
    todo(OdooFalso())
    cliente = OdooFalso()
    cliente._empleados[2]['department_id'] = False
    cliente._empleados[2]['parent_id'] = False
    todo(cliente)
    asesor = User.objects.get(odoo_id=3)
    assert asesor.area.name == 'SALES'
    assert asesor.manager.odoo_id == 2


def test_las_bajas_se_desactivan_sin_borrar():
    todo(OdooFalso())
    cliente = OdooFalso()
    cliente._empleados[2]['active'] = False
    registro = todo(cliente)
    assert registro.resumen['desactivados'] == 1
    assert User.objects.get(odoo_id=3).is_active is False


def test_sin_la_opcion_las_bajas_quedan_pendientes():
    todo(OdooFalso())
    cliente = OdooFalso()
    cliente._empleados[2]['active'] = False
    registro = sincronizar(cliente=cliente, crear=False, actualizar=True, desactivar=False)
    assert registro.resumen['bajas_sin_desactivar'] == 1
    assert User.objects.get(odoo_id=3).is_active is True


def test_un_admin_archivado_en_odoo_no_se_desactiva_solo():
    User.objects.create_superuser(
        email='jordonez@supli.tech', username='juan', first_name='Juan', password='clave-123456'
    )
    cliente = OdooFalso()
    cliente._empleados[0]['active'] = False
    registro = todo(cliente)
    assert User.objects.get(username='juan').is_active is True
    assert any(e['tipo'] == 'admin_de_baja' for e in registro.excepciones)


def test_el_admin_conserva_su_tipo():
    User.objects.create_superuser(
        email='jordonez@supli.tech', username='juan', first_name='Juan', password='clave-123456'
    )
    User.objects.filter(username='juan').update(kind=User.Kind.ADMIN)
    todo(OdooFalso())
    assert User.objects.get(username='juan').kind == User.Kind.ADMIN


def test_la_vista_previa_no_guarda_nada():
    registro = todo(OdooFalso(), simular=True)
    assert registro.resumen['creados'] == 4
    assert User.objects.count() == 0
    assert Departamento.objects.count() == 0
    assert Area.objects.count() == 0
    # Pero sí queda en la bitácora.
    assert Sincronizacion.objects.filter(simulacion=True).count() == 1


def test_si_odoo_falla_no_se_toca_nada_y_queda_registrado():
    existente = User.objects.create(username='alguien', first_name='Alguien')
    registro = todo(OdooFalso(error='No hubo conexión con Odoo'))
    assert registro.estado == 'error'
    assert 'conexión' in registro.error
    assert list(User.objects.all()) == [existente]


def test_reporta_excepciones_de_calidad():
    cliente = OdooFalso(empleados=[
        empleado(1, 'UNO PRIMERO JUAN', cedula='9', correo='x@supli.tech'),
        empleado(2, 'DOS SEGUNDO ANA', cedula='9', correo='x@supli.tech', jefe=1, depto=7),
        empleado(3, 'TRES TERCERO LUIS', jefe=1),
    ])
    registro = todo(cliente)
    tipos = {(e['odoo_id'], e['tipo']) for e in registro.excepciones}
    assert {(1, 'cedula_duplicada'), (2, 'cedula_duplicada'), (1, 'correo_duplicado')} <= tipos
    assert {(1, 'sin_jefe'), (1, 'sin_departamento'), (3, 'sin_cedula')} <= tipos
    # Un correo repetido en Odoo no se le pone a ninguno de los dos.
    assert not User.objects.filter(email='x@supli.tech').exists()


# ── API ──────────────────────────────────────────────────────────────────
def cliente_admin() -> APIClient:
    admin = User.objects.create_superuser(
        email='admin@supli.tech', username='admin', first_name='Admin', password='clave-123456'
    )
    api = APIClient()
    api.force_authenticate(user=admin)
    return api


def test_la_api_es_solo_para_admin():
    persona = User.objects.create_user(
        email='p@supli.tech', username='p', first_name='P', password='clave-123456'
    )
    api = APIClient()
    api.force_authenticate(user=persona)
    assert api.get('/api/admin/odoo/estado').status_code == 403
    assert api.post('/api/admin/odoo/sincronizar', {'crear': True}, format='json').status_code == 403


def test_estado_api(monkeypatch):
    monkeypatch.setattr(
        'apps.odoo.views.OdooClient.estado',
        lambda self: EstadoOdoo(True, True, 'https://x.odoo.com', 'x', 'u', '18.0+e', 120, 62, 'Conexión correcta.'),
    )
    datos = cliente_admin().get('/api/admin/odoo/estado').json()
    assert datos['conectado'] is True
    assert datos['empleadosActivos'] == 62
    assert datos['ultimaSincronizacion'] is None
    assert 'secreto' not in str(datos).lower()


def test_sincronizar_por_api(monkeypatch):
    monkeypatch.setattr('apps.odoo.sync.OdooClient', lambda: OdooFalso())
    api = cliente_admin()
    respuesta = api.post('/api/admin/odoo/sincronizar', {'crear': True, 'actualizar': True}, format='json')
    assert respuesta.status_code == 200, respuesta.data
    assert respuesta.json()['resumen']['creados'] == 4
    assert respuesta.json()['ejecutadaPorNombre'] == 'Admin'
    assert len(api.get('/api/admin/odoo/sincronizaciones').json()) == 1


def test_sincronizar_exige_una_opcion():
    respuesta = cliente_admin().post('/api/admin/odoo/sincronizar', {}, format='json')
    assert respuesta.status_code == 400


def test_sin_variables_el_estado_lo_dice(settings):
    settings.ODOO_URL = ''
    datos = cliente_admin().get('/api/admin/odoo/estado').json()
    assert datos['configurado'] is False and datos['conectado'] is False


def test_genera_accesos_solo_a_quien_no_tiene_clave():
    from io import BytesIO

    from openpyxl import load_workbook

    todo(OdooFalso())
    juan = User.objects.get(odoo_id=1)
    juan.set_password('ya-tenia-clave-1')
    juan.save()
    api = cliente_admin()

    assert api.get('/api/admin/odoo/accesos').json()['pendientes'] == 1
    respuesta = api.post('/api/admin/odoo/accesos')
    assert respuesta.status_code == 200
    filas = list(load_workbook(BytesIO(respuesta.content)).active.iter_rows(values_only=True))
    # Encabezado + Angie: es la única de Odoo con correo @supli.tech y sin clave.
    assert len(filas) == 2
    _, correo, clave, *_ = filas[1]
    assert correo == 'acalderon@supli.tech'
    assert User.objects.get(odoo_id=2).check_password(clave)
    juan.refresh_from_db()
    assert juan.check_password('ya-tenia-clave-1')
    # Una segunda vez ya no hay a quién.
    assert api.post('/api/admin/odoo/accesos').status_code == 400


def test_a_quien_viene_de_odoo_no_se_le_editan_sus_datos():
    todo(OdooFalso())
    angie = User.objects.get(odoo_id=2)
    respuesta = cliente_admin().patch(
        f'/api/admin/users/{angie.id}',
        {'firstName': 'Otra', 'position': 'Otro cargo', 'kind': 'colaborador', 'phone': '300'},
        format='json',
    )
    assert respuesta.status_code == 200, respuesta.data
    angie.refresh_from_db()
    assert (angie.first_name, angie.position, angie.phone) == ('Angie Lorena', 'ASESOR', '300')
    # Tiene equipo en Odoo: sigue siendo líder aunque el formulario diga otra cosa.
    assert angie.kind == User.Kind.LEADER


def test_las_areas_no_se_crean_desde_administracion():
    api = cliente_admin()
    assert api.post('/api/admin/areas', {'name': 'Nueva'}, format='json').status_code == 405
    todo(OdooFalso())
    assert set(Area.objects.filter(odoo=True).values_list('name', flat=True)) == {'SALES', 'TECH'}


# ── Eliminar a quien no está en Odoo ─────────────────────────────────────
def _objetivo(colaborador, registrado_por):
    from datetime import date

    from apps.performance.models import Objetivo

    return Objetivo.objects.create(
        colaborador=colaborador, registrado_por=registrado_por, periodo=date(2026, 10, 1),
        objetivo='Vender', kpi='Ventas', peso=100, tipo_medicion=Objetivo._meta.get_field('tipo_medicion').choices[0][0],
    )


def test_elimina_a_quien_no_esta_en_odoo_con_sus_registros():
    from apps.performance.models import Objetivo

    ejecutor = User.objects.create_superuser(
        email='admin@supli.tech', username='admin', first_name='Admin', password='clave-123456'
    )
    angie = User.objects.create_user(
        email='acalderon@supli.tech', username='angie', first_name='Angie', password='clave-123456'
    )
    viejo = User.objects.create_user(
        email='viejo@coltrade.com.co', username='viejo', first_name='Viejo', password='clave-123456'
    )
    otro_admin = User.objects.create_user(
        email='soporte@supli.tech', username='soporte', first_name='Soporte', password='clave-123456'
    )
    User.objects.filter(pk=otro_admin.pk).update(kind=User.Kind.ADMIN)
    # El viejo registró un objetivo para Angie: la base lo protege del borrado.
    _objetivo(angie, viejo)
    _objetivo(viejo, angie)

    registro = todo(OdooFalso(), eliminar=True, ejecutada_por=ejecutor)

    assert registro.estado == 'ok', registro.error
    assert registro.resumen['eliminados'] == 1
    assert not User.objects.filter(username='viejo').exists()
    assert Objetivo.objects.count() == 0
    # Angie emparejó por correo; los admins se conservan aunque no estén en Odoo.
    assert User.objects.filter(pk__in=[angie.pk, ejecutor.pk, otro_admin.pk]).count() == 3
    tipos = [e['tipo'] for e in registro.excepciones]
    assert tipos.count('conservado_sin_odoo') == 2 and tipos.count('eliminado') == 1


def test_no_elimina_a_quien_esta_archivado_en_odoo():
    cliente = OdooFalso()
    cliente._empleados.append(empleado(9, 'BAJA ANTIGUA PEDRO', correo='pedro@supli.tech', activo=False))
    pedro = User.objects.create_user(
        email='pedro@supli.tech', username='pedro', first_name='Pedro', password='clave-123456'
    )
    todo(cliente, eliminar=True)
    pedro.refresh_from_db()
    assert pedro.is_active is False


def test_la_vista_previa_de_eliminar_no_borra():
    User.objects.create_user(email='viejo@x.com', username='viejo', first_name='Viejo', password='clave-123456')
    registro = todo(OdooFalso(), eliminar=True, simular=True)
    assert registro.resumen['eliminados'] == 1
    assert User.objects.filter(username='viejo').exists()


def test_un_odoo_vacio_no_vacia_appsupli():
    User.objects.create_user(email='viejo@x.com', username='viejo', first_name='Viejo', password='clave-123456')
    registro = todo(OdooFalso(empleados=[]), eliminar=True)
    assert registro.estado == 'error'
    assert User.objects.filter(username='viejo').exists()


def test_empareja_por_nombre_y_conserva_el_correo_personal():
    asesor = User.objects.create_user(
        email='eberto123@gmail.com', username='eberto', first_name='Eberto de Jesus',
        last_name='Acuna Rodriguez', password='clave-123456',
    )
    registro = todo(OdooFalso(), eliminar=True)
    asesor.refresh_from_db()
    assert asesor.odoo_id == 3
    assert asesor.email == 'eberto123@gmail.com'
    assert asesor.check_password('clave-123456')
    assert registro.resumen.get('eliminados', 0) == 0
    assert User.objects.filter(odoo_id__isnull=False).count() == 4


def test_no_empareja_por_nombre_si_es_ambiguo():
    cliente = OdooFalso(empleados=[
        empleado(1, 'PEREZ GOMEZ ANA', cedula='1'),
        empleado(2, 'PEREZ GOMEZ ANA', cedula='2'),
    ])
    ana = User.objects.create(username='ana', first_name='Ana', last_name='Perez Gomez')
    todo(cliente)
    ana.refresh_from_db()
    assert ana.odoo_id is None


def test_empareja_si_el_nombre_de_aca_esta_incompleto():
    alejandro = User.objects.create(username='alejo', first_name='Alejandro', last_name='Trujillo', email='a@gmail.com')
    cliente = OdooFalso()
    cliente._empleados.append(empleado(7, 'ECHEVERRY TRUJILLO ALEJANDRO', cedula='700', jefe=2, depto=22))
    registro = todo(cliente, eliminar=True)
    alejandro.refresh_from_db()
    assert alejandro.odoo_id == 7
    assert registro.resumen.get('eliminados', 0) == 0


def test_el_jefe_emparejado_por_correo_queda_asignado_aunque_venga_despues():
    # El jefe ya existía (sin odoo_id) y en Odoo viene después de su equipo.
    juan = User.objects.create_user(
        email='jordonez@supli.tech', username='juan', first_name='Juan', password='clave-123456'
    )
    cliente = OdooFalso()
    cliente._empleados.reverse()
    todo(cliente)
    assert User.objects.get(odoo_id=2).manager == juan


def test_eliminar_borra_las_areas_que_no_son_de_odoo():
    Area.objects.create(name='Logistics')
    Area.objects.create(name='Tech')  # Misma área que en Odoo: se conserva y queda marcada.
    registro = todo(OdooFalso(), eliminar=True)
    assert set(Area.objects.values_list('name', flat=True)) == {'SALES', 'Tech'}
    assert Area.objects.filter(odoo=False).count() == 0
    assert registro.resumen['areas_eliminadas'] == 1


def test_sin_eliminar_las_areas_ajenas_se_quedan():
    Area.objects.create(name='Logistics')
    todo(OdooFalso())
    assert Area.objects.filter(name='Logistics').exists()
