"""
Pruebas de Objetivos y KPIs: permisos, reglas de negocio y motor de cálculo.

Las reglas que se prueban acá son las del §6 de la especificación técnica; el
número de cada una va en el nombre de la prueba cuando aplica.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from django.core.management import call_command
from rest_framework.test import APIClient

from apps.accounts.models import Application, Permission, Role, User
from apps.performance.cumplimiento import (
    FormulaInvalida,
    calcular_cumplimiento,
    cumplimiento_total,
    evaluar_formula,
    semaforo,
)
from apps.performance.models import (
    EstadoObjetivo,
    EstadoPeriodo,
    EstadoValidacion,
    Objetivo,
    Periodo,
    Semaforo,
    TipoMedicion,
    Unidad,
)
from apps.performance.periodos import etiqueta_corte, mes_siguiente, primer_dia, rango

pytestmark = pytest.mark.django_db

# El mes de las pruebas es siempre el entrante: los objetivos se editan hasta
# el último día del mes anterior (A9), así que fijar un mes del calendario
# haría fallar la suite sola al llegar esa fecha.
OCTUBRE = mes_siguiente(primer_dia(date.today()))
MES = OCTUBRE.strftime('%Y-%m')
RUTA = '/api/performance'


@pytest.fixture
def app_performance():
    contenedor = Application.objects.create(
        code='supli-performance',
        name='Supli Performance',
        base_path='/inicio/performance',
        order=300,
    )
    app = Application.objects.create(
        code='objetivos-kpis',
        name='Objetivos y KPIs',
        base_path='/inicio/performance/objetivos',
        order=10,
        parent=contenedor,
    )
    for code, name in (
        ('performance:objetivos:manage_all', 'Definir objetivos de cualquiera'),
        ('performance:objetivos:view_all', 'Ver toda la organización'),
        ('performance:periodos:manage', 'Gestionar el periodo'),
    ):
        Permission.objects.create(code=code, name=name, application=app)
    return app


def crear_usuario(email, app=None, permisos=(), jefe=None, cargo=''):
    usuario = User.objects.create_user(
        email=email,
        username=email.split('@')[0],
        first_name=email.split('@')[0].title(),
        password='clave-de-prueba-123',
        position=cargo,
    )
    if jefe:
        usuario.manager = jefe
        usuario.save(update_fields=['manager'])
    if app:
        usuario.applications.add(app)
    if permisos:
        rol = Role.objects.create(code=f'rol-{usuario.pk}', name='rol')
        rol.permissions.set(Permission.objects.filter(code__in=permisos))
        usuario.roles.add(rol)
    return usuario


def cliente_de(usuario) -> APIClient:
    cliente = APIClient()
    cliente.force_authenticate(user=usuario)
    return cliente


def objetivo_de(colaborador, jefe, peso, **extra) -> Objetivo:
    datos = {
        'objetivo': 'Entregar el portal de autogestión',
        'kpi': 'Portal en producción',
        'peso': Decimal(peso),
        'tipo_medicion': TipoMedicion.BINARIO,
        'unidad': Unidad.SI_NO,
        'periodo': OCTUBRE,
    }
    datos.update(extra)
    return Objetivo.objects.create(colaborador=colaborador, registrado_por=jefe, **datos)


@pytest.fixture
def equipo(app_performance):
    """Un jefe con dos personas a cargo, que es la jerarquía que usa el módulo."""
    jefe = crear_usuario('jefe@supli.tech', app_performance, cargo='Tech Manager')
    julian = crear_usuario('julian@supli.tech', app_performance, jefe=jefe, cargo='BI Analyst')
    daniel = crear_usuario('daniel@supli.tech', app_performance, jefe=jefe, cargo='BI Manager')
    return jefe, julian, daniel


# ── Motor de cálculo (§5) ──────────────────────────────────────────────────


def test_el_binario_es_todo_o_nada():
    objetivo = Objetivo(tipo_medicion=TipoMedicion.BINARIO, tope_cumplimiento=100)
    assert calcular_cumplimiento(objetivo, 1) == Decimal('100.00')
    assert calcular_cumplimiento(objetivo, 0) == Decimal('0.00')


def test_el_proporcional_divide_lo_logrado_entre_la_meta():
    objetivo = Objetivo(
        tipo_medicion=TipoMedicion.PROPORCIONAL, meta_valor=Decimal('8'), tope_cumplimiento=100
    )
    assert calcular_cumplimiento(objetivo, 6) == Decimal('75.00')
    # Sin sobrecumplimiento, pasarse de la meta se topa en 100.
    assert calcular_cumplimiento(objetivo, 10) == Decimal('100.00')

    objetivo.permite_sobrecumplimiento = True
    assert calcular_cumplimiento(objetivo, 10) == Decimal('125.00')


def test_el_proporcional_inverso_premia_el_valor_mas_bajo():
    """Menos es mejor: días, costos, tickets."""
    objetivo = Objetivo(
        tipo_medicion=TipoMedicion.PROPORCIONAL_INVERSO,
        meta_valor=Decimal('2'),
        tope_cumplimiento=100,
    )
    assert calcular_cumplimiento(objetivo, 2) == Decimal('100.00')
    assert calcular_cumplimiento(objetivo, 4) == Decimal('50.00')
    # Ejecutar cero es el mejor resultado posible, no una división por cero.
    assert calcular_cumplimiento(objetivo, 0) == Decimal('100.00')


def test_la_formula_solo_evalua_aritmetica_con_logrado_y_meta():
    assert evaluar_formula('logrado / meta * 100', {'logrado': 5, 'meta': 10}) == 50
    assert evaluar_formula('min(logrado / meta * 100, 120)', {'logrado': 20, 'meta': 10}) == 120

    for peligrosa in (
        '__import__("os").system("dir")',
        'open("password.txt").read()',
        'otra_variable * 2',
        'logrado.__class__',
    ):
        with pytest.raises(FormulaInvalida):
            evaluar_formula(peligrosa, {'logrado': 1, 'meta': 1})


def test_el_cumplimiento_de_la_persona_pondera_sus_objetivos():
    """Σ (peso/100 × %): es lo que después ordena el Top Performance."""
    assert cumplimiento_total([(30, 100), (25, 50), (45, 0)]) == Decimal('42.50')


# ── Definición de objetivos (§6, reglas 1 a 5) ─────────────────────────────


def test_el_jefe_define_los_objetivos_de_su_equipo(equipo):
    jefe, julian, _ = equipo

    respuesta = cliente_de(jefe).post(
        f'{RUTA}/objetivos',
        {
            'colaborador': julian.pk,
            'periodo': OCTUBRE.isoformat(),
            'objetivo': 'Entregar el portal de autogestión BI',
            'kpi': 'Portal en producción',
            'peso': '30',
            'tipoMedicion': 'binario',
            'unidad': 'si_no',
        },
        format='json',
    )

    assert respuesta.status_code == 201, respuesta.data
    creado = Objetivo.objects.get()
    assert creado.registrado_por == jefe
    assert creado.estado == EstadoObjetivo.BORRADOR
    # El responsable del resultado, si no se dice otra cosa, es el colaborador.
    assert creado.responsable_resultado == julian
    # El mes se abre solo con el primer objetivo.
    assert Periodo.objects.get(periodo=OCTUBRE).estado == EstadoPeriodo.DEFINICION


def test_nadie_define_sus_propios_objetivos(equipo):
    """Regla 3: el objetivo lo crea el jefe o People, nunca el colaborador."""
    _, julian, _ = equipo

    respuesta = cliente_de(julian).post(
        f'{RUTA}/objetivos',
        {
            'colaborador': julian.pk,
            'periodo': OCTUBRE.isoformat(),
            'objetivo': 'Auto-asignarme algo fácil',
            'kpi': 'Lo que yo diga',
            'peso': '100',
            'tipoMedicion': 'binario',
        },
        format='json',
    )

    assert respuesta.status_code == 403
    assert Objetivo.objects.count() == 0


def test_un_jefe_no_toca_el_equipo_de_otro(app_performance, equipo):
    jefe, julian, _ = equipo
    otro_jefe = crear_usuario('otro@supli.tech', app_performance)

    respuesta = cliente_de(otro_jefe).post(
        f'{RUTA}/objetivos',
        {
            'colaborador': julian.pk,
            'periodo': OCTUBRE.isoformat(),
            'objetivo': 'Objetivo de otro equipo',
            'kpi': 'KPI',
            'peso': '10',
            'tipoMedicion': 'binario',
        },
        format='json',
    )

    assert respuesta.status_code == 403


def test_people_define_objetivos_de_cualquiera(app_performance, equipo):
    _, julian, _ = equipo
    people = crear_usuario(
        'people@supli.tech', app_performance, ['performance:objetivos:manage_all']
    )

    respuesta = cliente_de(people).post(
        f'{RUTA}/objetivos',
        {
            'colaborador': julian.pk,
            'periodo': OCTUBRE.isoformat(),
            'objetivo': 'Objetivo transversal de People',
            'kpi': 'KPI',
            'peso': '20',
            'tipoMedicion': 'binario',
        },
        format='json',
    )

    assert respuesta.status_code == 201, respuesta.data


def test_no_se_pueden_pasar_de_seis_objetivos(equipo):
    """Regla 2: máximo seis objetivos por persona y mes."""
    jefe, julian, _ = equipo
    for _ in range(6):
        objetivo_de(julian, jefe, '10')

    respuesta = cliente_de(jefe).post(
        f'{RUTA}/objetivos',
        {
            'colaborador': julian.pk,
            'periodo': OCTUBRE.isoformat(),
            'objetivo': 'El séptimo',
            'kpi': 'KPI',
            'peso': '10',
            'tipoMedicion': 'binario',
        },
        format='json',
    )

    assert respuesta.status_code == 400
    assert 'máximo' in str(respuesta.data)
    assert Objetivo.objects.filter(colaborador=julian).count() == 6


def test_la_ponderacion_no_se_pasa_de_cien(equipo):
    jefe, julian, _ = equipo
    objetivo_de(julian, jefe, '80')

    respuesta = cliente_de(jefe).post(
        f'{RUTA}/objetivos',
        {
            'colaborador': julian.pk,
            'periodo': OCTUBRE.isoformat(),
            'objetivo': 'Uno que se pasa',
            'kpi': 'KPI',
            'peso': '30',
            'tipoMedicion': 'binario',
        },
        format='json',
    )

    assert respuesta.status_code == 400
    assert 'Solo queda 20%' in str(respuesta.data)


def test_la_meta_es_obligatoria_salvo_en_binario(equipo):
    jefe, julian, _ = equipo
    cliente = cliente_de(jefe)
    base = {
        'colaborador': julian.pk,
        'periodo': OCTUBRE.isoformat(),
        'objetivo': 'Publicar tableros validados',
        'kpi': 'Tableros publicados',
        'peso': '25',
        'tipoMedicion': 'proporcional',
        'unidad': 'unidades',
    }

    sin_meta = cliente.post(f'{RUTA}/objetivos', base, format='json')
    assert sin_meta.status_code == 400

    en_cero = cliente.post(f'{RUTA}/objetivos', {**base, 'metaValor': '0'}, format='json')
    assert en_cero.status_code == 400

    con_meta = cliente.post(f'{RUTA}/objetivos', {**base, 'metaValor': '8'}, format='json')
    assert con_meta.status_code == 201, con_meta.data


def test_la_formula_se_valida_al_guardarla(equipo):
    jefe, julian, _ = equipo
    base = {
        'colaborador': julian.pk,
        'periodo': OCTUBRE.isoformat(),
        'objetivo': 'Índice compuesto',
        'kpi': 'Índice',
        'peso': '25',
        'tipoMedicion': 'formula',
        'metaValor': '10',
    }

    mala = cliente_de(jefe).post(
        f'{RUTA}/objetivos', {**base, 'formula': '__import__("os")'}, format='json'
    )
    assert mala.status_code == 400

    buena = cliente_de(jefe).post(
        f'{RUTA}/objetivos', {**base, 'formula': 'logrado / meta * 100'}, format='json'
    )
    assert buena.status_code == 201, buena.data


# ── Periodo (§6, reglas 1 y 5) ─────────────────────────────────────────────


def test_activar_el_mes_exige_que_todos_sumen_cien(app_performance, equipo):
    jefe, julian, daniel = equipo
    people = crear_usuario('people@supli.tech', app_performance, ['performance:periodos:manage'])
    objetivo_de(julian, jefe, '100')
    objetivo_de(daniel, jefe, '60')

    respuesta = cliente_de(people).post(f'{RUTA}/periodos/{MES}/activar')

    assert respuesta.status_code == 422
    datos = respuesta.json()
    assert datos['error'] == 'PESO_INCOMPLETO'
    assert [fila['colaboradorNombre'] for fila in datos['pendientes']] == ['Daniel']
    assert 'suman 60%' in datos['pendientes'][0]['mensaje']
    # No se activa a medias: nada quedó congelado.
    assert Objetivo.objects.filter(estado=EstadoObjetivo.CONGELADO).count() == 0


def test_al_activar_el_mes_los_objetivos_se_congelan(app_performance, equipo):
    jefe, julian, daniel = equipo
    people = crear_usuario('people@supli.tech', app_performance, ['performance:periodos:manage'])
    objetivo_de(julian, jefe, '100')
    objetivo_de(daniel, jefe, '40')
    objetivo_de(daniel, jefe, '60', objetivo='Otro objetivo')

    respuesta = cliente_de(people).post(f'{RUTA}/periodos/{MES}/activar')

    assert respuesta.status_code == 200, respuesta.data
    assert respuesta.json()['congelados'] == 3
    assert Periodo.objects.get(periodo=OCTUBRE).estado == EstadoPeriodo.EN_MEDICION
    assert Objetivo.objects.exclude(estado=EstadoObjetivo.CONGELADO).count() == 0

    # Regla 5: con el mes en medición ya no se edita ni se borra.
    objetivo = Objetivo.objects.filter(colaborador=julian).first()
    edicion = cliente_de(jefe).patch(
        f'{RUTA}/objetivos/{objetivo.pk}', {'peso': '50'}, format='json'
    )
    assert edicion.status_code == 400
    borrado = cliente_de(jefe).delete(f'{RUTA}/objetivos/{objetivo.pk}')
    assert borrado.status_code == 400
    assert Objetivo.objects.filter(pk=objetivo.pk).exists()


def test_activar_el_mes_es_de_people_no_de_cada_lider(equipo):
    jefe, julian, _ = equipo
    objetivo_de(julian, jefe, '100')

    respuesta = cliente_de(jefe).post(f'{RUTA}/periodos/{MES}/activar')

    assert respuesta.status_code == 403
    assert Objetivo.objects.filter(estado=EstadoObjetivo.CONGELADO).count() == 0


# ── Consulta (§6, regla 6 y §7) ────────────────────────────────────────────


def test_el_colaborador_consulta_los_suyos_en_solo_lectura(equipo):
    jefe, julian, daniel = equipo
    objetivo_de(julian, jefe, '100')
    objetivo_de(daniel, jefe, '100')

    respuesta = cliente_de(julian).get(f'{RUTA}/mis-objetivos?periodo={MES}')

    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert len(datos) == 1
    assert datos[0]['colaborador'] == julian.pk
    # Y no ve los de su compañero por el listado general.
    del_equipo = cliente_de(julian).get(f'{RUTA}/objetivos?colaborador={daniel.pk}').json()
    assert del_equipo == []


def test_el_jefe_ve_el_equipo_y_la_direccion_ve_todo(app_performance, equipo):
    jefe, julian, daniel = equipo
    ajeno = crear_usuario('ajeno@supli.tech', app_performance)
    objetivo_de(julian, jefe, '100')
    objetivo_de(ajeno, jefe, '100')
    ceo = crear_usuario('ceo@supli.tech', app_performance, ['performance:objetivos:view_all'])

    del_jefe = cliente_de(jefe).get(f'{RUTA}/objetivos?periodo={MES}').json()
    assert {fila['colaborador'] for fila in del_jefe} == {julian.pk}

    del_ceo = cliente_de(ceo).get(f'{RUTA}/objetivos?periodo={MES}').json()
    assert {fila['colaborador'] for fila in del_ceo} == {julian.pk, ajeno.pk}


def test_el_historico_no_se_pisa_al_cambiar_de_mes(equipo):
    """Regla 6: cada (persona, mes) es un registro aparte."""
    jefe, julian, _ = equipo
    objetivo_de(julian, jefe, '100')
    objetivo_de(julian, jefe, '100', periodo=mes_siguiente(OCTUBRE))

    octubre = cliente_de(jefe).get(f'{RUTA}/objetivos?colaborador={julian.pk}&periodo={MES}')
    mes_dos = f'{mes_siguiente(OCTUBRE):%Y-%m}'
    noviembre = cliente_de(jefe).get(
        f'{RUTA}/objetivos?colaborador={julian.pk}&periodo={mes_dos}'
    )

    assert len(octubre.json()) == 1
    assert len(noviembre.json()) == 1
    assert Objetivo.objects.filter(colaborador=julian).count() == 2


def test_el_resumen_dice_cuanto_falta_para_el_cien(equipo):
    jefe, julian, _ = equipo
    objetivo_de(julian, jefe, '30')
    objetivo_de(julian, jefe, '25', objetivo='Otro')

    datos = cliente_de(jefe).get(f'{RUTA}/periodos/{MES}/resumen?colaborador={julian.pk}').json()

    fila = datos['colaboradores'][0]
    assert (fila['pesoAsignado'], fila['pesoDisponible'], fila['completo']) == (55.0, 45.0, False)
    assert datos['estado'] == EstadoPeriodo.DEFINICION


def test_sin_el_submodulo_no_se_entra(app_performance):
    fuera = crear_usuario('fuera@supli.tech')

    for ruta in ('opciones', 'resumen', 'mis-objetivos', 'objetivos'):
        assert cliente_de(fuera).get(f'{RUTA}/{ruta}').status_code == 403


def test_las_opciones_traen_el_equipo_y_los_catalogos(equipo):
    jefe, julian, daniel = equipo

    datos = cliente_de(jefe).get(f'{RUTA}/opciones').json()

    assert {persona['id'] for persona in datos['equipo']} == {julian.pk, daniel.pk}
    assert datos['maximoObjetivos'] == 6
    assert len(datos['tiposMedicion']) == len(TipoMedicion.choices)
    assert datos['capacidades']['esLider'] is True


# ── Estructura del menú (§3) ───────────────────────────────────────────────


def test_el_seed_deja_valoracion_dentro_del_contenedor():
    Application.objects.create(
        code='valoracion', name='Supli performance', base_path='/inicio/valoracion', order=300
    )

    call_command('seed_performance_app', verbosity=0)

    contenedor = Application.objects.get(code='supli-performance')
    objetivos = Application.objects.get(code='objetivos-kpis')
    valoracion = Application.objects.get(code='valoracion')
    assert objetivos.parent == contenedor
    assert valoracion.parent == contenedor
    # La valoración conserva su ruta: quien ya tenía acceso sigue entrando igual.
    assert valoracion.base_path == '/inicio/valoracion'
    assert valoracion.name == 'Valoración'


def test_el_submodulo_arrastra_al_contenedor_en_el_menu(app_performance):
    persona = crear_usuario('menu@supli.tech', app_performance)

    apps = persona.get_accessible_applications()

    codigos = set(apps.values_list('code', flat=True))
    assert codigos == {'objetivos-kpis', 'supli-performance'}


# ── Decisiones cerradas por People (documento de definiciones) ─────────────


def test_el_semaforo_usa_los_cortes_de_people():
    """A2: verde ≥ 100, naranja 85–99.9, rojo < 85."""
    assert semaforo(120) == Semaforo.VERDE
    assert semaforo(100) == Semaforo.VERDE
    assert semaforo(99.9) == Semaforo.NARANJA
    assert semaforo(85) == Semaforo.NARANJA
    assert semaforo(84.99) == Semaforo.ROJO
    assert semaforo(0) == Semaforo.ROJO
    # Sin resultado cargado no hay color: no está en rojo, está sin medir.
    assert semaforo(None) is None


def test_la_meta_cualitativa_vale_cien_cincuenta_o_cero():
    """A2: 2 de 2 criterios = 100%, 1 de 2 = 50%, 0 de 2 = 0%."""
    objetivo = Objetivo(
        tipo_medicion=TipoMedicion.CUALITATIVA,
        meta_valor=Decimal('2'),
        tope_cumplimiento=Decimal('100'),
        permite_sobrecumplimiento=False,
    )
    assert calcular_cumplimiento(objetivo, 2) == Decimal('100.00')
    assert calcular_cumplimiento(objetivo, 1) == Decimal('50.00')
    assert calcular_cumplimiento(objetivo, 0) == Decimal('0.00')


def test_la_cualitativa_se_guarda_con_dos_criterios_sin_pedir_meta(equipo):
    jefe, julian, _daniel = equipo

    respuesta = cliente_de(jefe).post(
        f'{RUTA}/objetivos',
        {
            'colaborador': julian.pk,
            'periodo': OCTUBRE.isoformat(),
            'objetivo': 'Documentar el modelo de datos',
            'kpi': 'Documento publicado y socializado',
            'peso': '100',
            'tipoMedicion': TipoMedicion.CUALITATIVA,
            'unidad': '',
            'metaValor': None,
        },
        format='json',
    )

    assert respuesta.status_code == 201, respuesta.data
    objetivo = Objetivo.objects.get()
    assert objetivo.meta_valor == Decimal('2.00')
    # Una meta cualitativa nunca sobrecumple: dos de dos ya es todo.
    assert objetivo.permite_sobrecumplimiento is False


def test_el_resultado_cualitativo_no_pasa_de_dos_criterios(equipo):
    jefe, julian, _daniel = equipo
    objetivo = objetivo_de(
        julian, jefe, '100', tipo_medicion=TipoMedicion.CUALITATIVA, meta_valor=Decimal('2')
    )
    _empezar_el_mes(objetivo.periodo)

    respuesta = cliente_de(julian).put(
        f'{RUTA}/objetivos/{objetivo.pk}/resultado',
        {'resultadoEjecutado': '3', 'evidencias': [{'linkSoporte': 'https://supli.sharepoint.com/x'}]},
        format='json',
    )

    assert respuesta.status_code == 400
    assert 'criterios' in str(respuesta.data).lower()


# ── Congelamiento por calendario (A9) ──────────────────────────────────────


def _empezar_el_mes(periodo):
    """Corre el reloj: deja el mes del objetivo en el pasado para las pruebas."""
    Objetivo.objects.filter(periodo=periodo).update(periodo=primer_dia(date.today()))
    Periodo.objects.filter(periodo=periodo).update(periodo=primer_dia(date.today()))
    return primer_dia(date.today())


def test_los_objetivos_del_mes_en_curso_estan_congelados(equipo):
    """A9: se editan hasta el último día del mes anterior; después, no."""
    jefe, julian, _daniel = equipo
    objetivo = objetivo_de(julian, jefe, '100')
    mes_en_curso = _empezar_el_mes(objetivo.periodo)
    objetivo.refresh_from_db()

    respuesta = cliente_de(jefe).patch(
        f'{RUTA}/objetivos/{objetivo.pk}', {'peso': '50'}, format='json'
    )

    assert respuesta.status_code == 400
    assert 'congelados' in str(respuesta.data)
    assert objetivo.periodo == mes_en_curso
    # Y tampoco se puede borrar.
    assert cliente_de(jefe).delete(f'{RUTA}/objetivos/{objetivo.pk}').status_code == 400


def test_el_mes_entrante_si_se_edita(equipo):
    jefe, julian, _daniel = equipo
    objetivo = objetivo_de(julian, jefe, '100')

    respuesta = cliente_de(jefe).patch(
        f'{RUTA}/objetivos/{objetivo.pk}', {'peso': '60'}, format='json'
    )

    assert respuesta.status_code == 200, respuesta.data
    assert respuesta.json()['editable'] is True


def test_people_habilita_la_edicion_de_un_mes_congelado(app_performance, equipo):
    """A9: la excepción autorizada por el CEO, con registro de quién y por qué."""
    jefe, julian, _daniel = equipo
    people = crear_usuario(
        'people@supli.tech', app_performance, ['performance:periodos:manage']
    )
    objetivo = objetivo_de(julian, jefe, '100')
    mes = _empezar_el_mes(objetivo.periodo)

    sin_permiso = cliente_de(jefe).post(
        f'{RUTA}/periodos/{mes:%Y-%m}/edicion', {'motivo': 'porque sí'}, format='json'
    )
    assert sin_permiso.status_code == 403

    sin_motivo = cliente_de(people).post(
        f'{RUTA}/periodos/{mes:%Y-%m}/edicion', {}, format='json'
    )
    assert sin_motivo.status_code == 400

    abierto = cliente_de(people).post(
        f'{RUTA}/periodos/{mes:%Y-%m}/edicion',
        {'motivo': 'Reestructuración de octubre, autorizada por el CEO.'},
        format='json',
    )
    assert abierto.status_code == 200, abierto.data

    registro = Periodo.objects.get(periodo=mes)
    assert registro.edicion_habilitada is True
    assert registro.habilitada_por == people
    assert registro.fecha_habilitacion is not None
    assert 'CEO' in registro.motivo_habilitacion

    # Con la excepción abierta, el jefe vuelve a editar.
    editado = cliente_de(jefe).patch(
        f'{RUTA}/objetivos/{objetivo.pk}', {'peso': '80'}, format='json'
    )
    assert editado.status_code == 200, editado.data

    # Y People la vuelve a cerrar.
    cerrado = cliente_de(people).post(
        f'{RUTA}/periodos/{mes:%Y-%m}/edicion', {'habilitada': False}, format='json'
    )
    assert cerrado.status_code == 200
    assert Periodo.objects.get(periodo=mes).edicion_habilitada is False


# ── Resultados y evidencias (A5) ───────────────────────────────────────────


def test_el_responsable_carga_el_resultado_y_el_sistema_calcula_el_porcentaje(equipo):
    jefe, julian, _daniel = equipo
    objetivo = objetivo_de(
        julian,
        jefe,
        '100',
        tipo_medicion=TipoMedicion.PROPORCIONAL,
        unidad=Unidad.UNIDADES,
        meta_valor=Decimal('10'),
    )
    _empezar_el_mes(objetivo.periodo)

    respuesta = cliente_de(julian).put(
        f'{RUTA}/objetivos/{objetivo.pk}/resultado',
        {
            'resultadoEjecutado': '9',
            'evidencias': [
                {'nombre': 'Tablero', 'linkSoporte': 'https://supli.sharepoint.com/tablero'}
            ],
        },
        format='json',
    )

    assert respuesta.status_code == 200, respuesta.data
    datos = respuesta.json()
    # El porcentaje no se digita: lo calcula el motor, y de ahí sale el color.
    assert datos['porcentajeCumplimiento'] == '90.00'
    assert datos['semaforo'] == Semaforo.NARANJA
    assert datos['estadoValidacion'] == EstadoValidacion.PENDIENTE
    assert datos['evidencias'][0]['linkSoporte'] == 'https://supli.sharepoint.com/tablero'


def test_el_resultado_lo_carga_el_responsable_no_cualquiera(app_performance, equipo):
    """Para los asesores de PDV el responsable es su Trade Leader, no ellos."""
    jefe, julian, daniel = equipo
    objetivo = objetivo_de(julian, jefe, '100', responsable_resultado=jefe)
    _empezar_el_mes(objetivo.periodo)

    ajeno = cliente_de(daniel).put(
        f'{RUTA}/objetivos/{objetivo.pk}/resultado', {'resultadoEjecutado': '1'}, format='json'
    )
    assert ajeno.status_code == 403

    del_responsable = cliente_de(jefe).put(
        f'{RUTA}/objetivos/{objetivo.pk}/resultado', {'resultadoEjecutado': '1'}, format='json'
    )
    assert del_responsable.status_code == 200, del_responsable.data


def test_no_se_carga_el_resultado_de_un_mes_que_no_ha_empezado(equipo):
    jefe, julian, _daniel = equipo
    objetivo = objetivo_de(julian, jefe, '100')

    respuesta = cliente_de(julian).put(
        f'{RUTA}/objetivos/{objetivo.pk}/resultado', {'resultadoEjecutado': '1'}, format='json'
    )

    assert respuesta.status_code == 400
    assert 'todavía no empieza' in str(respuesta.data)


def test_el_jefe_valida_el_resultado_y_quien_lo_cargo_no(equipo):
    jefe, julian, _daniel = equipo
    objetivo = objetivo_de(julian, jefe, '100')
    _empezar_el_mes(objetivo.periodo)
    cliente_de(julian).put(
        f'{RUTA}/objetivos/{objetivo.pk}/resultado', {'resultadoEjecutado': '1'}, format='json'
    )

    propio = cliente_de(julian).post(
        f'{RUTA}/objetivos/{objetivo.pk}/validar', {'estado': 'validado'}, format='json'
    )
    assert propio.status_code == 403

    sin_razon = cliente_de(jefe).post(
        f'{RUTA}/objetivos/{objetivo.pk}/validar', {'estado': 'rechazado'}, format='json'
    )
    assert sin_razon.status_code == 400

    validado = cliente_de(jefe).post(
        f'{RUTA}/objetivos/{objetivo.pk}/validar',
        {'estado': 'validado', 'observacion': 'Soporte revisado.'},
        format='json',
    )
    assert validado.status_code == 200
    assert validado.json()['estadoValidacion'] == EstadoValidacion.VALIDADO
    assert validado.json()['validadoPorNombre'] == jefe.full_name


def test_corregir_el_resultado_devuelve_la_validacion_a_pendiente(equipo):
    jefe, julian, _daniel = equipo
    objetivo = objetivo_de(julian, jefe, '100')
    _empezar_el_mes(objetivo.periodo)
    cliente = cliente_de(julian)
    cliente.put(
        f'{RUTA}/objetivos/{objetivo.pk}/resultado', {'resultadoEjecutado': '1'}, format='json'
    )
    cliente_de(jefe).post(
        f'{RUTA}/objetivos/{objetivo.pk}/validar', {'estado': 'validado'}, format='json'
    )

    corregido = cliente.put(
        f'{RUTA}/objetivos/{objetivo.pk}/resultado', {'resultadoEjecutado': '0'}, format='json'
    )

    assert corregido.json()['estadoValidacion'] == EstadoValidacion.PENDIENTE
    assert corregido.json()['porcentajeCumplimiento'] == '0.00'


# ── Cortes de tiempo (A8) ──────────────────────────────────────────────────


def test_el_q_es_trimestral_no_de_cuatro_meses():
    assert rango('trimestre', 2026, 1) == (date(2026, 1, 1), date(2026, 3, 1))
    assert rango('trimestre', 2026, 3) == (date(2026, 7, 1), date(2026, 9, 1))
    assert rango('trimestre', 2026, 4) == (date(2026, 10, 1), date(2026, 12, 1))
    assert rango('semestre', 2026, 2) == (date(2026, 7, 1), date(2026, 12, 1))
    assert rango('anio', 2026) == (date(2026, 1, 1), date(2026, 12, 1))
    assert etiqueta_corte('trimestre', 2026, 3) == '3Q · Jul–Sep 2026'


def test_el_acumulado_pondera_lo_medido_y_pinta_el_semaforo(equipo):
    jefe, julian, _daniel = equipo
    mes = primer_dia(date.today())
    dos = objetivo_de(
        julian, jefe, '50', periodo=mes, tipo_medicion=TipoMedicion.PROPORCIONAL,
        meta_valor=Decimal('10'),
    )
    objetivo_de(
        julian, jefe, '50', periodo=mes, tipo_medicion=TipoMedicion.PROPORCIONAL,
        meta_valor=Decimal('10'),
    )
    cliente_de(julian).put(
        f'{RUTA}/objetivos/{dos.pk}/resultado', {'resultadoEjecutado': '8'}, format='json'
    )

    datos = cliente_de(jefe).get(
        f'{RUTA}/acumulado?tipo=trimestre&anio={mes.year}&indice={(mes.month - 1) // 3 + 1}'
    ).json()

    fila = next(f for f in datos['colaboradores'] if f['colaborador'] == julian.pk)
    assert fila['objetivos'] == 2
    assert fila['medidos'] == 1
    # Se pondera sobre lo medido: el objetivo sin resultado no cuenta como cero.
    assert fila['cumplimiento'] == 80.0
    assert fila['semaforo'] == Semaforo.ROJO
    assert datos['cortes']['verdeDesde'] == 100.0
    assert datos['label'].endswith(str(mes.year))


# ── Carga de objetivos por Excel ───────────────────────────────────────────


ENCABEZADOS_EXCEL = (
    'objetivo',
    'kpi',
    'tipo_medicion',
    'unidad',
    'meta',
    'peso',
    'umbral_cumplimiento',
    'permite_sobrecumplimiento',
    'formula',
    'fuente_datos',
    'responsable_correo',
)


def _xlsx(filas, encabezados=ENCABEZADOS_EXCEL):
    """Arma en memoria un .xlsx como el que subiría una persona."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Datos'
    hoja.append(list(encabezados))
    for fila in filas:
        hoja.append(list(fila))
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    buffer.name = 'objetivos.xlsx'
    return buffer


def _fila(objetivo='Entregar el portal', peso='50', tipo='proporcional', meta='10', **extra):
    valores = {
        'objetivo': objetivo,
        'kpi': 'Portal en producción',
        'tipo_medicion': tipo,
        'unidad': 'porcentaje',
        'meta': meta,
        'peso': peso,
        'umbral_cumplimiento': None,
        'permite_sobrecumplimiento': 'No',
        'formula': None,
        'fuente_datos': 'Tablero de BI',
        'responsable_correo': None,
    }
    valores.update(extra)
    return tuple(valores[columna] for columna in ENCABEZADOS_EXCEL)


def _importar(cliente, colaborador, filas, modo='agregar'):
    return cliente.post(
        f'{RUTA}/objetivos/importar',
        {
            'archivo': _xlsx(filas),
            'colaborador': colaborador.pk,
            'periodo': OCTUBRE.isoformat(),
            'modo': modo,
        },
        format='multipart',
    )


def test_la_plantilla_de_objetivos_trae_encabezados_e_instrucciones(equipo):
    jefe, _julian, _daniel = equipo

    respuesta = cliente_de(jefe).get(f'{RUTA}/objetivos/plantilla')

    assert respuesta.status_code == 200
    assert respuesta['Content-Disposition'].endswith('.xlsx"')
    libro = openpyxl.load_workbook(BytesIO(respuesta.content))
    assert libro.sheetnames == ['Datos', 'Instrucciones']
    assert [celda.value for celda in libro['Datos'][1]] == list(ENCABEZADOS_EXCEL)
    texto = ' '.join(
        str(celda.value)
        for fila in libro['Instrucciones'].iter_rows()
        for celda in fila
        if celda.value
    )
    assert 'deben sumar 100' in texto
    assert 'proporcional_inverso' in texto


def test_el_jefe_carga_los_objetivos_de_su_equipo_por_excel(equipo):
    jefe, julian, _daniel = equipo

    respuesta = _importar(
        cliente_de(jefe),
        julian,
        [
            _fila('Entregar el portal', peso='60'),
            _fila('Publicar tableros', peso='40', tipo='binario', meta=None, unidad='si_no'),
        ],
    )

    assert respuesta.status_code == 200, respuesta.data
    datos = respuesta.json()
    assert datos['created'] == 2
    assert datos['completo'] is True
    assert datos['pesoAsignado'] == 100.0
    objetivos = Objetivo.objects.filter(colaborador=julian, periodo=OCTUBRE)
    assert objetivos.count() == 2
    # El registrado_por es quien sube el archivo, no quien lo llenó.
    assert objetivos.first().registrado_por == jefe
    # El binario no guarda meta, como en el formulario.
    assert objetivos.get(tipo_medicion='binario').meta_valor is None


def test_el_excel_no_pasa_de_cien_entre_todas_las_filas(equipo):
    """Fila por fila cada peso es válido; lo que no puede pasar es la suma."""
    jefe, julian, _daniel = equipo

    respuesta = _importar(
        cliente_de(jefe),
        julian,
        [_fila('Uno', peso='60'), _fila('Dos', peso='60')],
    )

    assert respuesta.status_code == 400
    assert '120' in str(respuesta.data)
    assert Objetivo.objects.count() == 0


def test_el_excel_cuenta_lo_que_la_persona_ya_tenia(equipo):
    jefe, julian, _daniel = equipo
    objetivo_de(julian, jefe, '70')

    respuesta = _importar(cliente_de(jefe), julian, [_fila('Nuevo', peso='40')])

    assert respuesta.status_code == 400
    # El aviso dice cuánto queda, no solo que no cabe.
    assert 'Solo queda 30%' in str(respuesta.data['filas'][0]['errores'])
    assert Objetivo.objects.count() == 1


def test_el_modo_reemplazar_deja_solo_lo_del_archivo(equipo):
    jefe, julian, _daniel = equipo
    objetivo_de(julian, jefe, '70')

    respuesta = _importar(
        cliente_de(jefe), julian, [_fila('Nuevo', peso='100')], modo='reemplazar'
    )

    assert respuesta.status_code == 200, respuesta.data
    assert respuesta.json()['deleted'] == 1
    assert [o.objetivo for o in Objetivo.objects.all()] == ['Nuevo']


def test_una_fila_mala_cancela_toda_la_carga(equipo):
    jefe, julian, _daniel = equipo

    respuesta = _importar(
        cliente_de(jefe),
        julian,
        [
            _fila('Buena', peso='50'),
            _fila('Sin meta', peso='50', tipo='proporcional', meta=None),
        ],
    )

    assert respuesta.status_code == 400
    assert respuesta.data['filas'][0]['fila'] == 3
    assert 'meta' in str(respuesta.data['filas'][0]['errores']).lower()
    assert Objetivo.objects.count() == 0


def test_el_excel_avisa_si_el_tipo_de_medicion_no_existe(equipo):
    jefe, julian, _daniel = equipo

    respuesta = _importar(cliente_de(jefe), julian, [_fila('Uno', tipo='a-ojo')])

    assert respuesta.status_code == 400
    assert 'tipo_medicion' in str(respuesta.data)


def test_el_responsable_del_resultado_llega_por_correo(equipo):
    """Para asesores y promotores, el resultado lo carga su Trade Leader."""
    jefe, julian, _daniel = equipo

    inexistente = _importar(
        cliente_de(jefe), julian, [_fila('Uno', responsable_correo='nadie@supli.tech')]
    )
    assert inexistente.status_code == 400
    assert 'nadie@supli.tech' in str(inexistente.data)

    respuesta = _importar(
        cliente_de(jefe), julian, [_fila('Uno', peso='100', responsable_correo=jefe.email)]
    )

    assert respuesta.status_code == 200, respuesta.data
    assert Objetivo.objects.get().responsable_resultado == jefe


def test_nadie_carga_por_excel_los_objetivos_de_otro_equipo(app_performance, equipo):
    jefe, julian, _daniel = equipo
    ajeno = crear_usuario('ajeno@supli.tech', app_performance)

    respuesta = _importar(cliente_de(ajeno), julian, [_fila('Uno')])

    assert respuesta.status_code == 403
    assert Objetivo.objects.count() == 0


def test_no_se_carga_el_excel_de_un_mes_congelado(app_performance, equipo):
    jefe, julian, _daniel = equipo
    cliente = cliente_de(jefe)
    # El mes en curso ya está congelado: se define hasta el último día del anterior.
    respuesta = cliente.post(
        f'{RUTA}/objetivos/importar',
        {
            'archivo': _xlsx([_fila('Uno')]),
            'colaborador': julian.pk,
            'periodo': primer_dia(date.today()).isoformat(),
        },
        format='multipart',
    )

    assert respuesta.status_code == 400
    assert 'congelados' in str(respuesta.data)


def test_la_plantilla_se_llena_y_se_sube_tal_cual(equipo):
    """El ciclo real: descargar, borrar el ejemplo, llenar y subir."""
    jefe, julian, _daniel = equipo
    cliente = cliente_de(jefe)

    descargada = cliente.get(f'{RUTA}/objetivos/plantilla')
    libro = openpyxl.load_workbook(BytesIO(descargada.content))
    hoja = libro['Datos']
    hoja.delete_rows(2)  # la fila de ejemplo
    hoja.append(
        ['Entregar el portal', 'Portal en producción', 'proporcional', 'porcentaje',
         90, 100, None, 'No', None, 'Tablero de BI', None]
    )
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    buffer.name = 'plantilla-llena.xlsx'

    respuesta = cliente.post(
        f'{RUTA}/objetivos/importar',
        {'archivo': buffer, 'colaborador': julian.pk, 'periodo': OCTUBRE.isoformat()},
        format='multipart',
    )

    assert respuesta.status_code == 200, respuesta.data
    objetivo = Objetivo.objects.get()
    assert objetivo.meta_valor == Decimal('90.00')
    assert objetivo.peso == Decimal('100.00')
