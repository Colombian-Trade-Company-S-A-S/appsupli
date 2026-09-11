"""
Lector del «Informe Inventarios Tecnología Coltrade».

Es el archivo que sale del ERP, no una plantilla nuestra: trae ventas e
inventario mezclados en la misma hoja y muchas columnas que aquí no se usan.
Lo que los separa es la clase de movimiento (`CMv`):

  · `601` → una venta. Trae fecha contable.
  · `1`   → una existencia en bodega. No trae fecha: es el stock de hoy.

Y algo que es fácil de pasar por alto: `601A` **también** dice «Venta» en la
columna de etiqueta, pero es de otro periodo y no entra. Solo `601`.

Cada fila es **una unidad**: el informe lista número de serie por número de
serie, así que la cantidad sale de contar filas, no de leer una columna.
"""
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

#: La hoja con los datos. Se busca sin distinguir mayúsculas («Base», «base»).
NOMBRE_HOJA = 'base'

#: Clases de movimiento que interesan.
CMV_VENTA = '601'
CMV_INVENTARIO = '1'

#: Encabezado de cada dato que se necesita. Se comparan normalizados, así que
#: los espacios de sobra del archivo («Ce. ») no importan.
COLUMNA_MATERIAL = 'Material'
COLUMNA_CENTRO = 'Ce.'
COLUMNA_FECHA = 'Fe/contab/'
COLUMNA_CLASE = 'CMv'
COLUMNAS = (COLUMNA_MATERIAL, COLUMNA_CENTRO, COLUMNA_FECHA, COLUMNA_CLASE)


class ErrorDeInforme(Exception):
    """El archivo no es el informe esperado. Lleva el mensaje para la persona."""


def _normalizar_encabezado(valor) -> str:
    return str(valor or '').strip().casefold()


def _texto(valor) -> str:
    """
    Un código tal como lo escribiría una persona.

    Excel devuelve los códigos numéricos como float (`7023988.0`), y ese `.0`
    haría que no cruce nunca con el código guardado en la base.
    """
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, int):
        return str(valor)
    return str(valor).strip()


def _fecha(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return None


@dataclass
class Lectura:
    """Lo que trae el archivo, ya separado y contado."""

    #: `(id_producto, id_punto_venta, fecha)` → unidades vendidas.
    ventas: Counter = field(default_factory=Counter)
    #: `(id_producto, id_punto_venta)` → unidades en inventario.
    inventario: Counter = field(default_factory=Counter)
    #: Filas leídas de cada clase, antes de descartar nada.
    filas_venta: int = 0
    filas_inventario: int = 0
    #: Filas de otras clases de movimiento («601A», «7»…).
    filas_ignoradas: int = 0
    #: Filas sin producto, sin centro o sin fecha legible.
    filas_incompletas: int = 0

    @property
    def dias_con_venta(self) -> set[date]:
        return {fecha for _, _, fecha in self.ventas}


def leer_informe(archivo) -> Lectura:
    """
    Lee la hoja «base» y devuelve las ventas y el inventario que trae.

    No mira todavía si el producto o el punto de venta existen: eso lo decide
    quien importa, que es el que sabe qué hay en la base.
    """
    try:
        # `read_only` para no cargar 11.000 filas a memoria y `data_only` para
        # leer el resultado de las fórmulas, no su texto.
        libro = load_workbook(archivo, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl lanza de todo
        raise ErrorDeInforme(
            'No se pudo abrir el archivo. Debe ser un Excel .xlsx sin contraseña.'
        ) from exc

    hoja = next(
        (libro[nombre] for nombre in libro.sheetnames if nombre.strip().casefold() == NOMBRE_HOJA),
        None,
    )
    if hoja is None:
        disponibles = ', '.join(libro.sheetnames) or 'ninguna'
        raise ErrorDeInforme(
            f'El archivo no tiene una hoja llamada «base». Hojas encontradas: {disponibles}.'
        )

    filas = hoja.iter_rows(values_only=True)
    try:
        encabezado = next(filas)
    except StopIteration as exc:
        raise ErrorDeInforme('La hoja «base» está vacía.') from exc

    posicion = {_normalizar_encabezado(nombre): indice for indice, nombre in enumerate(encabezado)}
    indices = {}
    faltantes = []
    for columna in COLUMNAS:
        indice = posicion.get(_normalizar_encabezado(columna))
        if indice is None:
            faltantes.append(columna)
        else:
            indices[columna] = indice
    if faltantes:
        raise ErrorDeInforme(
            'A la hoja «base» le faltan estas columnas: ' + ', '.join(f'«{c}»' for c in faltantes)
        )

    lectura = Lectura()
    for fila in filas:
        clase = _texto(fila[indices[COLUMNA_CLASE]])
        if clase not in (CMV_VENTA, CMV_INVENTARIO):
            lectura.filas_ignoradas += 1
            continue

        producto = _texto(fila[indices[COLUMNA_MATERIAL]])
        centro = _texto(fila[indices[COLUMNA_CENTRO]])

        if clase == CMV_VENTA:
            lectura.filas_venta += 1
            fecha = _fecha(fila[indices[COLUMNA_FECHA]])
            if not producto or not centro or fecha is None:
                lectura.filas_incompletas += 1
                continue
            lectura.ventas[(producto, centro, fecha)] += 1
        else:
            lectura.filas_inventario += 1
            if not producto or not centro:
                lectura.filas_incompletas += 1
                continue
            lectura.inventario[(producto, centro)] += 1

    return lectura
