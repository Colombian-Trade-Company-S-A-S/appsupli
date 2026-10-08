"""
Carga los puntos de venta del plan Belkin desde «INFORMACION CLARO.xlsx», el
archivo con el que Trade alimentaba el Power BI del plan.

    python manage.py cargar_catalogo_belkin --archivo "…/INFORMACION CLARO.xlsx" [--simular]

De la tabla «Puntos_Venta» toma, por centro de costos:

  · La regional. «Zona Norte» y «Zona Sur» son de Coltrade; «Cav Fuera de
    Coltrade» queda sin regional: no usa el formulario y sus registros llegan
    del informe de ventas de Claro.
  · La categoría del bono (1, 2 o 3).
  · El asesor Apple. «Agregar» es un hueco del archivo, no un nombre.

Al punto que ya existe le actualiza la regional y la categoría, no el nombre:
el de la app es el que ya se ve en el formulario. Al que falta lo crea con el
nombre del archivo, sin el código que algunos traen pegado al final.

De la tabla «Productos» crea los que falten, inactivos: cuentan para el
informe y para las ventas, pero no salen en el formulario hasta que alguien los
active en la lista. Lo que ya está en la lista no se toca.

No borra nada y se puede correr las veces que haga falta.
"""
import re

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.bi_trade.models import (
    AsesorApple,
    CategoriaBelkin,
    ProductoBelkin,
    PuntoVentaBelkin,
    RegionalBelkin,
)

FUERA_DE_COLTRADE = 'cav fuera de coltrade'
SIN_ASESOR = {'', 'agregar'}

#: El «Tipo» del archivo → la categoría de la lista. La que no exista se crea
#: inactiva, como los productos.
CATEGORIAS = {
    'case apple': 'Case Apple',
    'lamina belkin': 'Lámina',
    'cable belkin': 'Cable',
    'cargador': 'Cargador',
    'powerbank': 'Power Bank',
}


def _texto(valor) -> str:
    return ' '.join(str(valor or '').split())


def leer_tabla(libro, nombre: str) -> list[dict]:
    """Las filas de una tabla de Excel con nombre, como diccionarios."""
    for hoja in libro.worksheets:
        if nombre in hoja.tables:
            celdas = hoja[hoja.tables[nombre].ref]
            encabezado = [_texto(c.value) for c in celdas[0]]
            return [
                dict(zip(encabezado, (c.value for c in fila), strict=True)) for fila in celdas[1:]
            ]
    raise CommandError(f'El archivo no tiene la tabla «{nombre}».')


class Command(BaseCommand):
    help = 'Carga puntos, categorías, asesores y productos del plan Belkin desde INFORMACION CLARO.'

    def add_arguments(self, parser):
        parser.add_argument('--archivo', required=True, help='Ruta de INFORMACION CLARO.xlsx.')
        parser.add_argument(
            '--simular', action='store_true', help='Muestra lo que haría y no guarda nada.'
        )

    def handle(self, *args, archivo, simular, **options):
        try:
            libro = load_workbook(archivo, data_only=True)
        except Exception as exc:  # noqa: BLE001 — openpyxl lanza de todo
            raise CommandError(f'No se pudo abrir {archivo}: {exc}') from exc
        puntos = leer_tabla(libro, 'Puntos_Venta')
        productos = leer_tabla(libro, 'Productos')

        with transaction.atomic():
            self._puntos(puntos)
            self._productos(productos)
            if simular:
                transaction.set_rollback(True)
                self.stdout.write(self.style.WARNING('Simulación: no se guardó nada.'))
            else:
                self.stdout.write(self.style.SUCCESS('Catálogo del plan Belkin cargado.'))

    def _puntos(self, filas: list[dict]) -> None:
        creados = actualizados = asesores = 0
        for fila in filas:
            codigo = _texto(fila.get('Centro_Costos')).upper()
            if not codigo:
                continue
            nombre_regional = _texto(fila.get('Regional'))
            regional = None
            if nombre_regional.casefold() != FUERA_DE_COLTRADE:
                regional, _ = RegionalBelkin.objects.get_or_create(nombre=nombre_regional)
            categoria = fila.get('Categoria')
            categoria = int(categoria) if categoria in (1, 2, 3, '1', '2', '3') else None

            punto = PuntoVentaBelkin.objects.filter(pk=codigo).first()
            if punto is None:
                # «Cav Armenia Portal Quindio C304» → «Cav Armenia Portal Quindio».
                nombre = re.sub(rf'\s*{re.escape(codigo)}$', '', _texto(fila.get('Punto_Venta')))
                punto = PuntoVentaBelkin.objects.create(
                    id_punto_venta=codigo,
                    nombre_pdv=nombre or codigo,
                    id_regional=regional,
                    categoria=categoria,
                )
                creados += 1
            elif punto.id_regional != regional or punto.categoria != categoria:
                punto.id_regional = regional
                punto.categoria = categoria
                punto.save(update_fields=['id_regional', 'categoria', 'updated_at'])
                actualizados += 1

            asesor = _texto(fila.get('Asesor_Apple'))
            if asesor.casefold() not in SIN_ASESOR:
                _, nuevo = AsesorApple.objects.get_or_create(
                    nombre__iexact=asesor,
                    id_punto_venta=punto,
                    defaults={'nombre': asesor},
                )
                asesores += nuevo

            self.stdout.write(
                f'  {codigo:<6} {punto.nombre_pdv:<40} '
                f'{(regional.nombre if regional else "Fuera de Coltrade"):<18} '
                f'cat. {categoria or "-"}  {asesor}'
            )
        self.stdout.write(
            f'Puntos de venta: {creados} nuevos, {actualizados} actualizados. '
            f'Asesores nuevos: {asesores}.'
        )

    def _productos(self, filas: list[dict]) -> None:
        categorias = {}
        nuevos = []
        for fila in filas:
            codigo = _texto(fila.get('Material'))
            if not codigo or ProductoBelkin.objects.filter(pk=codigo).exists():
                continue
            tipo = _texto(fila.get('Tipo'))
            nombre_categoria = CATEGORIAS.get(tipo.casefold(), tipo)
            if nombre_categoria not in categorias:
                categoria = CategoriaBelkin.objects.filter(nombre__iexact=nombre_categoria).first()
                if categoria is None:
                    categoria = CategoriaBelkin.objects.create(
                        nombre=nombre_categoria, activa=False
                    )
                    self.stdout.write(f'  Categoría nueva (inactiva): {nombre_categoria}')
                categorias[nombre_categoria] = categoria
            producto = ProductoBelkin.objects.create(
                id_producto=codigo,
                nombre_producto=_texto(fila.get('Producto'))[:120],
                id_categoria=categorias[nombre_categoria],
                activo=False,
            )
            nuevos.append(producto)
            self.stdout.write(
                f'  {codigo}  {producto.nombre_producto:<50} {nombre_categoria} (inactivo)'
            )
        self.stdout.write(f'Productos nuevos: {len(nuevos)}.')
