"""
Plantillas e importación en Excel para los CRUD de BI Trade.

Cada viewset declara sus columnas una sola vez y de ahí salen las dos cosas:
la plantilla que se descarga y el lector que valida lo que se sube. Así la
plantilla nunca se desincroniza de lo que el importador acepta.
"""
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Encabezado de la plantilla: gris oscuro con letra blanca.
RELLENO_ENCABEZADO = PatternFill('solid', fgColor='1F2937')
LETRA_ENCABEZADO = Font(color='FFFFFF', bold=True)


def _a_bytes(libro: Workbook) -> bytes:
    """Serializa el libro en memoria: nunca se escribe un archivo en disco."""
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


@dataclass
class Columna:
    """Una columna de la plantilla y su regla de lectura."""

    nombre: str
    tipo: str = 'texto'  # texto | entero | fecha | opcion
    obligatoria: bool = True
    opciones: list[str] = field(default_factory=list)
    ayuda: str = ''
    ejemplo: str = ''

    @property
    def descripcion_tipo(self) -> str:
        return {
            'texto': 'Texto',
            'entero': 'Número entero',
            'fecha': 'Fecha (AAAA-MM-DD)',
            'opcion': 'Una de las opciones listadas',
        }[self.tipo]


class ErrorDeFila(Exception):
    """Una celda que no se pudo leer. Lleva el mensaje que verá la persona."""


def _leer_celda(valor, columna: Columna):
    """Convierte el contenido de una celda al tipo que espera el modelo."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        if columna.obligatoria:
            raise ErrorDeFila(f'«{columna.nombre}» es obligatorio.')
        return None

    if columna.tipo == 'entero':
        try:
            # Excel suele devolver los números como float (12.0).
            entero = int(float(valor))
        except (TypeError, ValueError) as exc:
            raise ErrorDeFila(f'«{columna.nombre}» debe ser un número entero.') from exc
        if entero < 0:
            raise ErrorDeFila(f'«{columna.nombre}» no puede ser negativo.')
        return entero

    if columna.tipo == 'fecha':
        if isinstance(valor, datetime):
            return valor.date()
        if isinstance(valor, date):
            return valor
        for formato in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
            try:
                return datetime.strptime(str(valor).strip(), formato).date()
            except ValueError:
                continue
        raise ErrorDeFila(
            f'«{columna.nombre}» no se entiende como fecha. Usa el formato AAAA-MM-DD.'
        )

    texto = str(valor).strip()

    if columna.tipo == 'opcion' and texto not in columna.opciones:
        raise ErrorDeFila(
            f'«{columna.nombre}»: «{texto}» no es un valor válido. '
            f'Opciones: {", ".join(columna.opciones)}.'
        )

    return texto


def construir_plantilla(titulo: str, columnas: list[Columna], nota: str = '') -> bytes:
    """Arma el .xlsx de la plantilla: una hoja de datos y una de instrucciones."""
    libro = Workbook()

    hoja = libro.active
    hoja.title = 'Datos'

    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=indice, value=columna.nombre)
        celda.fill = RELLENO_ENCABEZADO
        celda.font = LETRA_ENCABEZADO
        celda.alignment = Alignment(horizontal='center', vertical='center')
        hoja.column_dimensions[get_column_letter(indice)].width = max(len(columna.nombre) + 6, 18)

        # Fila de ejemplo, para que se vea el formato esperado.
        if columna.ejemplo:
            hoja.cell(row=2, column=indice, value=columna.ejemplo)

        # Desplegable con los valores válidos de las columnas con opciones.
        if columna.tipo == 'opcion' and columna.opciones:
            validacion = DataValidation(
                type='list',
                formula1='"{}"'.format(','.join(columna.opciones)),
                allow_blank=not columna.obligatoria,
                showDropDown=False,
            )
            hoja.add_data_validation(validacion)
            letra = get_column_letter(indice)
            validacion.add(f'{letra}2:{letra}1000')

    hoja.freeze_panes = 'A2'

    guia = libro.create_sheet('Instrucciones')
    guia.column_dimensions['A'].width = 28
    guia.column_dimensions['B'].width = 18
    guia.column_dimensions['C'].width = 14
    guia.column_dimensions['D'].width = 70

    guia['A1'] = titulo
    guia['A1'].font = Font(bold=True, size=13)
    guia['A3'] = (
        'Llena la hoja «Datos» y súbela con el botón «Importar Excel». '
        'La fila 1 son los encabezados: no la borres ni la renombres. '
        'La fila 2 es un ejemplo: bórrala antes de importar.'
    )
    guia['A3'].alignment = Alignment(wrap_text=True, vertical='top')
    guia.merge_cells('A3:D3')
    guia.row_dimensions[3].height = 45

    if nota:
        guia['A5'] = nota
        guia['A5'].alignment = Alignment(wrap_text=True, vertical='top')
        guia.merge_cells('A5:D5')
        guia.row_dimensions[5].height = 30

    encabezados = ['Columna', 'Tipo', '¿Obligatoria?', 'Detalle']
    for indice, texto in enumerate(encabezados, start=1):
        celda = guia.cell(row=7, column=indice, value=texto)
        celda.fill = RELLENO_ENCABEZADO
        celda.font = LETRA_ENCABEZADO

    for numero, columna in enumerate(columnas, start=8):
        guia.cell(row=numero, column=1, value=columna.nombre)
        guia.cell(row=numero, column=2, value=columna.descripcion_tipo)
        guia.cell(row=numero, column=3, value='Sí' if columna.obligatoria else 'No')
        detalle = columna.ayuda
        if columna.opciones:
            opciones = 'Opciones: ' + ', '.join(columna.opciones)
            detalle = (detalle + ' ' + opciones) if detalle else opciones
        celda = guia.cell(row=numero, column=4, value=detalle)
        celda.alignment = Alignment(wrap_text=True, vertical='top')

    return _a_bytes(libro)


def leer_archivo(archivo, columnas: list[Columna]) -> tuple[list[dict], list[dict]]:
    """
    Lee el .xlsx subido y devuelve `(filas, errores)`.

    Las columnas se localizan por el texto del encabezado, no por su posición:
    reordenar las columnas en Excel no debe romper la importación.
    """
    try:
        libro = load_workbook(archivo, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl lanza de todo con archivos corruptos
        raise ErrorDeFila(
            'No se pudo abrir el archivo. Debe ser un Excel .xlsx generado desde la plantilla.'
        ) from exc

    hoja = libro['Datos'] if 'Datos' in libro.sheetnames else libro.active

    encabezados = {}
    primera = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), ())
    for indice, celda in enumerate(primera):
        if celda:
            encabezados[str(celda).strip().lower()] = indice

    faltantes = [c.nombre for c in columnas if c.nombre.lower() not in encabezados]
    if faltantes:
        raise ErrorDeFila(
            'Al archivo le faltan columnas: ' + ', '.join(faltantes) + '. '
            'Descarga la plantilla y vuelve a intentarlo.'
        )

    filas: list[dict] = []
    errores: list[dict] = []

    for numero, valores in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        if valores is None or all(v is None or str(v).strip() == '' for v in valores):
            continue  # fila en blanco

        datos: dict = {}
        problemas: list[str] = []
        for columna in columnas:
            indice = encabezados[columna.nombre.lower()]
            valor = valores[indice] if indice < len(valores) else None
            try:
                leido = _leer_celda(valor, columna)
            except ErrorDeFila as exc:
                problemas.append(str(exc))
                continue
            if leido is not None:
                datos[columna.nombre] = leido

        if problemas:
            errores.append({'fila': numero, 'errores': problemas})
        else:
            filas.append({'_fila': numero, **datos})

    return filas, errores


@dataclass
class ColumnaExport:
    """Una columna del archivo exportado.

    `campo` es una ruta de `values()` (`id_producto__marca`), así que la
    exportación se arma con una sola consulta y no instanciando modelos.
    """

    nombre: str
    campo: str
    formato: str = 'texto'  # texto | entero | dinero | fecha


#: Formatos de celda de Excel por tipo de columna.
_FORMATOS = {
    'entero': '#,##0',
    'dinero': '"$" #,##0',
    'fecha': 'yyyy-mm-dd',
}


def _escribir_hoja(hoja, columnas: list[ColumnaExport], filas: list[dict]) -> None:
    """Encabezado, datos, formatos y autofiltro de una hoja."""
    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=indice, value=columna.nombre)
        celda.fill = RELLENO_ENCABEZADO
        celda.font = LETRA_ENCABEZADO
        celda.alignment = Alignment(horizontal='center', vertical='center')
        hoja.column_dimensions[get_column_letter(indice)].width = max(len(columna.nombre) + 4, 16)

    for numero, fila in enumerate(filas, start=2):
        for indice, columna in enumerate(columnas, start=1):
            valor = fila.get(columna.campo)
            celda = hoja.cell(row=numero, column=indice, value=valor)
            formato = _FORMATOS.get(columna.formato)
            if formato and valor is not None:
                celda.number_format = formato

    hoja.freeze_panes = 'A2'
    hoja.auto_filter.ref = f'A1:{get_column_letter(len(columnas))}{max(len(filas) + 1, 2)}'


def construir_export(titulo: str, columnas: list[ColumnaExport], filas: list[dict]) -> bytes:
    """
    Vuelca un listado ya filtrado a un .xlsx de una sola hoja.

    A diferencia de la plantilla, esto no lleva instrucciones ni ejemplos: es
    el dato tal cual, con autofiltro y encabezado congelado para poder
    trabajarlo en Excel de una vez.
    """
    return construir_export_multihoja([(titulo, columnas, filas)])


def construir_export_multihoja(
    hojas: list[tuple[str, list[ColumnaExport], list[dict]]],
) -> bytes:
    """
    Un .xlsx con una hoja por bloque: `(título, columnas, filas)`.

    Lo usa el tablero, donde un mismo mes se mira de tres formas —día por día,
    por regional y por punto de venta— y tenerlas en un solo archivo evita
    descargar tres.
    """
    libro = Workbook()
    # `Workbook()` ya trae una hoja: la primera se reutiliza y el resto se crea,
    # así no queda una «Sheet» vacía al final del archivo.
    for indice, (titulo, columnas, filas) in enumerate(hojas):
        nombre = (titulo[:31] or 'Datos')
        hoja = libro.active if indice == 0 else libro.create_sheet()
        hoja.title = nombre
        _escribir_hoja(hoja, columnas, filas)
    return _a_bytes(libro)
