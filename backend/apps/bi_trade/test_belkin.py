"""Plan Recomiéndame Belkin: el formulario, sus catálogos, el bono y el tablero."""
from datetime import date, datetime
from io import BytesIO, StringIO

import openpyxl
import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db.models import Count
from django.utils import timezone
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table
from rest_framework.test import APIClient

from apps.bi_trade.belkin import calcular_bono, siguiente_bono
from apps.bi_trade.models import (
    AsesorApple,
    CategoriaBelkin,
    EnlacePublico,
    Producto,
    ProductoBelkin,
    PuntoVenta,
    PuntoVentaBelkin,
    RegionalBelkin,
    RegistroBelkin,
    Venta,
)
from apps.bi_trade.tests import _importar, app_bi_trade, cliente_de, crear_usuario  # noqa: F401

pytestmark = pytest.mark.django_db

RUTA = '/api/bi-trade/belkin'


@pytest.fixture
def catalogo_belkin():
    norte = RegionalBelkin.objects.create(nombre='Zona Norte')
    sur = RegionalBelkin.objects.create(nombre='Zona Sur')
    andino = PuntoVentaBelkin.objects.create(
        id_punto_venta='C159', nombre_pdv='Cav Andino', id_regional=norte
    )
    plaza_claro = PuntoVentaBelkin.objects.create(
        id_punto_venta='C108', nombre_pdv='Cav Bogota Plaza Claro', id_regional=sur
    )
    case = CategoriaBelkin.objects.create(nombre='Case Apple')
    lamina = CategoriaBelkin.objects.create(nombre='Lámina')
    producto = ProductoBelkin.objects.create(
        id_producto='7020178',
        nombre_producto='Spigen Iphone 12 / 12 Pro Case Crystal Flex',
        id_categoria=case,
    )
    asesor = AsesorApple.objects.create(nombre='Duvan Riaño', id_punto_venta=plaza_claro)
    return {
        'norte': norte,
        'sur': sur,
        'andino': andino,
        'plaza_claro': plaza_claro,
        'case': case,
        'lamina': lamina,
        'producto': producto,
        'asesor': asesor,
    }


def _recomendacion(punto, producto, asesor=None):
    return {
        'idPuntoVenta': punto.pk,
        'idAsesor': asesor.pk if asesor else None,
        'idProducto': producto.pk,
        'fechaRecomendacion': str(timezone.localdate()),
        'observacion': '  Cliente volvió por otro  ',
    }


def test_el_formulario_guarda_el_codigo_y_lee_regional_y_categoria(
    app_bi_trade, catalogo_belkin  # noqa: F811
):
    c = catalogo_belkin
    promotor = crear_usuario('promotor@supli.tech', app_bi_trade)

    respuesta = cliente_de(promotor).post(
        f'{RUTA}/registros',
        _recomendacion(c['plaza_claro'], c['producto'], c['asesor']),
        format='json',
    )

    assert respuesta.status_code == 201, respuesta.data
    cuerpo = respuesta.json()
    assert cuerpo['puntoVentaEtiqueta'] == 'Cav Bogota Plaza Claro \\ C108'
    assert cuerpo['productoEtiqueta'] == '7020178 \\ Spigen Iphone 12 / 12 Pro Case Crystal Flex'
    assert cuerpo['regional'] == 'Zona Sur'
    assert cuerpo['categoria'] == 'Case Apple'
    assert cuerpo['asesor'] == 'Duvan Riaño'
    assert cuerpo['observacion'] == 'Cliente volvió por otro'
    assert RegistroBelkin.objects.get().registrado_por == promotor


def test_el_asesor_es_opcional(app_bi_trade, catalogo_belkin):  # noqa: F811
    c = catalogo_belkin
    respuesta = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade)).post(
        f'{RUTA}/registros', _recomendacion(c['andino'], c['producto']), format='json'
    )

    assert respuesta.status_code == 201, respuesta.data
    assert respuesta.json()['asesor'] == ''


def test_el_asesor_tiene_que_ser_del_punto_elegido(app_bi_trade, catalogo_belkin):  # noqa: F811
    c = catalogo_belkin
    respuesta = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade)).post(
        f'{RUTA}/registros',
        _recomendacion(c['andino'], c['producto'], c['asesor']),
        format='json',
    )

    assert respuesta.status_code == 400
    assert 'idAsesor' in respuesta.json()['errors']
    assert RegistroBelkin.objects.count() == 0


def test_mover_el_punto_o_el_producto_mueve_lo_ya_cargado(
    app_bi_trade, catalogo_belkin  # noqa: F811
):
    """Cav Andino pasa a Zona Sur y el case a Lámina: lo viejo se va con ellos."""
    c = catalogo_belkin
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)
    cliente.post(f'{RUTA}/registros', _recomendacion(c['andino'], c['producto']), format='json')

    assert cliente.patch(
        f'{RUTA}/puntos-venta/C159', {'idRegional': c['sur'].pk}, format='json'
    ).status_code == 200
    assert cliente.patch(
        f'{RUTA}/productos/7020178', {'idCategoria': c['lamina'].pk}, format='json'
    ).status_code == 200

    registro = cliente.get(f'{RUTA}/registros').json()['items'][0]
    assert registro['regional'] == 'Zona Sur'
    assert registro['categoria'] == 'Lámina'

    # Y el formulario también los muestra en su nuevo grupo.
    opciones = cliente.get(f'{RUTA}/opciones').json()
    andino = next(p for p in opciones['puntosVenta'] if p['value'] == 'C159')
    assert andino['idRegional'] == c['sur'].pk
    case = next(p for p in opciones['productos'] if p['value'] == '7020178')
    assert case['idCategoria'] == c['lamina'].pk


def test_las_opciones_traen_cada_asesor_con_su_punto(app_bi_trade, catalogo_belkin):  # noqa: F811
    c = catalogo_belkin
    AsesorApple.objects.create(nombre='Inactivo', id_punto_venta=c['plaza_claro'], activo=False)

    opciones = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade)).get(
        f'{RUTA}/opciones'
    ).json()

    assert opciones['asesores'] == [
        {'value': c['asesor'].pk, 'label': 'Duvan Riaño', 'idPuntoVenta': 'C108'}
    ]
    assert {r['label'] for r in opciones['regionales']} == {'Zona Norte', 'Zona Sur'}
    assert {cat['label'] for cat in opciones['categorias']} == {'Case Apple', 'Lámina'}


def test_un_punto_puede_tener_varios_asesores_sin_repetir_nombre(
    app_bi_trade, catalogo_belkin  # noqa: F811
):
    cliente = cliente_de(
        crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    )

    otro = cliente.post(
        f'{RUTA}/asesores', {'nombre': 'Kerlys Babilonia', 'idPuntoVenta': 'C108'}, format='json'
    )
    repetido = cliente.post(
        f'{RUTA}/asesores', {'nombre': 'DUVAN  Riaño', 'idPuntoVenta': 'C108'}, format='json'
    )

    assert otro.status_code == 201
    assert repetido.status_code == 400
    assert cliente.get(f'{RUTA}/puntos-venta/C108').json()['asesoresCount'] == 2


def test_administrar_los_catalogos_pide_permiso(app_bi_trade, catalogo_belkin):  # noqa: F811
    promotor = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade))
    for ruta, cuerpo in (
        ('regionales', {'nombre': 'Zona Costa'}),
        ('categorias', {'nombre': 'Audífonos'}),
        ('asesores', {'nombre': 'Ana', 'idPuntoVenta': 'C108'}),
    ):
        assert promotor.post(f'{RUTA}/{ruta}', cuerpo, format='json').status_code == 403


def test_no_se_borra_lo_que_ya_se_usa(app_bi_trade, catalogo_belkin):  # noqa: F811
    c = catalogo_belkin
    cliente = cliente_de(
        crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    )
    cliente.post(
        f'{RUTA}/registros',
        _recomendacion(c['plaza_claro'], c['producto'], c['asesor']),
        format='json',
    )

    for ruta in (
        f'{RUTA}/regionales/{c["sur"].pk}',
        f'{RUTA}/puntos-venta/C108',
        f'{RUTA}/asesores/{c["asesor"].pk}',
        f'{RUTA}/categorias/{c["case"].pk}',
        f'{RUTA}/productos/7020178',
    ):
        respuesta = cliente.delete(ruta)
        assert respuesta.status_code == 400, ruta
        assert 'Desactívalo' in respuesta.json()['message']

    # Lo que no se usa sí se borra.
    assert cliente.delete(f'{RUTA}/categorias/{c["lamina"].pk}').status_code == 204


def test_la_carga_inicial_es_idempotente():
    call_command('seed_plan_belkin', stdout=StringIO())
    call_command('seed_plan_belkin', stdout=StringIO())

    assert RegionalBelkin.objects.count() == 2
    assert PuntoVentaBelkin.objects.count() == 27
    assert PuntoVentaBelkin.objects.get(pk='C159').id_regional.nombre == 'Zona Norte'
    assert ProductoBelkin.objects.get(pk='7020178').id_categoria.nombre == 'Case Apple'
    assert CategoriaBelkin.objects.count() == 4



# ── El formulario abierto por enlace ───────────────────────────────────────
# Igual que el de Partners: sin cuenta ni contraseña, y solo para enviar.

RUTA_FORMULARIO = '/api/publico/formulario'


def _enlace(app, canal):
    editor = crear_usuario(f'editor-{canal}@supli.tech', app, ['bi-trade:data:manage'])
    datos = (
        cliente_de(editor)
        .post('/api/bi-trade/enlaces', {'nombre': 'Asesores Norte', 'canal': canal}, format='json')
        .json()
    )
    assert datos['clave'] == ''
    return datos['token']


def test_el_formulario_publico_se_diligencia_sin_cuenta_ni_clave(
    app_bi_trade, catalogo_belkin  # noqa: F811
):
    c = catalogo_belkin
    token = _enlace(app_bi_trade, 'belkin')
    anonimo = APIClient()

    opciones = anonimo.get(f'{RUTA_FORMULARIO}/belkin/{token}/opciones')
    assert opciones.status_code == 200
    assert len(opciones.json()['puntosVenta']) == 2

    respuesta = anonimo.post(
        f'{RUTA_FORMULARIO}/belkin/{token}/registros',
        _recomendacion(c['plaza_claro'], c['producto'], c['asesor']),
        format='json',
    )

    assert respuesta.status_code == 201, respuesta.data
    assert respuesta.json()['origen'] == 'Enlace · Asesores Norte'
    registro = RegistroBelkin.objects.get()
    assert registro.registrado_por is None
    assert EnlacePublico.objects.get(pk=registro.enlace_id).accesos == 1


def test_por_el_formulario_belkin_no_se_puede_hacer_nada_mas(app_bi_trade):  # noqa: F811
    token = _enlace(app_bi_trade, 'belkin')
    anonimo = APIClient()

    # Ni el tablero, ni lo ya cargado, ni el formulario del otro plan.
    assert anonimo.get(f'/api/publico/bi-trade/{token}').status_code == 404
    assert anonimo.get(f'{RUTA_FORMULARIO}/belkin/{token}/registros').status_code == 405
    assert anonimo.get(f'{RUTA_FORMULARIO}/{token}/opciones').status_code == 404
    # Y no tiene contraseña que regenerar.
    editor = cliente_de(EnlacePublico.objects.get(token=token).creado_por)
    enlace = EnlacePublico.objects.get(token=token)
    assert editor.post(f'/api/bi-trade/enlaces/{enlace.pk}/regenerar-clave').status_code == 400


def test_un_enlace_de_partners_no_abre_el_formulario_belkin(app_bi_trade):  # noqa: F811
    token = _enlace(app_bi_trade, 'partners')
    assert APIClient().get(f'{RUTA_FORMULARIO}/belkin/{token}/opciones').status_code == 404


def test_un_enlace_revocado_ya_no_abre(app_bi_trade, catalogo_belkin):  # noqa: F811
    token = _enlace(app_bi_trade, 'belkin')
    EnlacePublico.objects.filter(token=token).update(activo=False)
    assert APIClient().get(f'{RUTA_FORMULARIO}/belkin/{token}/opciones').status_code == 404


# ── El bono ────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ('recomendaciones', 'categoria', 'bono'),
    [
        # Lo que el Power BI liquidó en septiembre de 2026.
        (62, 2, 180_000),  # dos paquetes de 24 y uno de 14
        (58, 2, 140_000),  # dos de 24; los 10 que sobran no completan nada
        (73, 1, 140_000),
        (45, 2, 120_000),  # uno de 24 y uno de 16
        (14, 2, 40_000),
        (13, 1, 0),
        (3, 3, 0),
        # El tope y el punto sin categoría.
        (100, 3, 300_000),
        (50, None, 0),
    ],
)
def test_el_bono_reparte_en_paquetes_de_mayor_a_menor(recomendaciones, categoria, bono):
    assert calcular_bono(recomendaciones, categoria) == bono


def test_el_siguiente_bono_dice_cuanto_falta():
    assert siguiente_bono(58, 2) == {'faltan': 2, 'bono': 170_000}
    assert siguiente_bono(0, 3) == {'faltan': 10, 'bono': 30_000}
    assert siguiente_bono(100, 3) is None
    assert siguiente_bono(5, None) is None


# ── Puntos fuera de Coltrade ───────────────────────────────────────────────


@pytest.fixture
def fuera_de_coltrade(catalogo_belkin):
    """Un punto sin regional, con su único asesor."""
    punto = PuntoVentaBelkin.objects.create(
        id_punto_venta='C304', nombre_pdv='Cav Armenia Portal Quindio', categoria=3
    )
    asesor = AsesorApple.objects.create(nombre='Olga Diaz', id_punto_venta=punto)
    return punto, asesor


def test_el_formulario_no_carga_puntos_fuera_de_coltrade(
    app_bi_trade, catalogo_belkin, fuera_de_coltrade  # noqa: F811
):
    punto, _ = fuera_de_coltrade
    cliente = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade))

    respuesta = cliente.post(
        f'{RUTA}/registros', _recomendacion(punto, catalogo_belkin['producto']), format='json'
    )

    assert respuesta.status_code == 400
    assert 'informe' in respuesta.json()['errors']['idPuntoVenta'][0]
    opciones = cliente.get(f'{RUTA}/opciones').json()
    assert 'C304' not in {p['value'] for p in opciones['puntosVenta']}
    # En la lista sí está, con su categoría y sin regional.
    fila = cliente.get(f'{RUTA}/puntos-venta/C304').json()
    assert (fila['idRegional'], fila['regional'], fila['categoria']) == (None, '', 3)


def test_el_informe_de_claro_carga_las_ventas_de_los_puntos_fuera_de_coltrade(
    app_bi_trade, catalogo_belkin, fuera_de_coltrade  # noqa: F811
):
    c = catalogo_belkin
    _, olga = fuera_de_coltrade
    editor = cliente_de(
        crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    )
    formulario = RegistroBelkin.objects.create(
        id_punto_venta=c['plaza_claro'],
        id_producto=c['producto'],
        fecha_recomendacion=date(2026, 9, 4),
    )

    datos = _importar(
        editor,
        [
            ('7020178', 'C304', '601', date(2026, 9, 4)),
            ('7020178', 'C304', '601', date(2026, 9, 4)),
            ('7020178', 'C304', '601', date(2026, 9, 5)),
            # Punto de Coltrade: lo suyo llega por el formulario.
            ('7020178', 'C108', '601', date(2026, 9, 5)),
            # Material que no está en la lista del plan.
            ('9999999', 'C304', '601', date(2026, 9, 5)),
            # Otro mes.
            ('7020178', 'C304', '601', date(2026, 8, 30)),
        ],
        anio=2026,
        mes=9,
    ).json()

    assert datos['belkin']['creados'] == 3
    del_informe = RegistroBelkin.objects.filter(fuente='informe')
    assert del_informe.count() == 3
    # El punto tiene un solo asesor: los registros quedan a su nombre.
    assert set(del_informe.values_list('id_asesor', flat=True)) == {olga.pk}

    # Completar respeta los días que ya tienen registros del informe…
    nuevas = [
        ('7020178', 'C304', '601', date(2026, 9, 4)),
        ('7020178', 'C304', '601', date(2026, 9, 6)),
    ]
    datos = _importar(editor, nuevas, anio=2026, mes=9).json()
    assert (datos['belkin']['creados'], datos['belkin']['omitidosPorDia']) == (1, 1)
    assert del_informe.count() == 4

    # …y sobrescribir deja exactamente lo del archivo, sin tocar el formulario.
    datos = _importar(editor, nuevas, anio=2026, mes=9, modo='sobrescribir').json()
    assert (datos['belkin']['creados'], datos['belkin']['eliminados']) == (2, 4)
    assert del_informe.count() == 2
    assert RegistroBelkin.objects.filter(pk=formulario.pk).exists()


# ── El tablero ─────────────────────────────────────────────────────────────


def _registros(n, punto, producto, asesor=None, fuente='formulario', dia=3):
    RegistroBelkin.objects.bulk_create(
        RegistroBelkin(
            id_punto_venta=punto,
            id_asesor=asesor,
            id_producto=producto,
            fecha_recomendacion=date(2026, 9, dia),
            fuente=fuente,
        )
        for _ in range(n)
    )


def test_el_tablero_liquida_el_bono_de_cada_promotor_con_la_categoria_del_punto(
    app_bi_trade, catalogo_belkin, fuera_de_coltrade  # noqa: F811
):
    c = catalogo_belkin
    fuera, olga = fuera_de_coltrade
    PuntoVentaBelkin.objects.filter(pk='C108').update(categoria=1)
    soacha = PuntoVentaBelkin.objects.create(
        id_punto_venta='C104', nombre_pdv='Cav Soacha Mercurio', id_regional=c['sur'], categoria=2
    )
    cubillos = AsesorApple.objects.create(nombre='John Cubillos', id_punto_venta=soacha)
    _registros(58, soacha, c['producto'], cubillos)
    _registros(13, c['plaza_claro'], c['producto'], c['asesor'], dia=4)
    _registros(2, c['plaza_claro'], c['producto'], dia=4)
    _registros(20, fuera, c['producto'], olga, fuente='informe', dia=5)
    # Lo que Claro vendió del mismo producto en Soacha.
    producto_claro = Producto.objects.create(
        id_producto='7020178',
        nombre_producto='Spigen Iphone 12',
        marca='Spigen',
        precio_venta_claro=1,
        precio_venta_coltrade=1,
    )
    Venta.objects.create(
        id_producto=producto_claro,
        id_punto_venta=PuntoVenta.objects.create(id_punto_venta='C104', nombre_pdv='Soacha'),
        fecha_venta=date(2026, 9, 3),
        cantidad_vendida=5,
    )
    cliente = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade))

    datos = cliente.get(f'{RUTA}/dashboard', {'anio': 2026, 'mes': 9}).json()

    totales = datos['totales']
    assert (totales['recomendaciones'], totales['formulario'], totales['informe']) == (93, 73, 20)
    # 140.000 de Soacha y 70.000 del punto de fuera; Plaza Claro no llega.
    assert (totales['bono'], totales['promotoresConBono'], totales['promotores']) == (
        210_000,
        2,
        4,
    )
    assert totales['ventas'] == 25  # 5 de Claro en Soacha y 20 del informe
    primero = datos['promotores'][0]
    assert (primero['asesor'], primero['bono'], primero['siguiente']) == (
        'John Cubillos',
        140_000,
        {'faltan': 2, 'bono': 170_000},
    )
    sin_asesor = next(p for p in datos['promotores'] if p['idAsesor'] is None)
    assert (sin_asesor['codigo'], sin_asesor['recomendaciones']) == ('C108', 2)
    armenia = next(p for p in datos['porPunto'] if p['codigo'] == 'C304')
    assert (armenia['regional'], armenia['fueraDeColtrade'], armenia['ventas']) == (
        'Fuera de Coltrade',
        True,
        20,
    )
    dia3 = next(d for d in datos['porDia'] if d['fecha'] == '2026-09-03')
    assert (dia3['recomendaciones'], dia3['ventas']) == (58, 5)
    assert len(datos['porDia']) == 30

    # El filtro de regional también separa los puntos de fuera.
    solo_fuera = cliente.get(
        f'{RUTA}/dashboard', {'anio': 2026, 'mes': 9, 'regional': 'fuera'}
    ).json()
    assert [p['codigo'] for p in solo_fuera['porPunto']] == ['C304']


def test_el_excel_por_punto_trae_todos_los_puntos_aunque_esten_en_cero(
    app_bi_trade, catalogo_belkin, fuera_de_coltrade  # noqa: F811
):
    c = catalogo_belkin
    PuntoVentaBelkin.objects.filter(pk='C108').update(categoria=1)
    PuntoVentaBelkin.objects.create(
        id_punto_venta='C999', nombre_pdv='Cav Cerrado', id_regional=c['sur'], activo=False
    )
    _registros(13, c['plaza_claro'], c['producto'], c['asesor'])
    cliente = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade))

    respuesta = cliente.get(f'{RUTA}/dashboard/puntos/exportar', {'anio': 2026, 'mes': 9})

    assert respuesta.status_code == 200
    assert 'belkin-puntos-2026-09.xlsx' in respuesta['Content-Disposition']
    hoja = openpyxl.load_workbook(BytesIO(respuesta.content)).active
    filas = {fila[0]: fila for fila in hoja.iter_rows(min_row=2, values_only=True)}
    # Plaza Claro con lo suyo; Andino y el de fuera en cero; el inactivo sin datos no.
    assert set(filas) == {'C108', 'C159', 'C304'}
    assert filas['C108'][1:5] == ('Cav Bogota Plaza Claro', 'Zona Sur', 1, 'Duvan Riaño')
    assert filas['C108'][7] == 13
    assert filas['C159'][7] == 0
    assert (filas['C304'][2], filas['C304'][4]) == ('Fuera de Coltrade', 'Olga Diaz')


def _abrir_enlace(app, canal):
    """Crea un enlace de tablero y entra con su contraseña: devuelve token y acceso."""
    editor = crear_usuario(f'editor-{canal}@supli.tech', app, ['bi-trade:data:manage'])
    creado = (
        cliente_de(editor)
        .post('/api/bi-trade/enlaces', {'nombre': 'Gerencia', 'canal': canal}, format='json')
        .json()
    )
    assert creado['clave'], 'el tablero va con contraseña'
    acceso = APIClient().post(
        f'/api/publico/bi-trade/{creado["token"]}/acceso', {'clave': creado['clave']}
    )
    assert acceso.status_code == 200
    return creado['token'], acceso.json()['acceso']


def test_el_tablero_belkin_se_comparte_de_solo_lectura(
    app_bi_trade, catalogo_belkin  # noqa: F811
):
    c = catalogo_belkin
    _registros(13, c['plaza_claro'], c['producto'], c['asesor'])
    token, acceso = _abrir_enlace(app_bi_trade, 'belkin_bi')
    anonimo = APIClient(HTTP_X_ACCESO_PUBLICO=acceso)
    publico = f'/api/publico/bi-trade/{token}'

    datos = anonimo.get(f'{publico}/belkin', {'anio': 2026, 'mes': 9}).json()
    assert datos['totales']['recomendaciones'] == 13
    plaza_claro = {
        'value': 'C108',
        'label': 'Cav Bogota Plaza Claro \\ C108',
        'regional': str(c['sur'].pk),
    }
    assert plaza_claro in datos['opciones']['puntos']

    # Nada más: ni los tableros de los canales, ni las listas, ni el Excel.
    for tramo in ('avance-mensual', 'opciones', 'puntos-venta', 'productos', 'campanas'):
        assert anonimo.get(f'{publico}/{tramo}').status_code == 404, tramo
    assert anonimo.get(f'{RUTA}/puntos-venta').status_code == 401
    assert anonimo.get(f'{RUTA}/dashboard/puntos/exportar').status_code == 401
    # Y sin el acceso, ni el tablero.
    assert APIClient().get(f'{publico}/belkin').status_code == 403


def test_un_enlace_de_claro_no_abre_el_tablero_belkin(app_bi_trade):  # noqa: F811
    token, acceso = _abrir_enlace(app_bi_trade, 'claro')
    anonimo = APIClient(HTTP_X_ACCESO_PUBLICO=acceso)
    assert anonimo.get(f'/api/publico/bi-trade/{token}/belkin').status_code == 404


def test_sin_periodo_el_tablero_abre_en_el_ultimo_mes_con_registros(
    app_bi_trade, catalogo_belkin  # noqa: F811
):
    c = catalogo_belkin
    _registros(1, c['plaza_claro'], c['producto'])
    datos = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade)).get(
        f'{RUTA}/dashboard'
    ).json()

    assert (datos['filtros']['anio'], datos['filtros']['mes']) == (2026, 9)
    assert datos['filtros']['periodo'] == 'Septiembre 2026'
    assert {'anio': 2026, 'mes': 9, 'label': 'Septiembre 2026'} in datos['periodos']


# ── Las cargas desde Excel ─────────────────────────────────────────────────


def _libro_con_tablas(tablas: dict[str, list[list]], ruta) -> str:
    """Un .xlsx con tablas de Excel con nombre, como los archivos de Trade."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    columna = 1
    for nombre, filas in tablas.items():
        for i, fila in enumerate(filas, start=1):
            for j, valor in enumerate(fila):
                hoja.cell(row=i, column=columna + j, value=valor)
        inicio = get_column_letter(columna)
        fin = get_column_letter(columna + len(filas[0]) - 1)
        hoja.add_table(Table(displayName=nombre, ref=f'{inicio}1:{fin}{len(filas)}'))
        columna += len(filas[0]) + 1
    archivo = ruta / 'libro.xlsx'
    libro.save(archivo)
    return str(archivo)


def test_el_catalogo_se_carga_desde_informacion_claro(catalogo_belkin, tmp_path):
    archivo = _libro_con_tablas(
        {
            'Puntos_Venta': [
                ['Centro_Costos', 'Bloque', 'Regional', 'Punto_Venta', 'Asesor_Apple', 'Categoria'],
                ['C159', 'Bloque 1', 'Zona Norte', 'Cav Andino', 'Óscar Plata', 3],
                ['C304', 'Bloque 1', 'Cav Fuera de Coltrade', 'Cav Armenia C304', 'Olga Diaz', 3],
                ['C322', 'Bloque 2', 'Cav Fuera de Coltrade', 'Cav Quibdo', 'Agregar', 2],
            ],
            'Productos': [
                ['Material', 'Tipo', 'Producto'],
                [7020178, 'Case Apple', 'Otro nombre que no pisa el de la lista'],
                [7023246, 'Lamina Belkin', 'Lmna Tmpredglss S_F Iphone 17Promax Belk'],
                [7022267, 'Powerbank', 'Power Bank  Magnetico 7.5W 5000Mah Belk'],
            ],
        },
        tmp_path,
    )

    call_command('cargar_catalogo_belkin', archivo=archivo, stdout=StringIO())
    call_command('cargar_catalogo_belkin', archivo=archivo, stdout=StringIO())

    andino = PuntoVentaBelkin.objects.get(pk='C159')
    assert (andino.nombre_pdv, andino.id_regional.nombre, andino.categoria) == (
        'Cav Andino',
        'Zona Norte',
        3,
    )
    armenia = PuntoVentaBelkin.objects.get(pk='C304')
    assert (armenia.nombre_pdv, armenia.id_regional, armenia.categoria) == (
        'Cav Armenia',
        None,
        3,
    )
    assert list(armenia.asesores.values_list('nombre', flat=True)) == ['Olga Diaz']
    assert not PuntoVentaBelkin.objects.get(pk='C322').asesores.exists()
    assert AsesorApple.objects.filter(nombre='Óscar Plata').count() == 1
    # El producto que ya estaba no cambia; los nuevos entran inactivos.
    assert ProductoBelkin.objects.get(pk='7020178').nombre_producto.startswith('Spigen')
    lamina = ProductoBelkin.objects.get(pk='7023246')
    assert (lamina.id_categoria.nombre, lamina.activo) == ('Lámina', False)
    assert not CategoriaBelkin.objects.get(nombre='Power Bank').activa


def test_el_formulario_anterior_se_importa_cruzando_al_asesor_por_apellido(
    catalogo_belkin, tmp_path
):
    c = catalogo_belkin
    soacha = PuntoVentaBelkin.objects.create(
        id_punto_venta='C104', nombre_pdv='Cav Soacha Mercurio', id_regional=c['sur']
    )
    AsesorApple.objects.create(nombre='John Fredy Cubillos Galindo', id_punto_venta=soacha)
    encabezado = [
        'ID',
        'Zona Norte:',
        'Zona Sur:',
        'Fecha de la Recomendación:',
        'Nombre y Apellido Asesor Apple:',
        'Case Apple',
        'Lámina',
        'Cable',
        'Cargador',
        '¿Alguna Observación?, es opcional.',
    ]
    producto = '7020178 \\ Spigen Iphone 12 / 12 Pro Case Crystal Flex'

    def fila(n, punto, asesor, dia=1):
        return [n, None, punto, datetime(2026, 9, dia), asesor, producto, None, None, None, None]

    archivo = _libro_con_tablas(
        {
            'Ventas_Formulario': [
                encabezado,
                fila(1, 'Cav Soacha Mercurio \\ C104', 'Jhon Cubillos'),
                fila(2, 'Cav Soacha Mercurio \\ C104', 'JOHN  CUBILLOS'),
                fila(3, 'Cav Bogota Plaza Claro \\ C108', 'duvan riaño'),
                # Cubillos registró en otro punto: queda como asesor de ese punto.
                fila(4, 'Cav Bogota Plaza Claro \\ C108', 'Jhon Cubillos'),
                fila(5, 'Cav Bogota Plaza Claro \\ C108', 'Julián'),
                fila(6, 'Cav Inexistente \\ C999', 'Nadie'),
                fila(7, 'Cav Soacha Mercurio \\ C104', 'Jhon Cubillos', dia=30),
            ]
        },
        tmp_path,
    )
    opciones = {'archivo': archivo, 'desde': '2026-09-01', 'hasta': '2026-09-15'}

    call_command('importar_formulario_belkin', **opciones, stdout=StringIO())

    assert RegistroBelkin.objects.count() == 5
    asesores = dict(
        RegistroBelkin.objects.values('id_asesor__nombre')
        .annotate(n=Count('pk'))
        .values_list('id_asesor__nombre', 'n')
    )
    assert asesores == {
        'John Fredy Cubillos Galindo': 3,
        'Duvan Riaño': 1,
        'Julián': 1,
    }
    assert AsesorApple.objects.filter(
        nombre='John Fredy Cubillos Galindo', id_punto_venta='C108'
    ).exists()

    # Correrlo otra vez pide reemplazar, y reemplazar no duplica.
    with pytest.raises(CommandError):
        call_command('importar_formulario_belkin', **opciones, stdout=StringIO())
    call_command('importar_formulario_belkin', **opciones, reemplazar=True, stdout=StringIO())
    assert RegistroBelkin.objects.count() == 5
