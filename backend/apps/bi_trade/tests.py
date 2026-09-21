"""Pruebas de BI Trade Marketing: CRUD, permisos y agregados del tablero."""
from datetime import date, timedelta
from io import BytesIO, StringIO

import openpyxl
import pytest
from django.contrib.auth.hashers import check_password
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import Application, Permission, Role, User
from apps.bi_trade import publico
from apps.bi_trade.calendario import dias_habiles, es_habil, festivos, pascua
from apps.bi_trade.models import (
    Acelerador,
    Campana,
    EnlacePublico,
    EscalaTicket,
    Inventario,
    InventarioFalabella,
    InventarioHc,
    InventarioTmk,
    MarcaPartner,
    Materiales,
    MetaComercial,
    MetaComercialFalabella,
    MetaComercialHc,
    MetaComercialTmk,
    MetaPartner,
    Producto,
    ProductoFalabella,
    ProductoHc,
    ProductoPartner,
    ProductoTmk,
    PuntoVenta,
    PuntoVentaFalabella,
    PuntoVentaHc,
    PuntoVentaPartner,
    PuntoVentaTmk,
    Regional,
    RegionalFalabella,
    RegionalHc,
    RegionalPartner,
    RegionalTmk,
    RegistroPartner,
    Venta,
    VentaFalabella,
    VentaHc,
    VentaTmk,
)
from apps.bi_trade.tickets import calcular_tickets, participa

pytestmark = pytest.mark.django_db


@pytest.fixture
def app_bi_trade():
    app = Application.objects.create(
        code='bi-trade', name='BI Trade Marketing', base_path='/inicio/bi-trade', order=200
    )
    Permission.objects.create(
        code='bi-trade:data:manage', name='Editar datos', application=app
    )
    return app


def crear_usuario(email, app=None, permisos=()):
    usuario = User.objects.create_user(
        email=email,
        username=email.split('@')[0],
        first_name=email.split('@')[0].title(),
        password='clave-de-prueba-123',
    )
    if app:
        usuario.applications.add(app)
    if permisos:
        rol = Role.objects.create(code=f'rol-{usuario.pk}', name='rol')
        rol.permissions.set(Permission.objects.filter(code__in=permisos))
        usuario.roles.add(rol)
    return usuario


def cliente_de(usuario) -> APIClient:
    cliente = APIClient()
    cliente.force_authenticate(user=usuario)
    return cliente


@pytest.fixture
def catalogo():
    pdv_norte = PuntoVenta.objects.create(
        id_punto_venta='PDV-001',
        nombre_pdv='Claro Unicentro',
        regional=Regional.ZONA_NORTE,
        materiales=Materiales.TODOS,
    )
    pdv_sur = PuntoVenta.objects.create(
        id_punto_venta='PDV-002',
        nombre_pdv='Claro Jardín Plaza',
        regional=Regional.ZONA_SUR,
        materiales=Materiales.INCOMPLETOS,
    )
    caro = Producto.objects.create(
        id_producto='SKU-1',
        nombre_producto='iPhone 15',
        marca='Apple',
        precio_venta_claro=5_000_000,
        precio_venta_coltrade=4_000_000,
        puntaje=99,
    )
    barato = Producto.objects.create(
        id_producto='SKU-2',
        nombre_producto='Redmi Note 13',
        marca='Xiaomi',
        precio_venta_claro=900_000,
        precio_venta_coltrade=800_000,
    )
    return pdv_norte, pdv_sur, caro, barato


# ── Modelos ────────────────────────────────────────────────────────────────

def test_las_llaves_primarias_son_los_codigos_del_negocio(catalogo):
    pdv, _sur, producto, _barato = catalogo
    assert pdv.pk == 'PDV-001'
    assert producto.pk == 'SKU-1'


def test_el_total_de_una_venta_usa_el_precio_coltrade(catalogo):
    pdv, _sur, producto, _barato = catalogo
    venta = Venta.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 3, 15),
        cantidad_vendida=3,
    )
    assert venta.total_coltrade == 12_000_000


def test_regional_y_materiales_pueden_ir_vacios():
    punto = PuntoVenta.objects.create(id_punto_venta='PDV-X', nombre_pdv='Sin clasificar')
    assert punto.regional is None
    assert punto.materiales is None


# ── Permisos ───────────────────────────────────────────────────────────────

def test_sin_la_app_no_se_entra(app_bi_trade):
    fuera = crear_usuario('fuera@supli.tech')
    assert cliente_de(fuera).get('/api/bi-trade/dashboard').status_code == 403


def test_consultar_solo_pide_la_app_pero_escribir_pide_permiso(app_bi_trade, catalogo):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    assert cliente_de(lector).get('/api/bi-trade/productos').status_code == 200

    cuerpo = {
        'idProducto': 'SKU-9',
        'nombreProducto': 'Producto nuevo',
        'marca': 'Motorola',
        'precioVentaClaro': 1_000_000,
        'precioVentaColtrade': 900_000,
    }
    ruta = '/api/bi-trade/productos'
    assert cliente_de(lector).post(ruta, cuerpo, format='json').status_code == 403
    assert cliente_de(editor).post(ruta, cuerpo, format='json').status_code == 201


# ── CRUD ───────────────────────────────────────────────────────────────────

def test_no_se_repite_el_codigo_de_un_punto_de_venta(app_bi_trade, catalogo):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    respuesta = cliente_de(editor).post(
        '/api/bi-trade/puntos-venta',
        {'idPuntoVenta': 'PDV-001', 'nombrePdv': 'Otro nombre'},
        format='json',
    )
    assert respuesta.status_code == 400


def test_no_se_borra_un_producto_con_ventas(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    Venta.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 1, 5),
        cantidad_vendida=1,
    )
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).delete(f'/api/bi-trade/productos/{producto.pk}')
    assert respuesta.status_code == 400
    assert Producto.objects.filter(pk=producto.pk).exists()


def test_registrar_una_venta_devuelve_su_total(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/ventas',
        {
            'idProducto': producto.pk,
            'idPuntoVenta': pdv.pk,
            'fechaVenta': '2026-04-01',
            'cantidadVendida': 2,
        },
        format='json',
    )
    assert respuesta.status_code == 201
    assert respuesta.data['total_coltrade'] == 8_000_000
    assert respuesta.data['nombre_pdv'] == 'Claro Unicentro'


def test_la_cantidad_vendida_no_puede_ser_cero(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    respuesta = cliente_de(editor).post(
        '/api/bi-trade/ventas',
        {
            'idProducto': producto.pk,
            'idPuntoVenta': pdv.pk,
            'fechaVenta': '2026-04-01',
            'cantidadVendida': 0,
        },
        format='json',
    )
    assert respuesta.status_code == 400


# ── Tablero ────────────────────────────────────────────────────────────────

def test_el_tablero_suma_unidades_e_ingresos(app_bi_trade, catalogo):
    pdv_norte, pdv_sur, caro, barato = catalogo
    Venta.objects.create(  # 8.000.000
        id_producto=caro,
        id_punto_venta=pdv_norte,
        fecha_venta=date(2026, 1, 10),
        cantidad_vendida=2,
    )
    Venta.objects.create(  # 4.000.000
        id_producto=barato,
        id_punto_venta=pdv_sur,
        fecha_venta=date(2026, 2, 10),
        cantidad_vendida=5,
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/dashboard').data

    assert datos['totales']['unidades'] == 7
    assert datos['totales']['ingresos'] == 12_000_000
    assert datos['totales']['operaciones'] == 2
    # La marca más cara encabeza el ranking por ingresos.
    assert datos['por_marca'][0]['label'] == 'Apple'
    assert datos['por_marca'][0]['participacion'] == pytest.approx(66.7)
    # Dos meses distintos, en orden cronológico.
    assert [fila['unidades'] for fila in datos['evolucion']] == [2, 5]


def test_el_tablero_se_puede_filtrar_por_regional(app_bi_trade, catalogo):
    pdv_norte, pdv_sur, caro, barato = catalogo
    Venta.objects.create(
        id_producto=caro,
        id_punto_venta=pdv_norte,
        fecha_venta=date(2026, 1, 10),
        cantidad_vendida=2,
    )
    Venta.objects.create(
        id_producto=barato,
        id_punto_venta=pdv_sur,
        fecha_venta=date(2026, 2, 10),
        cantidad_vendida=5,
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/dashboard', {'regional': 'Zona Sur'}).data

    assert datos['totales']['unidades'] == 5
    assert datos['totales']['ingresos'] == 4_000_000


def test_el_tablero_no_revienta_sin_ventas(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/dashboard').data

    assert datos['totales']['unidades'] == 0
    assert datos['totales']['ingresos'] == 0
    assert datos['totales']['ticket_promedio'] == 0
    assert datos['evolucion'] == []


# ── Inventario y metas ─────────────────────────────────────────────────────

def test_no_se_repite_el_par_producto_punto_en_inventario(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    Inventario.objects.create(id_producto=producto, id_punto_venta=pdv, cantidad_inventario=10)
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/inventario',
        {'idProducto': producto.pk, 'idPuntoVenta': pdv.pk, 'cantidadInventario': 5},
        format='json',
    )
    assert respuesta.status_code == 400
    assert Inventario.objects.count() == 1


def test_se_puede_editar_el_inventario_existente(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    registro = Inventario.objects.create(
        id_producto=producto, id_punto_venta=pdv, cantidad_inventario=10
    )
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).patch(
        f'/api/bi-trade/inventario/{registro.pk}',
        {'cantidadInventario': 25},
        format='json',
    )
    assert respuesta.status_code == 200
    registro.refresh_from_db()
    assert registro.cantidad_inventario == 25


def test_crear_una_meta_pide_permiso(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    cuerpo = {
        'idProducto': producto.pk,
        'idPuntoVenta': pdv.pk,
        'fechaMeta': '2026-01-01',
        'metaCantidad': 10,
    }
    assert cliente_de(lector).post('/api/bi-trade/metas', cuerpo, format='json').status_code == 403
    assert cliente_de(editor).post('/api/bi-trade/metas', cuerpo, format='json').status_code == 201


# ── Cumplimiento ───────────────────────────────────────────────────────────

def test_el_cumplimiento_compara_cada_medida_con_su_meta(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo  # iPhone: 4.000.000 Coltrade, 99 puntos

    Venta.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 1, 10),
        cantidad_vendida=8,
    )
    # Meta de 10 unidades → 8/10 = 80% en las tres medidas.
    MetaComercial.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_meta=date(2026, 1, 1),
        meta_cantidad=10,
    )
    Inventario.objects.create(id_producto=producto, id_punto_venta=pdv, cantidad_inventario=16)

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/cumplimiento').data
    totales = datos['totales']

    assert totales['real_cantidad'] == 8
    assert totales['real_dinero'] == 32_000_000
    assert totales['real_puntos'] == 792
    assert totales['cumplimiento_cantidad'] == pytest.approx(80.0)
    assert totales['cumplimiento_dinero'] == pytest.approx(80.0)
    assert totales['cumplimiento_puntos'] == pytest.approx(80.0)
    # 16 en stock contra 8 vendidas: alcanza para dos vueltas más.
    assert totales['inventario_unidades'] == 16
    assert totales['cobertura'] == pytest.approx(2.0)


def test_un_producto_sin_puntaje_no_anula_la_suma_de_puntos(app_bi_trade, catalogo):
    pdv, _sur, caro, sin_puntaje = catalogo  # `sin_puntaje` no trae puntaje

    Venta.objects.create(
        id_producto=caro, id_punto_venta=pdv, fecha_venta=date(2026, 1, 10), cantidad_vendida=2
    )
    Venta.objects.create(
        id_producto=sin_puntaje,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 1, 11),
        cantidad_vendida=5,
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/cumplimiento').data

    # Solo aportan puntos las 2 unidades del producto que sí tiene puntaje.
    assert datos['totales']['real_puntos'] == 198


def test_sin_meta_cargada_el_cumplimiento_es_cero_y_no_revienta(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    Venta.objects.create(
        id_producto=producto, id_punto_venta=pdv, fecha_venta=date(2026, 1, 10), cantidad_vendida=3
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    totales = cliente_de(lector).get('/api/bi-trade/cumplimiento').data['totales']

    assert totales['meta_cantidad'] == 0
    assert totales['cumplimiento_cantidad'] == 0.0
    assert totales['cumplimiento_dinero'] == 0.0


def test_el_cumplimiento_aparece_aunque_no_haya_ventas(app_bi_trade, catalogo):
    """Una meta sin ventas debe verse como 0%, no desaparecer del listado."""
    pdv, _sur, producto, _barato = catalogo
    MetaComercial.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_meta=date(2026, 1, 1),
        meta_cantidad=10,
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/cumplimiento').data

    filas = datos['por_punto_venta']
    assert len(filas) == 1
    assert filas[0]['label'] == 'Claro Unicentro'
    assert filas[0]['meta_cantidad'] == 10
    assert filas[0]['real_cantidad'] == 0
    assert filas[0]['cumplimiento_cantidad'] == 0.0


def test_el_cumplimiento_se_filtra_por_regional(app_bi_trade, catalogo):
    pdv_norte, pdv_sur, caro, barato = catalogo

    Venta.objects.create(
        id_producto=caro,
        id_punto_venta=pdv_norte,
        fecha_venta=date(2026, 1, 10),
        cantidad_vendida=2,
    )
    Venta.objects.create(
        id_producto=barato,
        id_punto_venta=pdv_sur,
        fecha_venta=date(2026, 2, 10),
        cantidad_vendida=5,
    )
    MetaComercial.objects.create(
        id_producto=caro,
        id_punto_venta=pdv_norte,
        fecha_meta=date(2026, 1, 1),
        meta_cantidad=4,
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/cumplimiento', {'regional': 'Zona Norte'}).data

    assert datos['totales']['real_cantidad'] == 2
    assert datos['totales']['cumplimiento_cantidad'] == pytest.approx(50.0)


# ── Borrado masivo ─────────────────────────────────────────────────────────

def test_vaciar_productos_se_bloquea_y_dice_qué_lo_impide(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    Venta.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 1, 5),
        cantidad_vendida=1,
    )
    Inventario.objects.create(id_producto=producto, id_punto_venta=pdv, cantidad_inventario=4)
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).delete('/api/bi-trade/productos/eliminar-todos')

    assert respuesta.status_code == 400
    mensaje = respuesta.data['message']
    assert 'No es posible eliminar los productos' in mensaje
    assert '1 ventas' in mensaje
    assert '1 registros de inventario' in mensaje
    assert Producto.objects.count() == 2


def test_vaciar_puntos_de_venta_se_bloquea_por_metas(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    MetaComercial.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_meta=date(2026, 1, 1),
        meta_cantidad=5,
    )
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).delete('/api/bi-trade/puntos-venta/eliminar-todos')

    assert respuesta.status_code == 400
    assert '1 metas' in respuesta.data['message']
    assert PuntoVenta.objects.count() == 2


def test_vaciar_ventas_siempre_se_puede(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    for dia in (5, 6, 7):
        Venta.objects.create(
            id_producto=producto,
            id_punto_venta=pdv,
            fecha_venta=date(2026, 1, dia),
            cantidad_vendida=1,
        )
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).delete('/api/bi-trade/ventas/eliminar-todos')

    assert respuesta.status_code == 200
    assert respuesta.data['deleted'] == 3
    assert Venta.objects.count() == 0
    # Vaciar las ventas no se lleva por delante el catálogo.
    assert Producto.objects.count() == 2
    assert PuntoVenta.objects.count() == 2


def test_al_quedar_sin_dependencias_ya_se_pueden_vaciar_los_productos(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    Venta.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 1, 5),
        cantidad_vendida=1,
    )
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)

    assert cliente.delete('/api/bi-trade/productos/eliminar-todos').status_code == 400
    assert cliente.delete('/api/bi-trade/ventas/eliminar-todos').status_code == 200
    assert cliente.delete('/api/bi-trade/productos/eliminar-todos').status_code == 200
    assert Producto.objects.count() == 0


def test_vaciar_una_tabla_ya_vacía_no_es_un_error(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    respuesta = cliente_de(editor).delete('/api/bi-trade/ventas/eliminar-todos')

    assert respuesta.status_code == 200
    assert respuesta.data['deleted'] == 0


def test_vaciar_exige_el_permiso_de_edicion(app_bi_trade, catalogo):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    assert (
        cliente_de(lector).delete('/api/bi-trade/productos/eliminar-todos').status_code == 403
    )
    assert Producto.objects.count() == 2


def test_eliminar_todos_no_choca_con_el_detalle_de_un_registro(app_bi_trade, catalogo):
    """Un producto con ese código no debe capturar la ruta del borrado masivo."""
    Producto.objects.create(
        id_producto='eliminar-todos',
        nombre_producto='Producto con nombre traicionero',
        marca='Test',
        precio_venta_claro=1000,
        precio_venta_coltrade=900,
    )
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).delete('/api/bi-trade/productos/eliminar-todos')

    assert respuesta.status_code == 200
    assert Producto.objects.count() == 0


# ── Excel: plantilla e importación ─────────────────────────────────────────

def _hoja_de(respuesta):
    """Abre el .xlsx que devolvió la API."""
    return openpyxl.load_workbook(BytesIO(respuesta.content))


def _archivo_xlsx(encabezados, filas):
    """Arma en memoria un .xlsx como el que subiría una persona."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Datos'
    hoja.append(list(encabezados))
    for fila in filas:
        hoja.append(list(fila))
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    buffer.name = 'carga.xlsx'
    return buffer


def test_la_plantilla_trae_los_encabezados_y_las_instrucciones(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    respuesta = cliente_de(editor).get('/api/bi-trade/productos/plantilla')

    assert respuesta.status_code == 200
    assert respuesta['Content-Disposition'].endswith('.xlsx"')

    libro = _hoja_de(respuesta)
    assert libro.sheetnames == ['Datos', 'Instrucciones']
    encabezados = [celda.value for celda in libro['Datos'][1]]
    assert encabezados == [
        'id_producto',
        'nombre_producto',
        'marca',
        'precio_venta_claro',
        'precio_venta_coltrade',
        'puntaje',
    ]


def test_la_plantilla_de_puntos_lista_las_opciones_validas(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    libro = _hoja_de(cliente_de(editor).get('/api/bi-trade/puntos-venta/plantilla'))

    texto = ' '.join(
        str(celda.value)
        for fila in libro['Instrucciones'].iter_rows()
        for celda in fila
        if celda.value
    )
    assert 'Zona Sur' in texto
    assert 'Todos los materiales' in texto


def test_descargar_la_plantilla_exige_permiso(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    assert cliente_de(lector).get('/api/bi-trade/productos/plantilla').status_code == 403


def test_importar_productos_crea_y_luego_actualiza(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)
    encabezados = (
        'id_producto',
        'nombre_producto',
        'marca',
        'precio_venta_claro',
        'precio_venta_coltrade',
        'puntaje',
    )

    archivo = _archivo_xlsx(encabezados, [('SKU-9', 'Producto nuevo', 'Nokia', 500000, 450000, 60)])
    respuesta = cliente.post(
        '/api/bi-trade/productos/importar', {'archivo': archivo}, format='multipart'
    )
    assert respuesta.status_code == 200
    assert respuesta.data['created'] == 1
    assert Producto.objects.get(pk='SKU-9').marca == 'Nokia'

    # El mismo código no duplica: actualiza.
    archivo = _archivo_xlsx(
        encabezados, [('SKU-9', 'Producto corregido', 'Nokia', 600000, 550000, 65)]
    )
    respuesta = cliente.post(
        '/api/bi-trade/productos/importar', {'archivo': archivo}, format='multipart'
    )
    assert respuesta.status_code == 200
    assert respuesta.data['updated'] == 1
    assert Producto.objects.count() == 1
    assert Producto.objects.get(pk='SKU-9').nombre_producto == 'Producto corregido'


def test_una_fila_mala_cancela_toda_la_importacion(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    encabezados = (
        'id_producto',
        'nombre_producto',
        'marca',
        'precio_venta_claro',
        'precio_venta_coltrade',
        'puntaje',
    )
    archivo = _archivo_xlsx(
        encabezados,
        [
            ('SKU-A', 'Producto bueno', 'Nokia', 500000, 450000, 60),
            ('SKU-B', 'Producto malo', 'Nokia', 'no es un número', 450000, 60),
        ],
    )

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/productos/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 400
    assert respuesta.data['filas'][0]['fila'] == 3
    assert 'precio_venta_claro' in respuesta.data['filas'][0]['errores'][0]
    # Ni siquiera la fila buena se guardó.
    assert Producto.objects.count() == 0


def test_importar_avisa_si_falta_una_columna(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    archivo = _archivo_xlsx(('id_producto', 'nombre_producto'), [('SKU-1', 'Incompleto')])

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/productos/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 400
    assert 'le faltan columnas' in respuesta.data['message']
    assert 'marca' in respuesta.data['message']


def test_importar_ventas_lee_las_fechas(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    archivo = _archivo_xlsx(
        ('id_producto', 'id_punto_venta', 'fecha_venta', 'cantidad_vendida'),
        [(producto.pk, pdv.pk, '2026-03-15', 4)],
    )

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/ventas/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 200
    venta = Venta.objects.get()
    assert venta.fecha_venta == date(2026, 3, 15)
    assert venta.cantidad_vendida == 4


def test_importar_inventario_actualiza_el_par_existente(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    Inventario.objects.create(id_producto=producto, id_punto_venta=pdv, cantidad_inventario=5)
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    archivo = _archivo_xlsx(
        ('id_producto', 'id_punto_venta', 'cantidad_inventario'),
        [(producto.pk, pdv.pk, 40)],
    )

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/inventario/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 200
    assert respuesta.data['updated'] == 1
    assert Inventario.objects.count() == 1
    assert Inventario.objects.get().cantidad_inventario == 40


def test_importar_rechaza_un_archivo_que_no_es_xlsx(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    falso = BytesIO(b'columna;otra\n1;2')
    falso.name = 'datos.csv'

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/productos/importar', {'archivo': falso}, format='multipart'
    )

    assert respuesta.status_code == 400
    assert '.xlsx' in respuesta.data['message']


def test_importar_exige_permiso_de_edicion(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    archivo = _archivo_xlsx(('id_producto',), [('SKU-1',)])

    respuesta = cliente_de(lector).post(
        '/api/bi-trade/productos/importar', {'archivo': archivo}, format='multipart'
    )
    assert respuesta.status_code == 403


def test_la_plantilla_se_puede_llenar_y_reimportar_tal_cual(app_bi_trade):
    """El ciclo real: descargar, borrar el ejemplo, llenar y subir."""
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)

    libro = _hoja_de(cliente.get('/api/bi-trade/puntos-venta/plantilla'))
    hoja = libro['Datos']
    hoja.delete_rows(2)  # la fila de ejemplo
    hoja.append(['PDV-77', 'Claro Plaza Nueva', 'Zona Norte', 'Todos los materiales'])

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    buffer.name = 'plantilla-llena.xlsx'

    respuesta = cliente.post(
        '/api/bi-trade/puntos-venta/importar', {'archivo': buffer}, format='multipart'
    )

    assert respuesta.status_code == 200
    punto = PuntoVenta.objects.get(pk='PDV-77')
    assert punto.nombre_pdv == 'Claro Plaza Nueva'
    assert punto.regional == 'Zona Norte'


# ── La fecha de la meta ────────────────────────────────────────────────────

def test_la_meta_necesita_fecha(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/metas',
        {
            'idProducto': producto.pk,
            'idPuntoVenta': pdv.pk,
            'metaCantidad': 10,
        },
        format='json',
    )
    assert respuesta.status_code == 400
    assert 'fecha_meta' in respuesta.data['errors']


def test_el_mismo_par_admite_metas_de_periodos_distintos(app_bi_trade, catalogo):
    """Ese es el punto de la fecha: marzo y abril no se pisan."""
    pdv, _sur, producto, _barato = catalogo
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)

    def meta(fecha):
        return cliente.post(
            '/api/bi-trade/metas',
            {
                'idProducto': producto.pk,
                'idPuntoVenta': pdv.pk,
                'fechaMeta': fecha,
                'metaCantidad': 10,
            },
            format='json',
        )

    assert meta('2026-03-01').status_code == 201
    assert meta('2026-04-01').status_code == 201
    # Repetir el mismo periodo sí se rechaza.
    repetida = meta('2026-03-01')
    assert repetida.status_code == 400
    assert 'cambia la fecha' in repetida.data['errors']['non_field_errors'][0]
    assert MetaComercial.objects.count() == 2


def test_el_cumplimiento_se_puede_acotar_a_un_periodo(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo  # iPhone: 4.000.000 Coltrade

    # Marzo: se vendieron 8 contra una meta de 10 → 80%.
    Venta.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 3, 15),
        cantidad_vendida=8,
    )
    MetaComercial.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_meta=date(2026, 3, 1),
        meta_cantidad=10,
    )
    # Abril: se vendieron 20 contra una meta de 10 → 200%.
    Venta.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 4, 15),
        cantidad_vendida=20,
    )
    MetaComercial.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_meta=date(2026, 4, 1),
        meta_cantidad=10,
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    # Sin periodo se suma todo: 28 contra 20.
    todo = cliente.get('/api/bi-trade/cumplimiento').data['totales']
    assert todo['real_cantidad'] == 28
    assert todo['meta_cantidad'] == 20
    assert todo['cumplimiento_cantidad'] == pytest.approx(140.0)

    # Acotado a marzo, cada lado se recorta por su propia fecha.
    marzo = cliente.get(
        '/api/bi-trade/cumplimiento', {'desde': '2026-03-01', 'hasta': '2026-03-31'}
    ).data['totales']
    assert marzo['real_cantidad'] == 8
    assert marzo['meta_cantidad'] == 10
    assert marzo['cumplimiento_cantidad'] == pytest.approx(80.0)

    abril = cliente.get(
        '/api/bi-trade/cumplimiento', {'desde': '2026-04-01', 'hasta': '2026-04-30'}
    ).data['totales']
    assert abril['real_cantidad'] == 20
    assert abril['cumplimiento_cantidad'] == pytest.approx(200.0)


def test_una_fecha_ilegible_se_ignora_en_vez_de_reventar(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    Venta.objects.create(
        id_producto=producto,
        id_punto_venta=pdv,
        fecha_venta=date(2026, 3, 15),
        cantidad_vendida=8,
    )
    lector = crear_usuario('lector@supli.tech', app_bi_trade)

    respuesta = cliente_de(lector).get('/api/bi-trade/cumplimiento', {'desde': '15/03/2026'})

    assert respuesta.status_code == 200
    assert respuesta.data['totales']['real_cantidad'] == 8


def test_la_plantilla_de_metas_incluye_la_fecha(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    libro = _hoja_de(cliente_de(editor).get('/api/bi-trade/metas/plantilla'))

    encabezados = [celda.value for celda in libro['Datos'][1]]
    assert encabezados == [
        'id_producto',
        'id_punto_venta',
        'fecha_meta',
        'meta_cantidad',
    ]


def test_importar_metas_separa_los_periodos(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    encabezados = (
        'id_producto',
        'id_punto_venta',
        'fecha_meta',
        'meta_cantidad',
    )
    archivo = _archivo_xlsx(
        encabezados,
        [
            (producto.pk, pdv.pk, '2026-03-01', 10, 40_000_000, 990),
            (producto.pk, pdv.pk, '2026-04-01', 12, 48_000_000, 1188),
        ],
    )

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/metas/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 200
    assert respuesta.data['created'] == 2
    assert MetaComercial.objects.count() == 2

    # Reimportar el mismo periodo actualiza; no duplica.
    archivo = _archivo_xlsx(
        encabezados, [(producto.pk, pdv.pk, '2026-03-01', 99, 40_000_000, 990)]
    )
    respuesta = cliente_de(editor).post(
        '/api/bi-trade/metas/importar', {'archivo': archivo}, format='multipart'
    )
    assert respuesta.data['updated'] == 1
    assert MetaComercial.objects.count() == 2
    assert MetaComercial.objects.get(fecha_meta=date(2026, 3, 1)).meta_cantidad == 99


# ── Calendario laboral ─────────────────────────────────────────────────────

@pytest.mark.parametrize(
    ('anio', 'esperado'),
    [(2024, date(2024, 3, 31)), (2025, date(2025, 4, 20)), (2026, date(2026, 4, 5))],
)
def test_la_pascua_cae_donde_debe(anio, esperado):
    assert pascua(anio) == esperado


def test_colombia_tiene_dieciocho_festivos_al_año():
    for anio in (2024, 2026, 2027, 2028):
        assert len(festivos(anio)) == 18


def test_dos_festivos_pueden_caer_el_mismo_dia():
    """En 2025 el Sagrado Corazón y San Pedro y San Pablo coinciden el 30 de junio.

    Por eso `festivos` devuelve un conjunto de fechas y no una lista de nombres:
    lo que descuenta un día hábil es la fecha, y ese día se pierde una sola vez.
    """
    assert date(2025, 6, 30) in festivos(2025)
    assert len(festivos(2025)) == 17


def test_los_festivos_trasladables_caen_en_lunes():
    # San José 2026 es jueves 19 de marzo; la Ley Emiliani lo corre al 23.
    assert date(2026, 3, 23) in festivos(2026)
    assert date(2026, 3, 19) not in festivos(2026)


def test_el_sabado_es_habil_y_el_domingo_no():
    assert es_habil(date(2026, 3, 21))  # sábado
    assert not es_habil(date(2026, 3, 22))  # domingo


def test_marzo_de_2026_tiene_veinticinco_dias_habiles():
    # 31 días − 5 domingos − 1 festivo (San José corrido al lunes 23).
    assert len(dias_habiles(2026, 3)) == 25


def test_un_festivo_en_domingo_no_se_descuenta_dos_veces():
    # 1 de noviembre de 2026 es domingo; Todos los Santos se corre al lunes 2,
    # así que noviembre pierde 5 domingos y 2 festivos, no 3.
    dias = dias_habiles(2026, 11)
    assert len(dias) == 30 - 5 - 2


# ── Avance mensual ─────────────────────────────────────────────────────────

def test_la_meta_del_mes_se_reparte_entre_los_dias_habiles(app_bi_trade, catalogo):
    """El día de la fecha_meta es irrelevante: lo que importa es el mes."""
    pdv, _sur, caro, barato = catalogo
    # Tres metas de marzo con días distintos: juntas son la meta de marzo.
    # 25 × 4.000.000 + 125 × 800.000 + 25 × 4.000.000 = 300.000.000.
    for producto, dia, unidades in ((caro, 1, 25), (barato, 14, 125), (caro, 15, 25)):
        MetaComercial.objects.create(
            id_producto=producto,
            id_punto_venta=pdv,
            fecha_meta=date(2026, 3, dia),
            meta_cantidad=unidades,
        )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/avance-mensual?anio=2026&mes=3').json()

    assert datos['periodo']['diasHabiles'] == 25
    assert datos['totales']['metaDinero'] == 300_000_000
    assert datos['totales']['metaDiaria'] == 12_000_000
    # La serie cubre el mes completo, no solo los días hábiles.
    assert len(datos['serie']) == 31
    assert {fila['metaDiaria'] for fila in datos['serie']} == {12_000_000}


def test_el_cumplimiento_del_dia_compara_la_venta_con_la_cuota_diaria(app_bi_trade, catalogo):
    pdv, _sur, caro, _barato = catalogo
    MetaComercial.objects.create(
        id_producto=caro,
        id_punto_venta=pdv,
        fecha_meta=date(2026, 3, 1),
        # 25 × 4.000.000 = 100.000.000 de meta para el mes.
        meta_cantidad=25,
    )
    # Una venta el día 2: 5 × 4.000.000 = 20.000.000.
    Venta.objects.create(
        id_producto=caro, id_punto_venta=pdv, fecha_venta=date(2026, 3, 2), cantidad_vendida=5
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/avance-mensual?anio=2026&mes=3').json()

    diaria = datos['totales']['metaDiaria']
    assert diaria == round(100_000_000 / 25)
    dia2 = next(fila for fila in datos['serie'] if fila['dia'] == 2)
    assert dia2['ventas'] == 20_000_000
    assert dia2['cumplimiento'] == round(20_000_000 * 100 / diaria, 1)
    # Los días sin venta quedan en cero, no ausentes: la línea no se corta.
    assert next(fila for fila in datos['serie'] if fila['dia'] == 3)['ventas'] == 0


def test_el_avance_solo_toma_el_mes_pedido(app_bi_trade, catalogo):
    pdv, _sur, caro, _barato = catalogo
    for mes in (2, 3):
        MetaComercial.objects.create(
            id_producto=caro,
            id_punto_venta=pdv,
            fecha_meta=date(2026, mes, 10),
            meta_cantidad=10,
        )
        Venta.objects.create(
            id_producto=caro,
            id_punto_venta=pdv,
            fecha_venta=date(2026, mes, 10),
            cantidad_vendida=mes,
        )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    marzo = cliente_de(lector).get('/api/bi-trade/avance-mensual?anio=2026&mes=3').json()

    assert marzo['totales']['metaDinero'] == 40_000_000
    assert marzo['totales']['ventasCantidad'] == 3


def test_el_avance_agrupa_por_regional_y_por_punto_de_venta(app_bi_trade, catalogo):
    norte, sur, caro, barato = catalogo
    for punto, producto in ((norte, caro), (sur, barato)):
        MetaComercial.objects.create(
            id_producto=producto,
            id_punto_venta=punto,
            fecha_meta=date(2026, 3, 1),
            meta_cantidad=10,
        )
    Inventario.objects.create(id_producto=caro, id_punto_venta=norte, cantidad_inventario=7)
    # 10 × 4.000.000 = 40.000.000 → el norte cumple al 100%.
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=10
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/avance-mensual?anio=2026&mes=3').json()

    por_regional = {fila['label']: fila for fila in datos['porRegional']}
    assert por_regional['Zona Norte']['metaMensual'] == 40_000_000
    assert por_regional['Zona Norte']['importe'] == 40_000_000
    assert por_regional['Zona Norte']['cantidad'] == 10
    assert por_regional['Zona Norte']['cumplimiento'] == 100.0
    assert por_regional['Zona Sur']['cumplimiento'] == 0.0

    por_punto = {fila['label']: fila for fila in datos['porPuntoVenta']}
    assert por_punto['Claro Unicentro']['inventario'] == 7


def test_el_avance_se_filtra_por_regional(app_bi_trade, catalogo):
    norte, sur, caro, barato = catalogo
    for punto, producto in ((norte, caro), (sur, barato)):
        MetaComercial.objects.create(
            id_producto=producto,
            id_punto_venta=punto,
            fecha_meta=date(2026, 3, 1),
            meta_cantidad=10,
        )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get(
        '/api/bi-trade/avance-mensual?anio=2026&mes=3&regional=Zona Norte'
    ).json()

    assert datos['totales']['metaDinero'] == 40_000_000
    assert [fila['label'] for fila in datos['porRegional']] == ['Zona Norte']


def test_sin_meta_ni_ventas_el_avance_no_revienta(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).get('/api/bi-trade/avance-mensual?anio=2026&mes=3')

    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert datos['totales']['metaDiaria'] == 0
    assert datos['totales']['cumplimiento'] == 0.0
    assert len(datos['serie']) == 31


def test_un_mes_invalido_cae_en_el_mes_en_curso(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    hoy = timezone.localdate()

    for consulta in ('anio=2026&mes=13', 'anio=abc&mes=3', 'anio=1800&mes=3'):
        datos = cliente_de(lector).get(f'/api/bi-trade/avance-mensual?{consulta}').json()
        assert (datos['periodo']['anio'], datos['periodo']['mes']) == (hoy.year, hoy.month)


def test_el_avance_pide_la_app(app_bi_trade):
    fuera = crear_usuario('fuera@supli.tech')
    assert cliente_de(fuera).get('/api/bi-trade/avance-mensual').status_code == 403


# ── Listados: paginación, totales y filtros ────────────────────────────────

@pytest.fixture
def muchas_ventas(catalogo):
    """Cuarenta ventas: veinte del norte con el caro, veinte del sur con el barato.

    Son más de una página a propósito: es justo el caso donde sumar lo que se
    ve en la tabla daría un número equivocado.
    """
    norte, sur, caro, barato = catalogo
    for indice in range(20):
        Venta.objects.create(
            id_producto=caro,
            id_punto_venta=norte,
            fecha_venta=date(2026, 3, 2 + indice % 10),
            cantidad_vendida=2,
        )
        Venta.objects.create(
            id_producto=barato,
            id_punto_venta=sur,
            fecha_venta=date(2026, 4, 1 + indice % 10),
            cantidad_vendida=3,
        )
    return catalogo


def test_el_listado_de_ventas_trae_quince_filas_y_el_total_real(app_bi_trade, muchas_ventas):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/ventas').json()

    assert len(datos['items']) == 15
    assert datos['total'] == 40
    assert datos['pageSize'] == 15
    assert datos['page'] == 1


def test_la_ultima_pagina_trae_el_resto(app_bi_trade, muchas_ventas):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/ventas?page=3').json()

    assert len(datos['items']) == 10
    assert datos['total'] == 40


def test_el_resumen_suma_todo_lo_filtrado_no_solo_la_pagina(app_bi_trade, muchas_ventas):
    """Lo que pide el negocio: la tarjeta suma los 40 registros, no los 15 visibles."""
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    pagina = cliente.get('/api/bi-trade/ventas').json()
    resumen = cliente.get('/api/bi-trade/ventas/resumen').json()

    # 20 × 2 unidades del caro + 20 × 3 del barato.
    assert resumen['registros'] == 40
    assert resumen['unidades'] == 20 * 2 + 20 * 3
    # 40 × 4.000.000 + 60 × 800.000.
    assert resumen['ingresos'] == 40 * 4_000_000 + 60 * 800_000
    assert resumen['productos'] == 2
    assert resumen['puntosVenta'] == 2

    # Y es distinto de sumar la página, que es justo el error a evitar.
    de_la_pagina = sum(fila['cantidadVendida'] for fila in pagina['items'])
    assert de_la_pagina < resumen['unidades']


def test_el_resumen_respeta_los_filtros(app_bi_trade, muchas_ventas):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    norte = cliente.get('/api/bi-trade/ventas/resumen?regional=Zona Norte').json()
    assert norte['registros'] == 20
    assert norte['unidades'] == 40
    assert norte['ingresos'] == 40 * 4_000_000

    marca = cliente.get('/api/bi-trade/ventas/resumen?marca=Xiaomi').json()
    assert marca['registros'] == 20
    assert marca['unidades'] == 60


def test_las_ventas_se_filtran_por_rango_de_fechas(app_bi_trade, muchas_ventas):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    # Las del caro son de marzo; las del barato, de abril.
    marzo = cliente.get('/api/bi-trade/ventas/resumen?desde=2026-03-01&hasta=2026-03-31').json()
    assert marzo['registros'] == 20
    assert marzo['productos'] == 1

    por_mes = cliente.get('/api/bi-trade/ventas/resumen?anio=2026&mes=4').json()
    assert por_mes['registros'] == 20


def test_las_ventas_se_filtran_por_cantidad(app_bi_trade, muchas_ventas):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    assert cliente.get('/api/bi-trade/ventas/resumen?cantidad_min=3').json()['registros'] == 20
    assert cliente.get('/api/bi-trade/ventas/resumen?cantidad_max=2').json()['registros'] == 20
    assert cliente.get('/api/bi-trade/ventas/resumen?cantidad_min=4').json()['registros'] == 0


def test_la_busqueda_encuentra_por_codigo_y_por_marca(app_bi_trade, muchas_ventas):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    assert cliente.get('/api/bi-trade/ventas/resumen?search=SKU-1').json()['registros'] == 20
    assert cliente.get('/api/bi-trade/ventas/resumen?search=Xiaomi').json()['registros'] == 20
    assert cliente.get('/api/bi-trade/ventas/resumen?search=Unicentro').json()['registros'] == 20


def test_el_resumen_de_una_tabla_vacia_es_cero_y_no_nulo(app_bi_trade):
    """Sin filas, `Sum` devuelve None: el resumen tiene que traducirlo a 0."""
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    resumen = cliente_de(lector).get('/api/bi-trade/ventas/resumen').json()

    assert resumen == {
        'registros': 0,
        'unidades': 0,
        'ingresos': 0,
        'puntos': 0,
        'productos': 0,
        'puntosVenta': 0,
    }


def test_el_resumen_de_inventario_cuenta_agotados_y_valoriza(app_bi_trade, catalogo):
    norte, sur, caro, barato = catalogo
    Inventario.objects.create(id_producto=caro, id_punto_venta=norte, cantidad_inventario=3)
    Inventario.objects.create(id_producto=barato, id_punto_venta=sur, cantidad_inventario=0)

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)
    resumen = cliente.get('/api/bi-trade/inventario/resumen').json()

    assert resumen['registros'] == 2
    assert resumen['unidades'] == 3
    assert resumen['valorizadoTotal'] == 3 * 4_000_000
    assert resumen['agotados'] == 1

    solo_agotados = cliente.get('/api/bi-trade/inventario/resumen?agotado=true').json()
    assert solo_agotados['registros'] == 1
    assert solo_agotados['unidades'] == 0

    con_stock = cliente.get('/api/bi-trade/inventario/resumen?agotado=false').json()
    assert con_stock['registros'] == 1


def test_el_resumen_de_metas_suma_las_tres_medidas(app_bi_trade, catalogo):
    norte, sur, caro, barato = catalogo
    for punto, producto in ((norte, caro), (sur, barato)):
        MetaComercial.objects.create(
            id_producto=producto,
            id_punto_venta=punto,
            fecha_meta=date(2026, 3, 1),
            meta_cantidad=10,
        )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)
    resumen = cliente.get('/api/bi-trade/metas/resumen').json()

    assert resumen['registros'] == 2
    assert resumen['metaCantidad'] == 20
    assert resumen['metaDinero'] == 40_000_000 + 8_000_000
    assert resumen['metaPuntos'] == 10 * 99

    por_mes = cliente.get('/api/bi-trade/metas/resumen?anio=2026&mes=3').json()
    assert por_mes['registros'] == 2
    assert cliente.get('/api/bi-trade/metas/resumen?anio=2026&mes=4').json()['registros'] == 0


def test_el_resumen_pide_la_app(app_bi_trade):
    fuera = crear_usuario('fuera@supli.tech')
    assert cliente_de(fuera).get('/api/bi-trade/ventas/resumen').status_code == 403


# ── Exportación ────────────────────────────────────────────────────────────

def test_exportar_ventas_trae_todas_las_filas_no_solo_la_pagina(app_bi_trade, muchas_ventas):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).get('/api/bi-trade/ventas/exportar')

    assert respuesta.status_code == 200
    assert respuesta['Content-Disposition'].startswith('attachment; filename="ventas-')

    hoja = _hoja_de(respuesta).active
    # 40 filas de datos más el encabezado, aunque la tabla muestre 15.
    assert hoja.max_row == 41
    assert [celda.value for celda in hoja[1]][:4] == [
        'Fecha',
        'Código producto',
        'Producto',
        'Marca',
    ]


def test_exportar_respeta_el_filtro(app_bi_trade, muchas_ventas):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).get('/api/bi-trade/ventas/exportar?regional=Zona Sur')
    hoja = _hoja_de(respuesta).active

    assert hoja.max_row == 21
    assert {celda.value for celda in hoja['G'][1:]} == {'Zona Sur'}


def test_el_excel_exportado_calcula_el_total_de_cada_venta(app_bi_trade, catalogo):
    norte, _sur, caro, _barato = catalogo
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=3
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    hoja = _hoja_de(cliente_de(lector).get('/api/bi-trade/ventas/exportar')).active

    encabezados = [celda.value for celda in hoja[1]]
    fila = {nombre: celda.value for nombre, celda in zip(encabezados, hoja[2], strict=True)}
    assert fila['Cantidad vendida'] == 3
    assert fila['Total'] == 3 * 4_000_000


def test_exportar_una_tabla_vacia_devuelve_solo_el_encabezado(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    hoja = _hoja_de(cliente_de(lector).get('/api/bi-trade/metas/exportar')).active

    assert hoja.max_row == 1
    assert hoja['A1'].value == 'Periodo'


def test_exportar_pide_la_app(app_bi_trade):
    fuera = crear_usuario('fuera@supli.tech')
    assert cliente_de(fuera).get('/api/bi-trade/ventas/exportar').status_code == 403


def test_exportar_solo_pide_la_app_no_el_permiso_de_edicion(app_bi_trade, catalogo):
    """Descargar datos es lectura: no debe exigir el permiso de escribir."""
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    assert cliente_de(lector).get('/api/bi-trade/inventario/exportar').status_code == 200


# ── Exportación del tablero ────────────────────────────────────────────────

@pytest.fixture
def marzo_con_datos(catalogo):
    """Meta y ventas de marzo de 2026 en las dos regionales."""
    norte, sur, caro, barato = catalogo
    for punto, producto in ((norte, caro), (sur, barato)):
        MetaComercial.objects.create(
            id_producto=producto,
            id_punto_venta=punto,
            fecha_meta=date(2026, 3, 1),
            meta_cantidad=10,
        )
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=5
    )
    return catalogo


def test_el_tablero_se_exporta_en_tres_hojas(app_bi_trade, marzo_con_datos):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).get('/api/bi-trade/avance-mensual/exportar?anio=2026&mes=3')

    assert respuesta.status_code == 200
    assert respuesta['Content-Disposition'] == 'attachment; filename="avance-2026-03.xlsx"'

    libro = _hoja_de(respuesta)
    assert libro.sheetnames == ['Día por día', 'Por regional', 'Por punto de venta']
    # Marzo tiene 31 días, y la serie los trae todos.
    assert libro['Día por día'].max_row == 32
    assert libro['Por regional'].max_row == 3


def test_una_hoja_sola_baja_sin_las_otras(app_bi_trade, marzo_con_datos):
    """Es lo que pide el botón chico de cada tarjeta."""
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).get(
        '/api/bi-trade/avance-mensual/exportar?anio=2026&mes=3&hoja=regional'
    )

    libro = _hoja_de(respuesta)
    assert libro.sheetnames == ['Por regional']
    assert 'avance-2026-03-regional.xlsx' in respuesta['Content-Disposition']


def test_el_excel_del_tablero_coincide_con_el_json(app_bi_trade, marzo_con_datos):
    """El archivo y la pantalla salen del mismo cálculo: no pueden diferir."""
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    datos = cliente.get('/api/bi-trade/avance-mensual?anio=2026&mes=3').json()
    hoja = _hoja_de(
        cliente.get('/api/bi-trade/avance-mensual/exportar?anio=2026&mes=3&hoja=serie')
    )['Día por día']

    encabezados = [celda.value for celda in hoja[1]]
    assert encabezados[:3] == ['Día', 'Fecha', '¿Hábil?']

    # El día 4 de marzo de 2026 es miércoles: hábil, y con la venta cargada.
    fila = {nombre: celda.value for nombre, celda in zip(encabezados, hoja[5], strict=True)}
    del_json = next(d for d in datos['serie'] if d['dia'] == 4)
    assert fila['Día'] == 4
    assert fila['¿Hábil?'] == 'Sí'
    assert fila['Ventas del día'] == del_json['ventas'] == 20_000_000
    assert fila['Meta diaria'] == del_json['metaDiaria']
    assert fila['Cumplimiento %'] == del_json['cumplimiento']


def test_el_domingo_sale_como_no_habil_en_el_excel(app_bi_trade, marzo_con_datos):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    hoja = _hoja_de(
        cliente_de(lector).get('/api/bi-trade/avance-mensual/exportar?anio=2026&mes=3&hoja=serie')
    )['Día por día']

    # 1 de marzo de 2026 es domingo; el 2, lunes.
    assert [hoja.cell(row=2, column=1).value, hoja.cell(row=2, column=3).value] == [1, 'No']
    assert [hoja.cell(row=3, column=1).value, hoja.cell(row=3, column=3).value] == [2, 'Sí']


def test_la_exportacion_del_tablero_respeta_los_filtros(app_bi_trade, marzo_con_datos):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    hoja = _hoja_de(
        cliente_de(lector).get(
            '/api/bi-trade/avance-mensual/exportar?anio=2026&mes=3&hoja=regional'
            '&regional=Zona Norte'
        )
    )['Por regional']

    assert hoja.max_row == 2
    assert hoja['A2'].value == 'Zona Norte'


def test_una_hoja_que_no_existe_se_ignora(app_bi_trade, marzo_con_datos):
    """Un `?hoja=` inventado no revienta: baja el mes completo."""
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    libro = _hoja_de(
        cliente_de(lector).get('/api/bi-trade/avance-mensual/exportar?anio=2026&mes=3&hoja=xyz')
    )
    assert len(libro.sheetnames) == 3


def test_exportar_el_tablero_pide_la_app(app_bi_trade):
    fuera = crear_usuario('fuera@supli.tech')
    assert cliente_de(fuera).get('/api/bi-trade/avance-mensual/exportar').status_code == 403


# ── Importación del informe del ERP ────────────────────────────────────────

#: Los encabezados del informe, con los espacios de sobra que trae el archivo.
ENCABEZADOS_INFORME = [
    ' N°SerieFab        ',
    'Material',
    'Ce. ',
    'CMv',
    'Fe/contab/',
    'Marca',
]


def _informe(filas, hoja='Base', encabezados=None):
    """
    Arma en memoria un informe como el que sale del ERP.

    Cada `fila` es `(material, centro, cmv, fecha)`. Se rellenan las columnas
    de relleno para que la hoja se parezca a la real: los índices se buscan
    por el nombre del encabezado, no por su posición.
    """
    libro = openpyxl.Workbook()
    pagina = libro.active
    pagina.title = hoja
    pagina.append(encabezados if encabezados is not None else ENCABEZADOS_INFORME)
    for material, centro, cmv, fecha in filas:
        pagina.append(['SERIE-1', material, centro, cmv, fecha, 'MARCA'])
    buffer = BytesIO()
    libro.save(buffer)
    return SimpleUploadedFile(
        'informe.xlsx',
        buffer.getvalue(),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def _importar(cliente, filas, anio=2026, mes=3, modo='completar', **extra):
    return cliente.post(
        '/api/bi-trade/importar-informe',
        {'archivo': _informe(filas), 'anio': anio, 'mes': mes, 'modo': modo, **extra},
        format='multipart',
    )


@pytest.fixture
def editor(app_bi_trade):
    return cliente_de(
        crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    )


def test_cada_fila_del_informe_es_una_unidad(editor, catalogo):
    """El ERP lista serie por serie, así que la cantidad sale de contar filas."""
    respuesta = _importar(
        editor,
        [
            ('SKU-1', 'PDV-001', '601', date(2026, 3, 4)),
            ('SKU-1', 'PDV-001', '601', date(2026, 3, 4)),
            ('SKU-1', 'PDV-001', '601', date(2026, 3, 4)),
            ('SKU-1', 'PDV-001', '601', date(2026, 3, 5)),
        ],
    )

    assert respuesta.status_code == 200
    # Tres filas del día 4 se agrupan en una venta de 3 unidades.
    assert Venta.objects.count() == 2
    dia4 = Venta.objects.get(fecha_venta=date(2026, 3, 4))
    assert dia4.cantidad_vendida == 3
    assert Venta.objects.get(fecha_venta=date(2026, 3, 5)).cantidad_vendida == 1


def test_solo_el_601_es_venta_y_el_601a_no(editor, catalogo):
    """`601A` también dice «Venta» en el informe, pero es de otro periodo."""
    respuesta = _importar(
        editor,
        [
            ('SKU-1', 'PDV-001', '601', date(2026, 3, 4)),
            ('SKU-1', 'PDV-001', '601A', date(2026, 3, 4)),
            ('SKU-1', 'PDV-001', '601A', date(2026, 3, 5)),
            ('SKU-1', 'PDV-001', '7', date(2026, 3, 6)),
        ],
    )

    datos = respuesta.json()
    assert datos['ventas']['creadas'] == 1
    assert Venta.objects.count() == 1
    assert Venta.objects.first().cantidad_vendida == 1


def test_el_cmv_1_es_inventario_y_reemplaza_lo_que_habia(editor, catalogo):
    norte, _sur, caro, barato = catalogo
    # Un inventario viejo que debe desaparecer.
    Inventario.objects.create(id_producto=barato, id_punto_venta=norte, cantidad_inventario=99)

    datos = _importar(
        editor,
        [
            ('SKU-1', 'PDV-001', '1', None),
            ('SKU-1', 'PDV-001', '1', None),
            ('SKU-2', 'PDV-002', '1', None),
        ],
    ).json()

    assert datos['inventario'] == {
        'creados': 2,
        'unidades': 3,
        'eliminados': 1,
        'filasLeidas': 3,
        'sinProducto': 0,
        'sinPuntoVenta': 0,
        'reemplazado': True,
    }
    assert Inventario.objects.count() == 2
    assert Inventario.objects.get(id_producto=caro, id_punto_venta=norte).cantidad_inventario == 2
    # El registro viejo del par (barato, norte) ya no está.
    assert not Inventario.objects.filter(id_producto=barato, id_punto_venta=norte).exists()


def test_completar_respeta_los_dias_que_ya_tienen_ventas(editor, catalogo):
    """El caso de todos los días: la base va al 17 y el archivo trae hasta el 23."""
    norte, _sur, caro, _barato = catalogo
    for dia in range(1, 18):
        Venta.objects.create(
            id_producto=caro,
            id_punto_venta=norte,
            fecha_venta=date(2026, 3, dia),
            cantidad_vendida=1,
        )

    filas = [('SKU-1', 'PDV-001', '601', date(2026, 3, dia)) for dia in range(1, 24)]
    datos = _importar(editor, filas, modo='completar').json()['ventas']

    assert datos['creadas'] == 6  # del 18 al 23
    assert datos['omitidasPorDia'] == 17
    assert datos['eliminadas'] == 0
    assert datos['diasCargados'] == [f'2026-03-{dia}' for dia in range(18, 24)]
    assert Venta.objects.count() == 17 + 6


def test_completar_llena_los_huecos_del_medio(editor, catalogo):
    """Si falta el 9 y del 24 al 31, se llenan sin tocar el resto."""
    norte, _sur, caro, _barato = catalogo
    cargados = [*range(1, 9), *range(10, 24)]
    for dia in cargados:
        Venta.objects.create(
            id_producto=caro,
            id_punto_venta=norte,
            fecha_venta=date(2026, 3, dia),
            cantidad_vendida=1,
        )

    filas = [('SKU-1', 'PDV-001', '601', date(2026, 3, dia)) for dia in range(1, 32)]
    datos = _importar(editor, filas, modo='completar').json()['ventas']

    faltantes = [9, *range(24, 32)]
    assert datos['creadas'] == len(faltantes)
    assert datos['diasCargados'] == [f'2026-03-{dia:02d}' for dia in faltantes]
    assert Venta.objects.count() == len(cargados) + len(faltantes)


def test_sobrescribir_borra_el_mes_y_deja_solo_el_archivo(editor, catalogo):
    norte, _sur, caro, _barato = catalogo
    for dia in range(1, 32):
        Venta.objects.create(
            id_producto=caro,
            id_punto_venta=norte,
            fecha_venta=date(2026, 3, dia),
            cantidad_vendida=1,
        )

    filas = [('SKU-1', 'PDV-001', '601', date(2026, 3, dia)) for dia in range(1, 16)]
    datos = _importar(editor, filas, modo='sobrescribir').json()['ventas']

    assert datos['eliminadas'] == 31
    assert datos['creadas'] == 15
    assert datos['omitidasPorDia'] == 0
    # El resto del mes queda en cero, aunque antes tuviera ventas.
    assert Venta.objects.count() == 15
    assert not Venta.objects.filter(fecha_venta=date(2026, 3, 20)).exists()


def test_sobrescribir_solo_toca_el_mes_elegido(editor, catalogo):
    norte, _sur, caro, _barato = catalogo
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 2, 10), cantidad_vendida=1
    )
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 4, 10), cantidad_vendida=1
    )

    _importar(
        editor, [('SKU-1', 'PDV-001', '601', date(2026, 3, 4))], mes=3, modo='sobrescribir'
    )

    assert Venta.objects.filter(fecha_venta=date(2026, 2, 10)).exists()
    assert Venta.objects.filter(fecha_venta=date(2026, 4, 10)).exists()
    assert Venta.objects.count() == 3


def test_las_ventas_de_otro_mes_no_entran(editor, catalogo):
    """Se importa el mes elegido, no lo que traiga el archivo."""
    datos = _importar(
        editor,
        [
            ('SKU-1', 'PDV-001', '601', date(2026, 3, 4)),
            ('SKU-1', 'PDV-001', '601', date(2026, 4, 4)),
            ('SKU-1', 'PDV-001', '601', date(2026, 2, 4)),
        ],
        mes=3,
    ).json()['ventas']

    assert datos['creadas'] == 1
    assert datos['fueraDelMes'] == 2
    assert Venta.objects.count() == 1


def test_no_se_carga_lo_que_no_existe_en_la_app(editor, catalogo):
    """El informe no crea catálogo: si el código no está, la fila se omite."""
    datos = _importar(
        editor,
        [
            ('SKU-1', 'PDV-001', '601', date(2026, 3, 4)),
            ('SKU-9', 'PDV-001', '601', date(2026, 3, 5)),  # producto que no existe
            ('SKU-1', 'C900', '601', date(2026, 3, 6)),  # punto de venta que no existe
            ('SKU-9', 'C900', '1', None),
            ('SKU-1', 'PDV-001', '1', None),
        ],
    ).json()

    assert datos['ventas']['creadas'] == 1
    assert datos['ventas']['sinProducto'] == 1
    assert datos['ventas']['sinPuntoVenta'] == 1
    assert datos['inventario']['creados'] == 1
    assert datos['inventario']['sinProducto'] == 1
    assert Venta.objects.count() == 1
    assert Inventario.objects.count() == 1


def test_los_codigos_numericos_cruzan_con_los_de_la_base(app_bi_trade, editor):
    """Excel devuelve los códigos como número; sin normalizar no cruzarían."""
    punto = PuntoVenta.objects.create(id_punto_venta='C230', nombre_pdv='Cav Centro')
    Producto.objects.create(
        id_producto='7023988',
        nombre_producto='Torre de sonido',
        marca='Aiwa',
        precio_venta_claro=1_603_900,
        precio_venta_coltrade=1_347_815,
    )

    # `Material` llega como entero y `CMv` también, tal como los da el ERP.
    datos = _importar(editor, [(7023988, 'C230', 601, date(2026, 3, 4))]).json()

    assert datos['ventas']['creadas'] == 1
    assert Venta.objects.get().id_punto_venta_id == punto.pk


def test_un_informe_sin_inventario_no_borra_el_que_hay(editor, catalogo):
    """Vaciar la tabla sin con qué reemplazarla sería destruir el dato."""
    norte, _sur, caro, _barato = catalogo
    Inventario.objects.create(id_producto=caro, id_punto_venta=norte, cantidad_inventario=7)

    datos = _importar(editor, [('SKU-1', 'PDV-001', '601', date(2026, 3, 4))]).json()

    assert datos['inventario']['reemplazado'] is False
    assert datos['inventario']['eliminados'] == 0
    assert Inventario.objects.get().cantidad_inventario == 7


def test_la_hoja_base_se_encuentra_sin_importar_mayusculas(editor, catalogo):
    cliente, archivo = editor, _informe(
        [('SKU-1', 'PDV-001', '601', date(2026, 3, 4))], hoja='BASE'
    )
    respuesta = cliente.post(
        '/api/bi-trade/importar-informe',
        {'archivo': archivo, 'anio': 2026, 'mes': 3, 'modo': 'completar'},
        format='multipart',
    )
    assert respuesta.status_code == 200


def test_sin_hoja_base_se_dice_que_hojas_hay(editor, catalogo):
    archivo = _informe([('SKU-1', 'PDV-001', '601', date(2026, 3, 4))], hoja='Resumen')
    respuesta = editor.post(
        '/api/bi-trade/importar-informe',
        {'archivo': archivo, 'anio': 2026, 'mes': 3, 'modo': 'completar'},
        format='multipart',
    )

    assert respuesta.status_code == 400
    assert 'Resumen' in respuesta.json()['message']


def test_si_falta_una_columna_se_dice_cual(editor, catalogo):
    archivo = _informe(
        [('SKU-1', 'PDV-001', '601', date(2026, 3, 4))],
        encabezados=['N°Serie', 'Material', 'Ce. ', 'CMv', 'Otra cosa', 'Marca'],
    )
    respuesta = editor.post(
        '/api/bi-trade/importar-informe',
        {'archivo': archivo, 'anio': 2026, 'mes': 3, 'modo': 'completar'},
        format='multipart',
    )

    assert respuesta.status_code == 400
    assert 'Fe/contab/' in respuesta.json()['message']


def test_un_informe_sin_ventas_ni_inventario_se_rechaza(editor, catalogo):
    respuesta = _importar(editor, [('SKU-1', 'PDV-001', '601A', date(2026, 3, 4))])

    assert respuesta.status_code == 400
    assert 'CMv' in respuesta.json()['message']


def test_el_mes_es_obligatorio_y_valido(editor, catalogo):
    filas = [('SKU-1', 'PDV-001', '601', date(2026, 3, 4))]
    for anio, mes in ((2026, 13), ('abc', 3), (1800, 3)):
        respuesta = _importar(editor, filas, anio=anio, mes=mes)
        assert respuesta.status_code == 400
        assert 'mes' in respuesta.json()['message']


def test_un_modo_desconocido_se_rechaza(editor, catalogo):
    respuesta = _importar(
        editor, [('SKU-1', 'PDV-001', '601', date(2026, 3, 4))], modo='mezclar'
    )
    assert respuesta.status_code == 400


def test_importar_el_informe_exige_permiso_de_edicion(app_bi_trade, catalogo):
    lector = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    respuesta = _importar(lector, [('SKU-1', 'PDV-001', '601', date(2026, 3, 4))])
    assert respuesta.status_code == 403


def test_importar_el_informe_rechaza_un_archivo_que_no_es_xlsx(editor, catalogo):
    respuesta = editor.post(
        '/api/bi-trade/importar-informe',
        {
            'archivo': SimpleUploadedFile('informe.csv', b'Material,Ce.', 'text/csv'),
            'anio': 2026,
            'mes': 3,
            'modo': 'completar',
        },
        format='multipart',
    )
    assert respuesta.status_code == 400


def test_importar_dos_veces_en_completar_no_duplica(editor, catalogo):
    filas = [('SKU-1', 'PDV-001', '601', date(2026, 3, 4))]
    _importar(editor, filas)
    datos = _importar(editor, filas).json()['ventas']

    assert datos['creadas'] == 0
    assert Venta.objects.count() == 1


# ── Concurso de tickets ────────────────────────────────────────────────────

@pytest.fixture
def campana(catalogo):
    """La campaña del afiche: escalas 10/12/15 y acelerador 150/200."""
    _norte, _sur, caro, barato = catalogo
    c = Campana.objects.create(
        nombre='Vamos por todo',
        desde=date(2026, 3, 1),
        hasta=date(2026, 3, 31),
        ventas_minimas=120,
        tickets_minimos=3,
        foco_minimo=5,
        bono_ventas=20,
        bono_cargadores=5,
        bono_tickets=3,
    )
    EscalaTicket.objects.bulk_create(
        EscalaTicket(campana=c, ventas=v, tickets=t) for v, t in ((10, 1), (12, 2), (15, 3))
    )
    Acelerador.objects.bulk_create(
        Acelerador(campana=c, ventas_totales=v, tickets_por_dia=t) for v, t in ((150, 1), (200, 2))
    )
    # `caro` es el producto foco y `barato` el cargador, para poder probar
    # las dos reglas con el catálogo que ya existe.
    c.productos_foco.set([caro])
    c.productos_cargador.set([barato])
    return c


def _vender(punto, producto, dia, unidades, mes=3):
    return Venta.objects.create(
        id_producto=producto,
        id_punto_venta=punto,
        fecha_venta=date(2026, mes, dia),
        cantidad_vendida=unidades,
    )


def _fila(campana_obj, punto):
    filas = {f.key: f for f in calcular_tickets(campana_obj)}
    return filas.get(punto.pk)


def test_la_escala_aplica_el_tramo_mas_alto_que_alcanza(campana, catalogo):
    norte, _sur, _caro, barato = catalogo
    # 13 unidades caen en la escala de 12, no en la de 15.
    _vender(norte, barato, 4, 13)

    fila = _fila(campana, norte)
    assert fila.dias[0].unidades == 13
    assert fila.dias[0].tickets_escala == 2
    assert fila.tickets == 2


def test_por_debajo_del_primer_tramo_no_hay_tickets(campana, catalogo):
    norte, _sur, _caro, barato = catalogo
    _vender(norte, barato, 4, 9)

    fila = _fila(campana, norte)
    assert fila.tickets == 0
    assert fila.dias_cumplidos == 0
    assert not fila.dias[0].cumplido


def test_el_producto_foco_duplica_los_tickets_del_dia(campana, catalogo):
    norte, _sur, caro, barato = catalogo
    # 10 unidades en total, de las cuales 5 son foco: escala de 10 = 1 ticket,
    # duplicado a 2.
    _vender(norte, caro, 4, 5)
    _vender(norte, barato, 4, 5)

    fila = _fila(campana, norte)
    assert fila.dias[0].unidades == 10
    assert fila.dias[0].foco == 5
    assert fila.dias[0].duplicado
    assert fila.tickets == 2


def test_sin_el_minimo_de_foco_no_se_duplica(campana, catalogo):
    norte, _sur, caro, barato = catalogo
    _vender(norte, caro, 4, 4)  # uno menos del mínimo
    _vender(norte, barato, 4, 6)

    fila = _fila(campana, norte)
    assert fila.dias[0].foco == 4
    assert not fila.dias[0].duplicado
    assert fila.tickets == 1


def test_el_doble_no_crea_tickets_donde_la_escala_no_dio_ninguno(campana, catalogo):
    """Duplicar cero es cero: el foco premia un buen día, no lo inventa."""
    norte, _sur, caro, _barato = catalogo
    _vender(norte, caro, 4, 6)  # 6 unidades: no alcanza la escala de 10

    fila = _fila(campana, norte)
    assert fila.dias[0].foco == 6
    assert fila.tickets == 0
    assert not fila.dias[0].duplicado


def test_el_bono_se_suma_a_la_escala_no_la_reemplaza(campana, catalogo):
    norte, _sur, caro, barato = catalogo
    # 20 unidades con 5 cargadores: escala de 15 = 3 tickets, más 3 del bono.
    # Los 15 del foco además duplican la escala.
    _vender(norte, caro, 4, 15)
    _vender(norte, barato, 4, 5)

    dia = _fila(campana, norte).dias[0]
    assert dia.unidades == 20
    assert dia.cargador == 5
    assert dia.tickets_escala == 6  # 3 duplicados
    assert dia.tickets_bono == 3
    assert dia.tickets == 9


def test_sin_los_cargadores_no_hay_bono(campana, catalogo):
    norte, _sur, caro, barato = catalogo
    _vender(norte, caro, 4, 16)
    _vender(norte, barato, 4, 4)  # uno menos del mínimo de cargadores

    dia = _fila(campana, norte).dias[0]
    assert dia.unidades == 20
    assert dia.tickets_bono == 0


def test_el_acelerador_suma_por_cada_dia_cumplido(campana, catalogo):
    """No premia el día sino el acumulado, y se resuelve al final."""
    norte, _sur, _caro, barato = catalogo
    # 15 días de 10 unidades: 150 en total, 15 días cumplidos de 1 ticket.
    for dia in range(1, 16):
        _vender(norte, barato, dia, 10)

    fila = _fila(campana, norte)
    assert fila.unidades == 150
    assert fila.dias_cumplidos == 15
    assert fila.tickets_escala == 15
    assert fila.acelerador_alcanzado == 150
    assert fila.tickets_acelerador == 15  # +1 por cada uno de los 15 días
    assert fila.tickets == 30


def test_el_acelerador_toma_el_umbral_mas_alto(campana, catalogo):
    norte, _sur, _caro, barato = catalogo
    for dia in range(1, 21):
        _vender(norte, barato, dia, 10)  # 200 unidades en 20 días

    fila = _fila(campana, norte)
    assert fila.unidades == 200
    assert fila.acelerador_alcanzado == 200
    assert fila.tickets_acelerador == 40  # +2 por cada uno de los 20 días


def test_un_dia_sin_tickets_no_cuenta_para_el_acelerador(campana, catalogo):
    norte, _sur, _caro, barato = catalogo
    for dia in range(1, 16):
        _vender(norte, barato, dia, 10)
    # Un día flojo: suma ventas pero no gana tickets ni acelera.
    _vender(norte, barato, 20, 5)

    fila = _fila(campana, norte)
    assert fila.unidades == 155
    assert fila.dias_cumplidos == 15
    assert fila.tickets_acelerador == 15


def test_solo_cuentan_las_ventas_dentro_de_la_vigencia(campana, catalogo):
    norte, _sur, _caro, barato = catalogo
    _vender(norte, barato, 4, 15)  # marzo, dentro
    _vender(norte, barato, 4, 15, mes=2)  # febrero, fuera
    _vender(norte, barato, 4, 15, mes=4)  # abril, fuera

    fila = _fila(campana, norte)
    assert fila.unidades == 15
    assert len(fila.dias) == 1


def test_participa_quien_cumple_las_dos_condiciones(campana, catalogo):
    norte, sur, _caro, barato = catalogo
    # El norte llega a 120 unidades y a 12 tickets.
    for dia in range(1, 13):
        _vender(norte, barato, dia, 10)
    # El sur tiene tickets pero no llega a las 120 ventas.
    for dia in range(1, 4):
        _vender(sur, barato, dia, 10)

    filas = {f.key: f for f in calcular_tickets(campana)}
    assert participa(campana, filas[norte.pk])
    assert not participa(campana, filas[sur.pk])
    assert filas[sur.pk].tickets == 3  # tiene los tickets, le faltan las ventas


def test_apagar_el_doble_y_el_bono_con_cero(campana, catalogo):
    """`foco_minimo` y `bono_tickets` en cero desactivan esas reglas."""
    norte, _sur, caro, barato = catalogo
    campana.foco_minimo = 0
    campana.bono_tickets = 0
    campana.save()
    _vender(norte, caro, 4, 15)
    _vender(norte, barato, 4, 5)

    dia = _fila(campana, norte).dias[0]
    assert dia.tickets_escala == 3  # sin duplicar
    assert dia.tickets_bono == 0
    assert not dia.duplicado


def test_sin_escalas_no_hay_tickets_de_escala_pero_el_bono_sigue(campana, catalogo):
    """El bono es independiente de la escala: en el afiche dice «adicional».

    Un día así gana los tickets del bono pero no cuenta como día cumplido, así
    que tampoco alimenta el acelerador: el acelerador premia cumplir la escala.
    """
    norte, _sur, _caro, barato = catalogo
    campana.escalas.all().delete()
    _vender(norte, barato, 4, 50)  # 50 unidades, todas cargadores

    fila = _fila(campana, norte)
    assert fila.unidades == 50
    assert fila.tickets_escala == 0
    assert fila.tickets_bono == 3
    assert fila.dias_cumplidos == 0
    assert fila.tickets_acelerador == 0


def test_sin_escalas_ni_bono_no_hay_tickets(campana, catalogo):
    norte, _sur, _caro, barato = catalogo
    campana.escalas.all().delete()
    campana.bono_tickets = 0
    campana.save()
    _vender(norte, barato, 4, 50)

    assert _fila(campana, norte).tickets == 0


def test_el_tablero_de_tickets_ordena_por_tickets(app_bi_trade, campana, catalogo):
    norte, sur, _caro, barato = catalogo
    for dia in range(1, 6):
        _vender(sur, barato, dia, 10)
    _vender(norte, barato, 1, 10)

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/tickets').json()

    assert [f['label'] for f in datos['filas']] == ['Claro Jardín Plaza', 'Claro Unicentro']
    assert datos['filas'][0]['tickets'] == 5
    assert datos['totales']['tickets'] == 6
    assert datos['totales']['puntosVenta'] == 2


def test_el_tablero_toma_la_campana_activa_mas_reciente(app_bi_trade, campana, catalogo):
    vieja = Campana.objects.create(
        nombre='Campaña vieja', desde=date(2025, 1, 1), hasta=date(2025, 1, 31)
    )
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    assert cliente.get('/api/bi-trade/tickets').json()['campana']['nombre'] == 'Vamos por todo'
    # Y con `?campana=` se puede mirar otra.
    pedida = cliente.get(f'/api/bi-trade/tickets?campana={vieja.pk}').json()
    assert pedida['campana']['nombre'] == 'Campaña vieja'


def test_una_campana_inactiva_no_se_elige_sola(app_bi_trade, campana, catalogo):
    campana.activa = False
    campana.save()
    lector = crear_usuario('lector@supli.tech', app_bi_trade)

    assert cliente_de(lector).get('/api/bi-trade/tickets').json()['campana'] is None


def test_sin_campanas_el_tablero_lo_dice_en_vez_de_reventar(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/tickets').json()

    assert datos['campana'] is None
    assert datos['filas'] == []
    assert 'campañas' in datos['message']


def test_el_concurso_se_filtra_por_regional(app_bi_trade, campana, catalogo):
    norte, sur, _caro, barato = catalogo
    _vender(norte, barato, 1, 15)
    _vender(sur, barato, 1, 15)

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/tickets?regional=Zona Norte').json()

    assert [f['label'] for f in datos['filas']] == ['Claro Unicentro']


def test_crear_una_campana_guarda_sus_escalas_de_una_vez(app_bi_trade, catalogo):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    respuesta = cliente_de(editor).post(
        '/api/bi-trade/campanas',
        {
            'nombre': 'Nueva',
            'desde': '2026-03-01',
            'hasta': '2026-03-31',
            'escalas': [{'ventas': 10, 'tickets': 1}, {'ventas': 20, 'tickets': 4}],
            'aceleradores': [{'ventasTotales': 100, 'ticketsPorDia': 1}],
            'productosFoco': ['SKU-1'],
        },
        format='json',
    )

    assert respuesta.status_code == 201
    creada = Campana.objects.get(nombre='Nueva')
    assert list(creada.escalas.values_list('ventas', 'tickets')) == [(10, 1), (20, 4)]
    assert list(creada.aceleradores.values_list('ventas_totales', 'tickets_por_dia')) == [(100, 1)]
    assert list(creada.productos_foco.values_list('pk', flat=True)) == ['SKU-1']
    assert respuesta.json()['dias'] == 31


def test_editar_una_campana_reemplaza_sus_escalas(app_bi_trade, campana):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    respuesta = cliente_de(editor).patch(
        f'/api/bi-trade/campanas/{campana.pk}',
        {'escalas': [{'ventas': 5, 'tickets': 1}]},
        format='json',
    )

    assert respuesta.status_code == 200
    assert list(campana.escalas.values_list('ventas', 'tickets')) == [(5, 1)]


def test_una_campana_no_puede_terminar_antes_de_empezar(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    respuesta = cliente_de(editor).post(
        '/api/bi-trade/campanas',
        {'nombre': 'Al revés', 'desde': '2026-03-31', 'hasta': '2026-03-01'},
        format='json',
    )
    assert respuesta.status_code == 400


def test_crear_una_campana_pide_permiso(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).post(
        '/api/bi-trade/campanas',
        {'nombre': 'No', 'desde': '2026-03-01', 'hasta': '2026-03-31'},
        format='json',
    )
    assert respuesta.status_code == 403


def test_el_concurso_se_exporta_en_dos_hojas(app_bi_trade, campana, catalogo):
    norte, _sur, _caro, barato = catalogo
    _vender(norte, barato, 4, 15)

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).get('/api/bi-trade/tickets/exportar')

    assert respuesta.status_code == 200
    libro = _hoja_de(respuesta)
    assert libro.sheetnames == ['Tickets por punto', 'Día por día']
    assert libro['Tickets por punto'].max_row == 2
    encabezados = [celda.value for celda in libro['Tickets por punto'][1]]
    assert 'Tickets totales' in encabezados
    assert '¿Participa?' in encabezados


def test_exportar_sin_campanas_avisa(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    assert cliente_de(lector).get('/api/bi-trade/tickets/exportar').status_code == 400


def test_el_concurso_pide_la_app(app_bi_trade):
    fuera = crear_usuario('fuera@supli.tech')
    assert cliente_de(fuera).get('/api/bi-trade/tickets').status_code == 403


# ── Cumplimiento diario ────────────────────────────────────────────────────

@pytest.fixture
def marzo_para_el_dia(catalogo):
    """Meta de marzo y una venta el día 4, para medir ese día.

    25 unidades del caro = 100.000.000 de meta al mes, que entre los 25 días
    hábiles de marzo de 2026 da 4.000.000 diarios.
    """
    norte, sur, caro, barato = catalogo
    MetaComercial.objects.create(
        id_producto=caro,
        id_punto_venta=norte,
        fecha_meta=date(2026, 3, 1),
        meta_cantidad=25,
    )
    MetaComercial.objects.create(
        id_producto=barato,
        id_punto_venta=sur,
        fecha_meta=date(2026, 3, 1),
        meta_cantidad=50,
    )
    return catalogo


def _dia(cliente, consulta=''):
    ruta = '/api/bi-trade/cumplimiento-diario'
    return cliente.get(f'{ruta}?{consulta}' if consulta else ruta).json()


def test_la_cuota_del_dia_es_la_meta_del_mes_entre_los_habiles(
    app_bi_trade, marzo_para_el_dia
):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = _dia(cliente_de(lector), 'fecha=2026-03-04')

    assert datos['dia']['diasHabiles'] == 25
    # 25 × 4.000.000 + 50 × 800.000 = 140.000.000 de meta al mes.
    assert datos['totales']['metaMensualDinero'] == 140_000_000
    assert datos['totales']['metaDinero'] == 140_000_000 // 25
    # Las unidades y los puntos se reparten igual.
    assert datos['totales']['metaCantidad'] == round(75 / 25)
    assert datos['totales']['metaPuntos'] == round(25 * 99 / 25)


def test_el_dia_solo_cuenta_las_ventas_de_ese_dia(app_bi_trade, marzo_para_el_dia):
    norte, _sur, caro, _barato = marzo_para_el_dia
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=1
    )
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 5), cantidad_vendida=9
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    dia4 = _dia(cliente, 'fecha=2026-03-04')['totales']
    assert dia4['realCantidad'] == 1
    assert dia4['realDinero'] == 4_000_000
    assert dia4['operaciones'] == 1

    dia5 = _dia(cliente, 'fecha=2026-03-05')['totales']
    assert dia5['realCantidad'] == 9


def test_el_cumplimiento_del_dia_compara_contra_la_cuota_del_dia(
    app_bi_trade, marzo_para_el_dia
):
    norte, _sur, caro, _barato = marzo_para_el_dia
    # La cuota total del día es 5.600.000; se venden 4.000.000.
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=1
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = _dia(cliente_de(lector), 'fecha=2026-03-04')['totales']

    assert datos['metaDinero'] == 5_600_000
    assert datos['cumplimientoDinero'] == round(4_000_000 * 100 / 5_600_000, 1)


def test_las_tres_medidas_tienen_su_propia_meta(app_bi_trade, marzo_para_el_dia):
    """Dinero, unidades y puntos no se miden con el mismo objetivo."""
    norte, _sur, caro, _barato = marzo_para_el_dia
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=2
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = _dia(cliente_de(lector), 'fecha=2026-03-04')['totales']

    assert datos['realDinero'] == 8_000_000
    assert datos['realCantidad'] == 2
    assert datos['realPuntos'] == 2 * 99
    # Cada porcentaje sale de su propia meta, así que son distintos.
    assert datos['cumplimientoDinero'] != datos['cumplimientoCantidad']


def test_el_dia_se_desglosa_por_las_cuatro_dimensiones(app_bi_trade, marzo_para_el_dia):
    norte, sur, caro, barato = marzo_para_el_dia
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=1
    )
    Venta.objects.create(
        id_producto=barato, id_punto_venta=sur, fecha_venta=date(2026, 3, 4), cantidad_vendida=3
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = _dia(cliente_de(lector), 'fecha=2026-03-04')

    por_regional = {f['label']: f for f in datos['porRegional']}
    assert por_regional['Zona Norte']['realDinero'] == 4_000_000
    assert por_regional['Zona Sur']['realDinero'] == 2_400_000
    # La cuota del norte: 100.000.000 del mes entre 25 días.
    assert por_regional['Zona Norte']['metaDinero'] == 4_000_000

    assert {f['label'] for f in datos['porMarca']} == {'Apple', 'Xiaomi'}
    assert {f['label'] for f in datos['porProducto']} == {'iPhone 15', 'Redmi Note 13'}
    assert {f['label'] for f in datos['porPuntoVenta']} == {
        'Claro Unicentro',
        'Claro Jardín Plaza',
    }


def test_el_dia_se_filtra_por_cada_campo(app_bi_trade, marzo_para_el_dia):
    norte, sur, caro, barato = marzo_para_el_dia
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=1
    )
    Venta.objects.create(
        id_producto=barato, id_punto_venta=sur, fecha_venta=date(2026, 3, 4), cantidad_vendida=3
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    cliente = cliente_de(lector)

    for consulta, esperado in (
        ('regional=Zona Norte', 4_000_000),
        ('marca=Xiaomi', 2_400_000),
        ('id_producto=SKU-1', 4_000_000),
        ('id_punto_venta=PDV-002', 2_400_000),
    ):
        datos = _dia(cliente, f'fecha=2026-03-04&{consulta}')['totales']
        assert datos['realDinero'] == esperado, consulta

    # El filtro también recorta la meta, no solo lo vendido.
    solo_norte = _dia(cliente, 'fecha=2026-03-04&regional=Zona Norte')['totales']
    assert solo_norte['metaMensualDinero'] == 100_000_000
    assert solo_norte['metaDinero'] == 4_000_000


def test_un_domingo_se_marca_como_no_habil(app_bi_trade, marzo_para_el_dia):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    # 1 de marzo de 2026 es domingo.
    datos = _dia(cliente_de(lector), 'fecha=2026-03-01')['dia']

    assert datos['habil'] is False
    assert datos['esDomingo'] is True
    assert datos['esFestivo'] is False
    assert datos['nombre'] == 'domingo 1 de marzo de 2026'


def test_un_festivo_se_marca_como_no_habil(app_bi_trade, marzo_para_el_dia):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    # San José, corrido al lunes 23 de marzo de 2026.
    datos = _dia(cliente_de(lector), 'fecha=2026-03-23')['dia']

    assert datos['habil'] is False
    assert datos['esFestivo'] is True
    assert datos['esDomingo'] is False


def test_el_sabado_si_es_habil(app_bi_trade, marzo_para_el_dia):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = _dia(cliente_de(lector), 'fecha=2026-03-21')['dia']

    assert datos['habil'] is True
    assert datos['nombre'].startswith('sábado')


def test_el_dia_trae_los_no_habiles_del_mes_para_el_selector(app_bi_trade, marzo_para_el_dia):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = _dia(cliente_de(lector), 'fecha=2026-03-04')['dia']

    # 5 domingos y el festivo del 23: 31 − 6 = 25 hábiles.
    assert datos['diasNoHabiles'] == [1, 8, 15, 22, 23, 29]
    assert datos['diasDelMes'] == 31
    assert len(datos['diasNoHabiles']) == 31 - 25


def test_un_dia_sin_ventas_da_cero_y_no_revienta(app_bi_trade, marzo_para_el_dia):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = _dia(cliente_de(lector), 'fecha=2026-03-04')['totales']

    assert datos['realDinero'] == 0
    assert datos['cumplimientoDinero'] == 0.0
    # La cuota sigue estando: es lo que no se cumplió.
    assert datos['metaDinero'] > 0


def test_un_mes_sin_metas_no_divide_por_cero(app_bi_trade, catalogo):
    norte, _sur, caro, _barato = catalogo
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 5, 4), cantidad_vendida=2
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = _dia(cliente_de(lector), 'fecha=2026-05-04')['totales']

    assert datos['metaDinero'] == 0
    assert datos['cumplimientoDinero'] == 0.0
    assert datos['realDinero'] == 8_000_000


def test_una_fecha_ilegible_cae_en_hoy(app_bi_trade, marzo_para_el_dia):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    hoy = timezone.localdate()

    for consulta in ('fecha=no-es-fecha', 'fecha=2026-13-40', ''):
        datos = _dia(cliente_de(lector), consulta)['dia']
        assert datos['fecha'] == hoy.isoformat(), consulta


def test_el_dia_se_exporta_en_cuatro_hojas(app_bi_trade, marzo_para_el_dia):
    norte, _sur, caro, _barato = marzo_para_el_dia
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=1
    )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).get(
        '/api/bi-trade/cumplimiento-diario/exportar?fecha=2026-03-04'
    )

    assert respuesta.status_code == 200
    assert 'cumplimiento-2026-03-04.xlsx' in respuesta['Content-Disposition']
    libro = _hoja_de(respuesta)
    assert libro.sheetnames == [
        'Por regional',
        'Por punto de venta',
        'Por marca',
        'Por producto',
    ]
    encabezados = [celda.value for celda in libro['Por regional'][1]]
    assert encabezados[0] == 'Regional'
    assert 'Meta del día ($)' in encabezados
    assert 'Cumplimiento pts (%)' in encabezados


def test_un_corte_solo_baja_sin_los_otros(app_bi_trade, marzo_para_el_dia):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).get(
        '/api/bi-trade/cumplimiento-diario/exportar?fecha=2026-03-04&corte=marcas'
    )

    libro = _hoja_de(respuesta)
    assert libro.sheetnames == ['Por marca']
    assert 'cumplimiento-2026-03-04-marcas.xlsx' in respuesta['Content-Disposition']


def test_exportar_el_dia_respeta_los_filtros(app_bi_trade, marzo_para_el_dia):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    hoja = _hoja_de(
        cliente_de(lector).get(
            '/api/bi-trade/cumplimiento-diario/exportar'
            '?fecha=2026-03-04&corte=regional&regional=Zona Norte'
        )
    )['Por regional']

    assert hoja.max_row == 2
    assert hoja['A2'].value == 'Zona Norte'


def test_el_cumplimiento_diario_pide_la_app(app_bi_trade):
    fuera = crear_usuario('fuera@supli.tech')
    assert cliente_de(fuera).get('/api/bi-trade/cumplimiento-diario').status_code == 403
    assert (
        cliente_de(fuera).get('/api/bi-trade/cumplimiento-diario/exportar').status_code == 403
    )


# ── Enlace público del tablero ─────────────────────────────────────────────

RUTA_PUBLICA = '/api/publico/bi-trade'


@pytest.fixture
def enlace(app_bi_trade):
    """Un enlace recién creado: `(token, clave, id)`.

    Limpia la caché antes: el límite de intentos vive ahí, y sin limpiarla una
    prueba que agota los intentos dejaría bloqueadas a las siguientes.
    """
    cache.clear()
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    datos = cliente_de(editor).post(
        '/api/bi-trade/enlaces', {'nombre': 'Gerencia Claro'}, format='json'
    ).json()
    return datos['token'], datos['clave'], datos['idEnlace']


def _entrar(token, clave):
    return APIClient().post(f'{RUTA_PUBLICA}/{token}/acceso', {'clave': clave}, format='json')


def _con_acceso(token, clave):
    """Un cliente anónimo con un acceso vigente para ese enlace."""
    acceso = _entrar(token, clave).json()['acceso']
    cliente = APIClient()
    cliente.credentials(HTTP_X_ACCESO_PUBLICO=acceso)
    return cliente


def test_crear_un_enlace_exige_el_permiso_de_edicion(app_bi_trade):
    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    respuesta = cliente_de(lector).post(
        '/api/bi-trade/enlaces', {'nombre': 'No'}, format='json'
    )
    assert respuesta.status_code == 403
    assert cliente_de(lector).get('/api/bi-trade/enlaces').status_code == 403


def test_la_clave_se_muestra_una_vez_y_se_guarda_como_hash(app_bi_trade, enlace):
    token, clave, id_enlace = enlace
    guardado = EnlacePublico.objects.get(pk=id_enlace)

    # Nunca en texto plano: ni con guiones ni sin ellos.
    assert clave not in guardado.clave_hash
    assert clave.replace('-', '') not in guardado.clave_hash
    assert check_password(clave.replace('-', ''), guardado.clave_hash)

    editor = crear_usuario('otro@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    listado = cliente_de(editor).get('/api/bi-trade/enlaces').json()
    assert 'clave' not in listado[0]
    assert 'claveHash' not in listado[0]


def test_la_clave_tiene_tres_bloques_sin_caracteres_ambiguos(enlace):
    _token, clave, _id = enlace
    bloques = clave.split('-')
    assert [len(bloque) for bloque in bloques] == [4, 4, 4]
    assert not set(''.join(bloques)) & set('0O1lI')


def test_dos_enlaces_no_comparten_token_ni_clave(app_bi_trade, enlace):
    editor = crear_usuario('otro@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    otro = cliente_de(editor).post(
        '/api/bi-trade/enlaces', {'nombre': 'Otro'}, format='json'
    ).json()
    assert otro['token'] != enlace[0]
    assert otro['clave'] != enlace[1]


def test_sin_acceso_no_se_ven_los_datos(enlace):
    token, _clave, _id = enlace
    anonimo = APIClient()
    for tramo in ('avance-mensual', 'cumplimiento-diario', 'tickets', 'opciones', 'productos'):
        assert anonimo.get(f'{RUTA_PUBLICA}/{token}/{tramo}').status_code == 403, tramo


def test_la_puerta_solo_dice_el_nombre_y_el_canal(enlace):
    token, _clave, _id = enlace
    respuesta = APIClient().get(f'{RUTA_PUBLICA}/{token}')
    assert respuesta.status_code == 200
    assert respuesta.json() == {'nombre': 'Gerencia Claro', 'canal': 'claro'}


def test_un_token_inventado_responde_404(enlace):
    assert APIClient().get(f'{RUTA_PUBLICA}/no-existe').status_code == 404
    assert _entrar('no-existe', 'lo-que-sea').status_code == 404


def test_una_clave_mala_no_abre(enlace):
    token, _clave, _id = enlace
    respuesta = _entrar(token, 'AAAA-BBBB-CCCC')
    assert respuesta.status_code == 403
    assert 'acceso' not in respuesta.json()
    assert _entrar(token, '').status_code == 403


def test_la_clave_se_acepta_sin_guiones_ni_espacios(enlace):
    token, clave, _id = enlace
    assert _entrar(token, clave).status_code == 200
    assert _entrar(token, clave.replace('-', '')).status_code == 200
    assert _entrar(token, f'  {clave.replace("-", " ")}  ').status_code == 200


def test_la_clave_distingue_mayusculas(enlace):
    token, clave, _id = enlace
    cambiada = clave.swapcase()
    if cambiada != clave:
        assert _entrar(token, cambiada).status_code == 403


def test_con_acceso_se_ven_las_tres_hojas(enlace):
    token, clave, _id = enlace
    cliente = _con_acceso(token, clave)
    for tramo in (
        'avance-mensual?anio=2026&mes=3',
        'cumplimiento-diario?fecha=2026-03-04',
        'tickets',
        'opciones',
        'productos',
        'puntos-venta',
        'campanas',
    ):
        assert cliente.get(f'{RUTA_PUBLICA}/{token}/{tramo}').status_code == 200, tramo


def test_los_datos_publicos_son_los_mismos_de_la_app(app_bi_trade, enlace, catalogo):
    norte, _sur, caro, _barato = catalogo
    MetaComercial.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_meta=date(2026, 3, 1), meta_cantidad=25
    )
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=2
    )
    token, clave, _id = enlace
    lector = crear_usuario('lector@supli.tech', app_bi_trade)

    consulta = 'avance-mensual?anio=2026&mes=3'
    de_la_app = cliente_de(lector).get(f'/api/bi-trade/{consulta}').json()
    publico_ = _con_acceso(token, clave).get(f'{RUTA_PUBLICA}/{token}/{consulta}').json()
    assert publico_ == de_la_app


def test_los_catalogos_publicos_no_llevan_precios(enlace, catalogo):
    token, clave, _id = enlace
    productos = _con_acceso(token, clave).get(f'{RUTA_PUBLICA}/{token}/productos').json()
    assert set(productos[0]) == {'idProducto', 'nombreProducto', 'marca'}


def test_el_acceso_de_un_enlace_no_abre_otro(app_bi_trade, enlace):
    token, clave, _id = enlace
    editor = crear_usuario('otro@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    otro = cliente_de(editor).post(
        '/api/bi-trade/enlaces', {'nombre': 'Otro'}, format='json'
    ).json()

    cliente = _con_acceso(token, clave)
    assert cliente.get(f'{RUTA_PUBLICA}/{otro["token"]}/opciones').status_code == 403


def test_un_acceso_falsificado_no_abre(enlace):
    token, _clave, _id = enlace
    cliente = APIClient()
    cliente.credentials(HTTP_X_ACCESO_PUBLICO='eyJlIjoxLCJ2IjoxfQ:falso:firma')
    assert cliente.get(f'{RUTA_PUBLICA}/{token}/opciones').status_code == 403


def test_un_acceso_vencido_vuelve_a_pedir_la_clave(enlace, monkeypatch):
    token, clave, _id = enlace
    cliente = _con_acceso(token, clave)
    # Con duración negativa, cualquier firma ya está vencida.
    monkeypatch.setattr(publico, 'DURACION_ACCESO', -1)
    assert cliente.get(f'{RUTA_PUBLICA}/{token}/opciones').status_code == 403


def test_regenerar_la_clave_invalida_los_accesos_abiertos(app_bi_trade, enlace):
    token, clave_vieja, id_enlace = enlace
    cliente = _con_acceso(token, clave_vieja)
    editor = crear_usuario('otro@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    nueva = cliente_de(editor).post(f'/api/bi-trade/enlaces/{id_enlace}/regenerar-clave').json()

    assert nueva['token'] == token  # la URL no cambia
    assert nueva['clave'] != clave_vieja
    assert cliente.get(f'{RUTA_PUBLICA}/{token}/opciones').status_code == 403
    assert _entrar(token, clave_vieja).status_code == 403
    assert _entrar(token, nueva['clave']).status_code == 200


def test_un_enlace_revocado_no_abre_aunque_el_acceso_siga_firmado(app_bi_trade, enlace):
    token, clave, id_enlace = enlace
    cliente = _con_acceso(token, clave)
    editor = crear_usuario('otro@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    cliente_de(editor).patch(f'/api/bi-trade/enlaces/{id_enlace}', {'activo': False}, format='json')

    assert cliente.get(f'{RUTA_PUBLICA}/{token}/opciones').status_code == 404
    assert APIClient().get(f'{RUTA_PUBLICA}/{token}').status_code == 404
    assert _entrar(token, clave).status_code == 404

    # Reactivarlo lo devuelve, con la misma clave.
    cliente_de(editor).patch(f'/api/bi-trade/enlaces/{id_enlace}', {'activo': True}, format='json')
    assert _entrar(token, clave).status_code == 200


def test_un_enlace_vencido_no_abre(enlace):
    token, clave, id_enlace = enlace
    EnlacePublico.objects.filter(pk=id_enlace).update(
        expira=timezone.now() - timedelta(minutes=1)
    )
    assert _entrar(token, clave).status_code == 404
    assert APIClient().get(f'{RUTA_PUBLICA}/{token}').status_code == 404


def test_no_se_crea_un_enlace_ya_vencido(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    respuesta = cliente_de(editor).post(
        '/api/bi-trade/enlaces',
        {'nombre': 'Viejo', 'expira': (timezone.now() - timedelta(days=1)).isoformat()},
        format='json',
    )
    assert respuesta.status_code == 400


def test_cada_ingreso_se_cuenta(enlace):
    token, clave, id_enlace = enlace
    _entrar(token, clave)
    _entrar(token, clave)
    _entrar(token, 'mala-mala-mala')  # un intento fallido no cuenta

    guardado = EnlacePublico.objects.get(pk=id_enlace)
    assert guardado.accesos == 2
    assert guardado.ultimo_acceso is not None


def test_la_clave_tiene_limite_de_intentos(enlace):
    """Contra la fuerza bruta: al pasar el límite, ni la clave buena entra."""
    token, clave, _id = enlace
    for _ in range(10):
        assert _entrar(token, 'AAAA-BBBB-CCCC').status_code == 403
    assert _entrar(token, clave).status_code == 429


def test_el_enlace_publico_solo_lee(enlace):
    """Nada de escribir ni exportar: solo las consultas de las tres hojas."""
    token, clave, _id = enlace
    cliente = _con_acceso(token, clave)

    assert cliente.post(f'{RUTA_PUBLICA}/{token}/avance-mensual', {}).status_code == 405
    for tramo in (
        'avance-mensual/exportar',
        'cumplimiento-diario/exportar',
        'tickets/exportar',
        'ventas',
        'importar-informe',
        'enlaces',
    ):
        assert cliente.get(f'{RUTA_PUBLICA}/{token}/{tramo}').status_code == 404, tramo


def test_el_acceso_publico_no_sirve_en_la_api_de_la_app(enlace):
    """El acceso del enlace no es una sesión: no abre las rutas con cuenta."""
    token, clave, _id = enlace
    cliente = _con_acceso(token, clave)
    for ruta in (
        '/api/bi-trade/avance-mensual',
        '/api/bi-trade/avance-mensual/exportar',
        '/api/bi-trade/ventas',
        '/api/bi-trade/enlaces',
    ):
        assert cliente.get(ruta).status_code == 401, ruta


def test_eliminar_un_enlace(app_bi_trade, enlace):
    token, clave, id_enlace = enlace
    editor = crear_usuario('otro@supli.tech', app_bi_trade, ['bi-trade:data:manage'])

    assert cliente_de(editor).delete(f'/api/bi-trade/enlaces/{id_enlace}').status_code == 204
    assert not EnlacePublico.objects.filter(pk=id_enlace).exists()
    assert _entrar(token, clave).status_code == 404


# ── Filtro por punto de venta en el mes y en el concurso ───────────────────

def test_el_avance_del_mes_se_filtra_por_punto_de_venta(app_bi_trade, catalogo):
    norte, sur, caro, barato = catalogo
    for punto, producto in ((norte, caro), (sur, barato)):
        MetaComercial.objects.create(
            id_producto=producto,
            id_punto_venta=punto,
            fecha_meta=date(2026, 3, 1),
            meta_cantidad=10,
        )
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=2
    )
    Venta.objects.create(
        id_producto=barato, id_punto_venta=sur, fecha_venta=date(2026, 3, 4), cantidad_vendida=5
    )
    Inventario.objects.create(id_producto=caro, id_punto_venta=norte, cantidad_inventario=7)

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get(
        '/api/bi-trade/avance-mensual?anio=2026&mes=3&id_punto_venta=PDV-001'
    ).json()

    # Solo el norte, en los tres lados: 10 × 4.000.000 de meta, 2 × 4.000.000
    # vendidos y su propio inventario.
    assert datos['totales']['metaDinero'] == 40_000_000
    assert datos['totales']['ventasDinero'] == 8_000_000
    assert datos['totales']['inventarioUnidades'] == 7
    assert [f['label'] for f in datos['porPuntoVenta']] == ['Claro Unicentro']
    assert datos['filtros']['idPuntoVenta'] == 'PDV-001'


def test_el_concurso_se_filtra_por_punto_de_venta(app_bi_trade, campana, catalogo):
    norte, sur, _caro, barato = catalogo
    _vender(norte, barato, 1, 15)
    _vender(sur, barato, 1, 15)

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    datos = cliente_de(lector).get('/api/bi-trade/tickets?id_punto_venta=PDV-002').json()

    assert [f['label'] for f in datos['filas']] == ['Claro Jardín Plaza']


def test_la_descarga_del_mes_respeta_el_punto_de_venta(app_bi_trade, catalogo):
    norte, sur, caro, barato = catalogo
    for punto, producto in ((norte, caro), (sur, barato)):
        MetaComercial.objects.create(
            id_producto=producto,
            id_punto_venta=punto,
            fecha_meta=date(2026, 3, 1),
            meta_cantidad=10,
        )

    lector = crear_usuario('lector@supli.tech', app_bi_trade)
    hoja = _hoja_de(
        cliente_de(lector).get(
            '/api/bi-trade/avance-mensual/exportar'
            '?anio=2026&mes=3&hoja=puntos&id_punto_venta=PDV-002'
        )
    )['Por punto de venta']

    assert hoja.max_row == 2
    assert hoja['A2'].value == 'Claro Jardín Plaza'


# ── Homecenter ─────────────────────────────────────────────────────────────

@pytest.fixture
def catalogo_hc():
    """Dos tiendas y dos productos de Homecenter."""
    tienda = PuntoVentaHc.objects.create(
        id_punto_venta='HC-101',
        nombre_pdv='Homecenter Calle 80',
        regional=RegionalHc.ZONA_NORTE,
    )
    centro = PuntoVentaHc.objects.create(
        id_punto_venta='HC-102',
        nombre_pdv='Homecenter Avenida 68',
        regional=RegionalHc.ZONA_CENTRO,
    )
    torre = ProductoHc.objects.create(
        id_producto='HCP-001',
        nombre_producto='Torre de sonido',
        marca='Aiwa',
        precio_venta_hc=1_200_000,
        precio_venta_coltrade=1_000_000,
    )
    bombillo = ProductoHc.objects.create(
        id_producto='HCP-002',
        nombre_producto='Bombillo inteligente',
        marca='Sylvania',
        precio_venta_hc=40_000,
        precio_venta_coltrade=30_000,
    )
    return tienda, centro, torre, bombillo


def _meta_hc(punto, producto, unidades, fecha=date(2026, 3, 10)):
    return MetaComercialHc.objects.create(
        id_producto=producto, id_punto_venta=punto, fecha_meta=fecha, meta_cantidad=unidades
    )


def _venta_hc(punto, producto, unidades, fecha=date(2026, 3, 4)):
    return VentaHc.objects.create(
        id_producto=producto, id_punto_venta=punto, fecha_venta=fecha, cantidad_vendida=unidades
    )


def _editor(app):
    return cliente_de(crear_usuario('editor@supli.tech', app, ['bi-trade:data:manage']))


def _xlsx(encabezados, filas):
    """Un .xlsx en memoria con una hoja «Datos», como el que sube una persona."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Datos'
    hoja.append(encabezados)
    for fila in filas:
        hoja.append(fila)
    buffer = BytesIO()
    libro.save(buffer)
    return SimpleUploadedFile(
        'datos.xlsx',
        buffer.getvalue(),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def test_las_tablas_y_llaves_de_hc_terminan_en_hc():
    for modelo in (PuntoVentaHc, ProductoHc, VentaHc, InventarioHc, MetaComercialHc):
        assert modelo._meta.db_table.endswith('_hc'), modelo
        assert modelo._meta.pk.column.endswith('_hc'), modelo
    for modelo in (VentaHc, InventarioHc, MetaComercialHc):
        assert modelo._meta.get_field('id_producto').column == 'id_producto_hc'
        assert modelo._meta.get_field('id_punto_venta').column == 'id_punto_venta_hc'
    assert ProductoHc._meta.get_field('precio_venta_hc').column == 'precio_venta_hc'


def test_claro_y_hc_no_se_mezclan(app_bi_trade, catalogo, catalogo_hc):
    norte, _sur, caro, _barato = catalogo
    tienda, _centro, torre, _bombillo = catalogo_hc
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=1
    )
    _venta_hc(tienda, torre, 3)

    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    claro = cliente.get('/api/bi-trade/avance-mensual?anio=2026&mes=3').json()
    hc = cliente.get('/api/bi-trade/hc/avance-mensual?anio=2026&mes=3').json()

    assert claro['totales']['ventasDinero'] == 4_000_000
    assert hc['totales']['ventasDinero'] == 3_000_000
    assert cliente.get('/api/bi-trade/ventas').json()['total'] == 1
    assert cliente.get('/api/bi-trade/hc/ventas').json()['total'] == 1
    tiendas = cliente.get('/api/bi-trade/hc/puntos-venta').json()
    assert [p['idPuntoVenta'] for p in tiendas] == ['HC-102', 'HC-101']


def test_la_cuota_diaria_de_hc_se_reparte_igual_que_en_claro(app_bi_trade, catalogo_hc):
    tienda, _centro, torre, _bombillo = catalogo_hc
    # 25 × 1.000.000, cargada con un día cualquiera de marzo.
    _meta_hc(tienda, torre, 25)

    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    datos = cliente.get('/api/bi-trade/hc/avance-mensual?anio=2026&mes=3').json()

    assert datos['periodo']['diasHabiles'] == 25
    assert datos['totales']['metaDinero'] == 25_000_000
    assert datos['totales']['metaDiaria'] == 1_000_000


def test_el_cumplimiento_diario_de_hc(app_bi_trade, catalogo_hc):
    tienda, _centro, torre, _bombillo = catalogo_hc
    _meta_hc(tienda, torre, 25)
    _venta_hc(tienda, torre, 2)

    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    datos = cliente.get('/api/bi-trade/hc/cumplimiento-diario?fecha=2026-03-04').json()

    assert datos['totales']['realDinero'] == 2_000_000
    assert datos['totales']['metaDinero'] == 1_000_000
    assert datos['totales']['cumplimientoDinero'] == 200.0
    assert [f['label'] for f in datos['porRegional']] == ['Zona Norte']


def test_los_crud_de_hc_escriben_en_sus_tablas(app_bi_trade):
    editor = _editor(app_bi_trade)
    punto = editor.post(
        '/api/bi-trade/hc/puntos-venta',
        {'idPuntoVenta': 'HC-900', 'nombrePdv': 'Homecenter Prueba', 'regional': 'Zona Centro'},
        format='json',
    )
    assert punto.status_code == 201, punto.json()
    producto = editor.post(
        '/api/bi-trade/hc/productos',
        {
            'idProducto': 'HCP-900',
            'nombreProducto': 'Taladro',
            'marca': 'Black+Decker',
            'precioVentaClaro': 50_000,
            'precioVentaColtrade': 40_000,
        },
        format='json',
    )
    assert producto.status_code == 201, producto.json()
    venta = editor.post(
        '/api/bi-trade/hc/ventas',
        {
            'idProducto': 'HCP-900',
            'idPuntoVenta': 'HC-900',
            'fechaVenta': '2026-03-04',
            'cantidadVendida': 2,
        },
        format='json',
    )
    assert venta.status_code == 201, venta.json()
    assert venta.data['total_coltrade'] == 80_000

    # En sus tablas, no en las de Claro.
    assert ProductoHc.objects.get(pk='HCP-900').precio_venta_hc == 50_000
    assert not PuntoVenta.objects.filter(pk='HC-900').exists()
    assert not Producto.objects.filter(pk='HCP-900').exists()
    assert VentaHc.objects.count() == 1
    assert not Venta.objects.exists()


def test_hc_no_acepta_la_regional_plaza_claro(app_bi_trade):
    respuesta = _editor(app_bi_trade).post(
        '/api/bi-trade/hc/puntos-venta',
        {'idPuntoVenta': 'HC-901', 'nombrePdv': 'Homecenter X', 'regional': 'Plaza Claro'},
        format='json',
    )
    assert respuesta.status_code == 400


def test_una_venta_hc_no_acepta_un_producto_de_claro(app_bi_trade, catalogo, catalogo_hc):
    respuesta = _editor(app_bi_trade).post(
        '/api/bi-trade/hc/ventas',
        {
            'idProducto': 'SKU-1',
            'idPuntoVenta': 'HC-101',
            'fechaVenta': '2026-03-04',
            'cantidadVendida': 1,
        },
        format='json',
    )
    assert respuesta.status_code == 400


def test_no_se_borra_un_producto_hc_con_ventas(app_bi_trade, catalogo_hc):
    tienda, _centro, torre, _bombillo = catalogo_hc
    _venta_hc(tienda, torre, 1)

    assert _editor(app_bi_trade).delete('/api/bi-trade/hc/productos/HCP-001').status_code == 400
    assert ProductoHc.objects.filter(pk='HCP-001').exists()


def test_las_plantillas_de_hc_usan_encabezados_hc(app_bi_trade):
    editor = _editor(app_bi_trade)
    esperados = {
        'puntos-venta': [
            'id_punto_venta_hc',
            'nombre_pdv',
            'regional',
            'materiales',
            'categoria',
        ],
        'productos': [
            'id_producto_hc',
            'ean',
            'sku_coltrade',
            'nombre_producto',
            'marca',
            'precio_venta_coltrade',
            'precio_venta_hc',
        ],
        'ventas': ['id_producto_hc', 'id_punto_venta_hc', 'fecha_venta', 'cantidad_vendida'],
        'inventario': ['id_producto_hc', 'id_punto_venta_hc', 'cantidad_inventario'],
        'metas': ['id_producto_hc', 'id_punto_venta_hc', 'fecha_meta', 'meta_cantidad'],
    }
    for recurso, encabezados in esperados.items():
        respuesta = editor.get(f'/api/bi-trade/hc/{recurso}/plantilla')
        assert respuesta.status_code == 200, recurso
        assert '-hc-' in respuesta['Content-Disposition'], recurso
        hoja = _hoja_de(respuesta)['Datos']
        assert [celda.value for celda in hoja[1]] == encabezados, recurso


def test_un_producto_hc_solo_exige_codigo_y_nombre(app_bi_trade):
    editor = _editor(app_bi_trade)
    respuesta = editor.post(
        '/api/bi-trade/hc/productos',
        {'idProducto': 'HCP-910', 'nombreProducto': 'Cable HDMI', 'marca': '  '},
        format='json',
    )
    assert respuesta.status_code == 201, respuesta.json()
    producto = ProductoHc.objects.get(pk='HCP-910')
    # Vacío es «sin dato»: NULL, no un texto en blanco.
    assert producto.marca is None
    assert producto.ean is None
    assert producto.precio_venta_hc is None
    assert producto.precio_venta_coltrade is None

    sin_nombre = editor.post(
        '/api/bi-trade/hc/productos', {'idProducto': 'HCP-911'}, format='json'
    )
    assert sin_nombre.status_code == 400


def test_la_categoria_de_un_punto_de_venta_hc(app_bi_trade):
    editor = _editor(app_bi_trade)
    ruta = '/api/bi-trade/hc/puntos-venta'
    sin_categoria = editor.post(
        ruta, {'idPuntoVenta': 'HC-950', 'nombrePdv': 'Homecenter Norte'}, format='json'
    )
    assert sin_categoria.status_code == 201, sin_categoria.json()
    assert PuntoVentaHc.objects.get(pk='HC-950').categoria is None

    minuscula = editor.patch(f'{ruta}/HC-950', {'categoria': 'a'}, format='json')
    assert minuscula.status_code == 200, minuscula.json()
    assert PuntoVentaHc.objects.get(pk='HC-950').categoria == 'A'
    assert editor.patch(f'{ruta}/HC-950', {'categoria': 'D'}, format='json').status_code == 400

    encabezados = ['id_punto_venta_hc', 'nombre_pdv', 'regional', 'materiales', 'categoria']
    filas = [
        ['HC-951', 'Homecenter Sur', None, None, 'c'],
        ['HC-952', 'Homecenter 80', None, None, None],
    ]
    importar = editor.post(
        f'{ruta}/importar', {'archivo': _xlsx(encabezados, filas)}, format='multipart'
    )
    assert importar.status_code == 200, importar.json()
    assert PuntoVentaHc.objects.get(pk='HC-951').categoria == 'C'
    assert PuntoVentaHc.objects.get(pk='HC-952').categoria is None
    assert 'categoria' not in [f.name for f in ProductoHc._meta.get_fields()]


def _query_hc(filas):
    """Un querie como el del portal de Homecenter: encabezado, fila de tipos y datos."""
    encabezados = [
        'Fecha', 'Código Proveedor', 'Código SKU', 'Desc. SKU', 'Ubicación', 'EAN Tienda',
        'Unidades',
    ]
    tipos = ['Fecha', 'Texto', 'Texto', 'Texto', 'Texto', 'Texto', 'Número']
    return _xlsx(encabezados, [tipos, *filas])


def test_el_querie_reemplaza_el_inventario_hc(app_bi_trade, catalogo_hc):
    tienda, centro, torre, bombillo = catalogo_hc
    proveedor = PuntoVentaHc.objects.create(
        id_punto_venta='7703670900993', nombre_pdv='Proveedor'
    )
    InventarioHc.objects.create(id_producto=torre, id_punto_venta=centro, cantidad_inventario=99)

    archivo = _query_hc(
        [
            # Día viejo: no entra.
            ['2026-09-18', '5112', 'HCP-001', 'Torre', 'Calle 80', 'HC-101', 50],
            # Día más reciente.
            ['2026-09-20', '5112', 'HCP-001', 'Torre', 'Calle 80', 'HC-101', 3],
            ['2026-09-20', '5112', 'HCP-002', 'Bombillo', 'Proveedor', None, 4],
            ['2026-09-20', '5112', 'HCP-002', 'Bombillo', 'Avenida 68', 'HC-102', 0],
            ['2026-09-20', '5112', 'HCP-001', 'Torre', 'Avenida 68', 'HC-102', -2],
            ['2026-09-20', '5112', 'NO-EXISTE', 'Otro', 'Calle 80', 'HC-101', 7],
            ['2026-09-20', '5112', 'HCP-002', 'Bombillo', 'Tienda nueva', '7700000000000', 2],
        ]
    )
    respuesta = _editor(app_bi_trade).post(
        '/api/bi-trade/hc/inventario/importar-query', {'archivo': archivo}, format='multipart'
    )
    assert respuesta.status_code == 200, respuesta.json()
    datos = respuesta.json()
    assert datos['fecha'] == '2026-09-20'
    assert datos['creados'] == 2
    assert datos['eliminados'] == 1
    assert datos['sinUnidades'] == 2
    assert datos['sinTienda'] == 1
    assert datos['productosFaltantes'] == ['NO-EXISTE']
    assert datos['puntosFaltantes'] == ['7700000000000']

    # Lo de antes se borró; queda solo el día más reciente, con 1 o más unidades.
    stock = {
        (i.id_producto_id, i.id_punto_venta_id): i.cantidad_inventario
        for i in InventarioHc.objects.all()
    }
    assert stock == {('HCP-001', 'HC-101'): 3, ('HCP-002', proveedor.pk): 4}


def test_un_querie_que_no_cruza_no_borra_el_inventario(app_bi_trade, catalogo_hc):
    tienda, _centro, torre, _bombillo = catalogo_hc
    InventarioHc.objects.create(id_producto=torre, id_punto_venta=tienda, cantidad_inventario=5)
    archivo = _query_hc([['2026-09-20', '5112', 'NO-EXISTE', 'Otro', 'X', 'HC-101', 7]])
    respuesta = _editor(app_bi_trade).post(
        '/api/bi-trade/hc/inventario/importar-query', {'archivo': archivo}, format='multipart'
    )
    assert respuesta.status_code == 400
    assert InventarioHc.objects.get().cantidad_inventario == 5


def test_un_querie_sin_sus_columnas_se_rechaza(app_bi_trade, catalogo_hc):
    archivo = _xlsx(['id_producto_hc', 'cantidad_inventario'], [['HCP-001', 3]])
    respuesta = _editor(app_bi_trade).post(
        '/api/bi-trade/hc/inventario/importar-query', {'archivo': archivo}, format='multipart'
    )
    assert respuesta.status_code == 400
    assert 'Código SKU' in respuesta.json()['message']


RUTA_QUERY_VENTAS = '/api/bi-trade/hc/ventas/importar-query'


def _query_ventas_hc(filas):
    """Un querie de ventas como el del portal: encabezado, fila de tipos y datos."""
    encabezados = [
        'Fecha', 'Tipo Registro', 'Código SKU', 'Nombre Tienda', 'EAN Tienda',
        'Unidades Vendidas', 'Ventas Pesos',
    ]
    tipos = ['Fecha', 'Texto', 'Texto', 'Texto', 'Texto', 'Número', 'Número - Moneda']
    return _xlsx(encabezados, [tipos, *filas])


def _archivo_ventas():
    return _query_ventas_hc(
        [
            ['2026-09-01', 'Venta', 'HCP-001', 'Calle 80', 'HC-101', 2, 100],
            ['2026-09-01', 'Venta', 'HCP-002', 'Venta Distancia Bogota', None, 1, 100],
            # Dos filas iguales el mismo día (otro canal): son dos ventas.
            ['2026-09-02', 'Venta', 'HCP-001', 'Calle 80', 'HC-101', 1, 100],
            ['2026-09-02', 'Venta', 'HCP-001', 'Calle 80', 'HC-101', 1, 100],
            ['2026-09-02', 'Devoluciones', 'HCP-001', 'Calle 80', 'HC-101', -1, -100],
            ['2026-09-02', 'Ajustes', 'HCP-002', 'Calle 80', 'HC-101', 0, 0],
            ['2026-09-02', 'Venta', 'NO-EXISTE', 'Calle 80', 'HC-101', 3, 100],
        ]
    )


@pytest.fixture
def tienda_sin_ean():
    return PuntoVentaHc.objects.create(id_punto_venta='7703670900993', nombre_pdv='Distancia')


def test_el_querie_de_ventas_carga_los_dias_del_archivo(
    app_bi_trade, catalogo_hc, tienda_sin_ean
):
    respuesta = _editor(app_bi_trade).post(
        RUTA_QUERY_VENTAS, {'archivo': _archivo_ventas()}, format='multipart'
    )
    assert respuesta.status_code == 200, respuesta.json()
    datos = respuesta.json()
    assert datos['creadas'] == 4
    assert datos['unidades'] == 5
    assert datos['dias'] == ['2026-09-01', '2026-09-02']
    assert datos['sinUnidades'] == 2
    assert datos['sinTienda'] == 1
    assert datos['productosFaltantes'] == ['NO-EXISTE']
    assert VentaHc.objects.filter(id_punto_venta=tienda_sin_ean).count() == 1
    assert VentaHc.objects.filter(fecha_venta=date(2026, 9, 2)).count() == 2


def test_el_querie_de_ventas_no_duplica_un_dia_ya_cargado(
    app_bi_trade, catalogo_hc, tienda_sin_ean
):
    tienda, _centro, torre, _bombillo = catalogo_hc
    editor = _editor(app_bi_trade)
    # Una venta de otro día no se toca nunca.
    _venta_hc(tienda, torre, 9, fecha=date(2026, 9, 3))
    assert editor.post(
        RUTA_QUERY_VENTAS, {'archivo': _archivo_ventas()}, format='multipart'
    ).status_code == 200

    # Subir el mismo archivo otra vez no guarda nada: avisa qué días ya tienen ventas.
    repetido = editor.post(RUTA_QUERY_VENTAS, {'archivo': _archivo_ventas()}, format='multipart')
    assert repetido.status_code == 409
    assert repetido.json()['code'] == 'dias_con_ventas'
    assert repetido.json()['diasConVentas'] == [
        {'fecha': '2026-09-01', 'registros': 2},
        {'fecha': '2026-09-02', 'registros': 2},
    ]
    assert VentaHc.objects.count() == 5

    # Sobrescribir borra esos días y los vuelve a cargar: el total no cambia.
    sobrescrito = editor.post(
        RUTA_QUERY_VENTAS,
        {'archivo': _archivo_ventas(), 'modo': 'sobrescribir'},
        format='multipart',
    )
    assert sobrescrito.status_code == 200, sobrescrito.json()
    assert sobrescrito.json()['eliminadas'] == 4
    assert VentaHc.objects.count() == 5
    assert VentaHc.objects.get(fecha_venta=date(2026, 9, 3)).cantidad_vendida == 9


def test_un_querie_de_ventas_que_no_cruza_no_guarda_nada(app_bi_trade, catalogo_hc):
    archivo = _query_ventas_hc(
        [['2026-09-01', 'Venta', 'NO-EXISTE', 'Calle 80', 'HC-101', 2, 100]]
    )
    respuesta = _editor(app_bi_trade).post(
        RUTA_QUERY_VENTAS, {'archivo': archivo}, format='multipart'
    )
    assert respuesta.status_code == 400
    assert not VentaHc.objects.exists()


def test_importar_productos_hc_con_ean_y_campos_vacios(app_bi_trade):
    encabezados = [
        'id_producto_hc',
        'ean',
        'sku_coltrade',
        'nombre_producto',
        'marca',
        'precio_venta_coltrade',
        'precio_venta_hc',
    ]
    filas = [
        ['HCP-920', 7701234567890, 'CT-920', 'Parlante', 'Aiwa', 90_000, 120_000],
        ['HCP-921', None, None, 'Linterna', None, None, None],
    ]
    respuesta = _editor(app_bi_trade).post(
        '/api/bi-trade/hc/productos/importar',
        {'archivo': _xlsx(encabezados, filas)},
        format='multipart',
    )
    assert respuesta.status_code == 200, respuesta.json()
    assert respuesta.json()['created'] == 2
    parlante = ProductoHc.objects.get(pk='HCP-920')
    assert parlante.ean == '7701234567890'
    assert parlante.sku_coltrade == 'CT-920'
    linterna = ProductoHc.objects.get(pk='HCP-921')
    assert linterna.marca is None
    assert linterna.precio_venta_coltrade is None


def test_una_venta_hc_sin_precio_cuenta_unidades_pero_no_dinero(app_bi_trade, catalogo_hc):
    tienda, _centro, torre, _bombillo = catalogo_hc
    linterna = ProductoHc.objects.create(id_producto='HCP-930', nombre_producto='Linterna')
    _venta_hc(tienda, torre, 1)
    _venta_hc(tienda, linterna, 5)

    editor = _editor(app_bi_trade)
    ventas = editor.get('/api/bi-trade/hc/ventas')
    assert ventas.status_code == 200
    sin_precio = next(v for v in ventas.json()['items'] if v['idProducto'] == 'HCP-930')
    assert sin_precio['totalColtrade'] is None

    avance = editor.get('/api/bi-trade/hc/avance-mensual', {'anio': 2026, 'mes': 3})
    assert avance.status_code == 200, avance.json()
    opciones = editor.get('/api/bi-trade/hc/opciones').json()
    assert None not in opciones['marcas']


def test_importar_en_hc_con_encabezados_hc(app_bi_trade, catalogo_hc):
    editor = _editor(app_bi_trade)
    ruta = '/api/bi-trade/hc/inventario/importar'
    encabezados = ['id_producto_hc', 'id_punto_venta_hc', 'cantidad_inventario']

    primera = editor.post(
        ruta, {'archivo': _xlsx(encabezados, [['HCP-001', 'HC-101', 7]])}, format='multipart'
    )
    assert primera.status_code == 200, primera.json()
    assert primera.json()['created'] == 1

    # La misma fila otra vez actualiza en vez de duplicar.
    segunda = editor.post(
        ruta, {'archivo': _xlsx(encabezados, [['HCP-001', 'HC-101', 9]])}, format='multipart'
    )
    assert segunda.json()['updated'] == 1
    assert InventarioHc.objects.get().cantidad_inventario == 9


def test_las_descargas_de_hc_llevan_su_prefijo(app_bi_trade, catalogo_hc):
    tienda, _centro, torre, _bombillo = catalogo_hc
    _venta_hc(tienda, torre, 1)
    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))

    ventas = cliente.get('/api/bi-trade/hc/ventas/exportar')
    assert 'filename="ventas-hc-' in ventas['Content-Disposition']
    assert _hoja_de(ventas).active.max_row == 2

    mes = cliente.get('/api/bi-trade/hc/avance-mensual/exportar?anio=2026&mes=3')
    assert mes['Content-Disposition'] == 'attachment; filename="hc-avance-2026-03.xlsx"'

    dia = cliente.get('/api/bi-trade/hc/cumplimiento-diario/exportar?fecha=2026-03-04')
    assert dia['Content-Disposition'] == (
        'attachment; filename="hc-cumplimiento-2026-03-04.xlsx"'
    )


def test_cada_modulo_de_hc_se_exporta_con_todos_sus_campos(app_bi_trade, catalogo_hc):
    tienda, _centro, torre, _bombillo = catalogo_hc
    tienda.categoria = 'A'
    tienda.save()
    torre.ean = '7701234567890'
    torre.save()
    _venta_hc(tienda, torre, 2)
    _meta_hc(tienda, torre, 5)
    InventarioHc.objects.create(id_producto=torre, id_punto_venta=tienda, cantidad_inventario=4)
    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))

    def exportar(recurso, **filtros):
        respuesta = cliente.get(f'/api/bi-trade/hc/{recurso}/exportar', filtros)
        assert respuesta.status_code == 200, recurso
        assert f'filename="{recurso}-hc-' in respuesta['Content-Disposition'], recurso
        hoja = _hoja_de(respuesta).active
        filas = [[celda.value for celda in fila] for fila in hoja.iter_rows()]
        return [dict(zip(filas[0], fila, strict=True)) for fila in filas[1:]]

    productos = exportar('productos')
    assert len(productos) == 2
    torre_export = next(f for f in productos if f['Código producto (SKU)'] == 'HCP-001')
    assert torre_export['EAN'] == '7701234567890'
    assert torre_export['Precio Homecenter'] == 1_200_000
    assert torre_export['Ventas registradas'] == 1
    # El filtro de la pantalla también filtra el archivo.
    assert len(exportar('productos', search='Bombillo')) == 1

    puntos = exportar('puntos-venta')
    assert {p['Código punto de venta']: p['Categoría'] for p in puntos}['HC-101'] == 'A'

    [venta] = exportar('ventas')
    assert venta['Categoría'] == 'A'
    assert venta['EAN'] == '7701234567890'
    assert venta['Total'] == 2_000_000
    assert 'Meta puntos' not in venta

    [inventario] = exportar('inventario')
    assert inventario['Valorizado'] == 4_000_000

    [meta] = exportar('metas')
    assert meta['Meta dinero'] == 5_000_000
    assert 'Meta puntos' not in meta


def test_las_opciones_de_hc_traen_sus_regionales_y_marcas(app_bi_trade, catalogo_hc):
    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    datos = cliente.get('/api/bi-trade/hc/opciones').json()

    regionales = [r['value'] for r in datos['regionales']]
    assert 'Zona Centro' in regionales
    assert 'Plaza Claro' not in regionales
    assert datos['marcas'] == ['Aiwa', 'Sylvania']


def test_escribir_en_hc_pide_permiso(app_bi_trade):
    lector = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    assert lector.get('/api/bi-trade/hc/productos').status_code == 200
    respuesta = lector.post('/api/bi-trade/hc/productos', {'idProducto': 'X'}, format='json')
    assert respuesta.status_code == 403


def test_el_enlace_publico_de_hc_muestra_hc(app_bi_trade, catalogo_hc):
    cache.clear()
    tienda, _centro, torre, _bombillo = catalogo_hc
    _venta_hc(tienda, torre, 3)
    editor = _editor(app_bi_trade)
    creado = editor.post(
        '/api/bi-trade/enlaces', {'nombre': 'Gerencia HC', 'canal': 'hc'}, format='json'
    ).json()
    token, clave = creado['token'], creado['clave']

    puerta = APIClient().get(f'{RUTA_PUBLICA}/{token}').json()
    assert puerta == {'nombre': 'Gerencia HC', 'canal': 'hc'}

    cliente = _con_acceso(token, clave)
    avance = cliente.get(f'{RUTA_PUBLICA}/{token}/avance-mensual?anio=2026&mes=3').json()
    assert avance['totales']['ventasDinero'] == 3_000_000
    # Homecenter no tiene concurso.
    assert cliente.get(f'{RUTA_PUBLICA}/{token}/tickets').status_code == 404
    assert cliente.get(f'{RUTA_PUBLICA}/{token}/campanas').json() == []
    productos = cliente.get(f'{RUTA_PUBLICA}/{token}/productos').json()
    assert {p['idProducto'] for p in productos} == {'HCP-001', 'HCP-002'}
    assert set(productos[0]) == {'idProducto', 'nombreProducto', 'marca'}

    # Cada tablero lista sus propios enlaces.
    de_hc = editor.get('/api/bi-trade/enlaces?canal=hc').json()
    assert [e['nombre'] for e in de_hc] == ['Gerencia HC']
    assert editor.get('/api/bi-trade/enlaces?canal=claro').json() == []


def test_la_semilla_de_hc_carga_datos_y_no_duplica():
    call_command('seed_bi_trade_hc_demo', stdout=StringIO())
    ventas = VentaHc.objects.count()
    metas = MetaComercialHc.objects.count()

    assert PuntoVentaHc.objects.count() == 5
    assert ProductoHc.objects.count() == 8
    assert InventarioHc.objects.count() == 5 * 8
    assert metas == 5 * 8 * 2  # el mes anterior y el actual
    assert ventas > 0

    call_command('seed_bi_trade_hc_demo', stdout=StringIO())
    assert VentaHc.objects.count() == ventas
    assert MetaComercialHc.objects.count() == metas
    # Lo de Claro no se toca.
    assert not Venta.objects.exists()
    assert not PuntoVenta.objects.exists()


# ── Falabella ──────────────────────────────────────────────────────────────

@pytest.fixture
def catalogo_falabella():
    """Una tienda y un producto de Falabella."""
    tienda = PuntoVentaFalabella.objects.create(
        id_punto_venta='FAL-201',
        nombre_pdv='Falabella Andino',
        regional=RegionalFalabella.ZONA_NORTE,
    )
    audifonos = ProductoFalabella.objects.create(
        id_producto='FALP-001',
        nombre_producto='Audífonos JBL',
        marca='JBL',
        precio_venta_falabella=300_000,
        precio_venta_coltrade=250_000,
    )
    return tienda, audifonos


def _venta_falabella(punto, producto, unidades, fecha=date(2026, 3, 4)):
    return VentaFalabella.objects.create(
        id_producto=producto, id_punto_venta=punto, fecha_venta=fecha, cantidad_vendida=unidades
    )


def test_las_tablas_y_llaves_de_falabella_terminan_en_falabella():
    modelos = (
        PuntoVentaFalabella,
        ProductoFalabella,
        VentaFalabella,
        InventarioFalabella,
        MetaComercialFalabella,
    )
    for modelo in modelos:
        assert modelo._meta.db_table.endswith('_falabella'), modelo
        assert modelo._meta.pk.column.endswith('_falabella'), modelo
    for modelo in modelos[2:]:
        assert modelo._meta.get_field('id_producto').column == 'id_producto_falabella'
        assert modelo._meta.get_field('id_punto_venta').column == 'id_punto_venta_falabella'
    campo = ProductoFalabella._meta.get_field('precio_venta_falabella')
    assert campo.column == 'precio_venta_falabella'


def test_los_tres_canales_no_se_mezclan(app_bi_trade, catalogo_hc, catalogo_falabella):
    tienda_hc, _centro, torre, _bombillo = catalogo_hc
    tienda, audifonos = catalogo_falabella
    _venta_hc(tienda_hc, torre, 3)
    _venta_falabella(tienda, audifonos, 2)

    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    mes = '?anio=2026&mes=3'
    claro = cliente.get(f'/api/bi-trade/avance-mensual{mes}').json()
    hc = cliente.get(f'/api/bi-trade/hc/avance-mensual{mes}').json()
    falabella = cliente.get(f'/api/bi-trade/falabella/avance-mensual{mes}').json()

    assert claro['totales']['ventasDinero'] == 0
    assert hc['totales']['ventasDinero'] == 3_000_000
    assert falabella['totales']['ventasDinero'] == 500_000
    assert cliente.get('/api/bi-trade/falabella/ventas').json()['total'] == 1
    tiendas = cliente.get('/api/bi-trade/falabella/puntos-venta').json()
    assert [p['idPuntoVenta'] for p in tiendas] == ['FAL-201']


def test_la_cuota_y_el_dia_de_falabella(app_bi_trade, catalogo_falabella):
    tienda, audifonos = catalogo_falabella
    # 25 × 250.000 en marzo de 2026, que tiene 25 días hábiles.
    MetaComercialFalabella.objects.create(
        id_producto=audifonos, id_punto_venta=tienda, fecha_meta=date(2026, 3, 10),
        meta_cantidad=25,
    )
    _venta_falabella(tienda, audifonos, 2)

    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    mes = cliente.get('/api/bi-trade/falabella/avance-mensual?anio=2026&mes=3').json()
    assert mes['periodo']['diasHabiles'] == 25
    assert mes['totales']['metaDinero'] == 6_250_000
    assert mes['totales']['metaDiaria'] == 250_000

    dia = cliente.get('/api/bi-trade/falabella/cumplimiento-diario?fecha=2026-03-04').json()
    assert dia['totales']['realDinero'] == 500_000
    assert dia['totales']['metaDinero'] == 250_000
    assert dia['totales']['cumplimientoDinero'] == 200.0


def test_los_crud_de_falabella_escriben_en_sus_tablas(app_bi_trade):
    editor = _editor(app_bi_trade)
    assert editor.post(
        '/api/bi-trade/falabella/puntos-venta',
        {'idPuntoVenta': 'FAL-900', 'nombrePdv': 'Falabella Prueba', 'regional': 'Zona Sur'},
        format='json',
    ).status_code == 201
    assert editor.post(
        '/api/bi-trade/falabella/productos',
        {
            'idProducto': 'FALP-900',
            'nombreProducto': 'Producto de prueba',
            'marca': 'JBL',
            'precioVentaClaro': 50_000,
            'precioVentaColtrade': 40_000,
        },
        format='json',
    ).status_code == 201

    assert ProductoFalabella.objects.get(pk='FALP-900').precio_venta_falabella == 50_000
    assert PuntoVentaFalabella.objects.filter(pk='FAL-900').exists()
    assert not PuntoVenta.objects.filter(pk='FAL-900').exists()
    assert not PuntoVentaHc.objects.filter(pk='FAL-900').exists()
    assert not ProductoHc.objects.filter(pk='FALP-900').exists()


def test_falabella_no_acepta_plaza_claro_ni_productos_de_otro_canal(
    app_bi_trade, catalogo_hc, catalogo_falabella
):
    editor = _editor(app_bi_trade)
    plaza = editor.post(
        '/api/bi-trade/falabella/puntos-venta',
        {'idPuntoVenta': 'FAL-901', 'nombrePdv': 'Falabella X', 'regional': 'Plaza Claro'},
        format='json',
    )
    assert plaza.status_code == 400

    de_hc = editor.post(
        '/api/bi-trade/falabella/ventas',
        {
            'idProducto': 'HCP-001',
            'idPuntoVenta': 'FAL-201',
            'fechaVenta': '2026-03-04',
            'cantidadVendida': 1,
        },
        format='json',
    )
    assert de_hc.status_code == 400
    assert not VentaFalabella.objects.exists()


def test_las_plantillas_de_falabella_usan_encabezados_falabella(app_bi_trade):
    editor = _editor(app_bi_trade)
    producto, punto = 'id_producto_falabella', 'id_punto_venta_falabella'
    esperados = {
        'puntos-venta': [punto, 'nombre_pdv', 'regional', 'materiales'],
        'productos': [
            producto,
            'nombre_producto',
            'marca',
            'precio_venta_falabella',
            'precio_venta_coltrade',
        ],
        'ventas': [producto, punto, 'fecha_venta', 'cantidad_vendida'],
        'inventario': [producto, punto, 'cantidad_inventario'],
        'metas': [producto, punto, 'fecha_meta', 'meta_cantidad'],
    }
    for recurso, encabezados in esperados.items():
        respuesta = editor.get(f'/api/bi-trade/falabella/{recurso}/plantilla')
        assert respuesta.status_code == 200, recurso
        assert '-falabella' in respuesta['Content-Disposition'], recurso
        hoja = _hoja_de(respuesta)['Datos']
        assert [celda.value for celda in hoja[1]] == encabezados, recurso


def test_importar_en_falabella_con_encabezados_falabella(app_bi_trade, catalogo_falabella):
    editor = _editor(app_bi_trade)
    ruta = '/api/bi-trade/falabella/inventario/importar'
    encabezados = ['id_producto_falabella', 'id_punto_venta_falabella', 'cantidad_inventario']

    primera = editor.post(
        ruta, {'archivo': _xlsx(encabezados, [['FALP-001', 'FAL-201', 7]])}, format='multipart'
    )
    assert primera.status_code == 200, primera.json()
    assert primera.json()['created'] == 1

    segunda = editor.post(
        ruta, {'archivo': _xlsx(encabezados, [['FALP-001', 'FAL-201', 9]])}, format='multipart'
    )
    assert segunda.json()['updated'] == 1
    assert InventarioFalabella.objects.get().cantidad_inventario == 9


def test_las_descargas_y_opciones_de_falabella(app_bi_trade, catalogo_falabella):
    tienda, audifonos = catalogo_falabella
    _venta_falabella(tienda, audifonos, 1)
    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))

    ventas = cliente.get('/api/bi-trade/falabella/ventas/exportar')
    assert 'filename="ventas-falabella-' in ventas['Content-Disposition']
    mes = cliente.get('/api/bi-trade/falabella/avance-mensual/exportar?anio=2026&mes=3')
    assert mes['Content-Disposition'] == 'attachment; filename="falabella-avance-2026-03.xlsx"'
    dia = cliente.get('/api/bi-trade/falabella/cumplimiento-diario/exportar?fecha=2026-03-04')
    assert dia['Content-Disposition'] == (
        'attachment; filename="falabella-cumplimiento-2026-03-04.xlsx"'
    )

    opciones = cliente.get('/api/bi-trade/falabella/opciones').json()
    regionales = [r['value'] for r in opciones['regionales']]
    assert 'Zona Centro' in regionales
    assert 'Plaza Claro' not in regionales
    assert opciones['marcas'] == ['JBL']


def test_escribir_en_falabella_pide_permiso(app_bi_trade):
    lector = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    assert lector.get('/api/bi-trade/falabella/productos').status_code == 200
    respuesta = lector.post('/api/bi-trade/falabella/productos', {'idProducto': 'X'}, format='json')
    assert respuesta.status_code == 403


def test_el_enlace_publico_de_falabella_muestra_falabella(
    app_bi_trade, catalogo_hc, catalogo_falabella
):
    cache.clear()
    tienda, audifonos = catalogo_falabella
    _venta_falabella(tienda, audifonos, 2)
    editor = _editor(app_bi_trade)
    creado = editor.post(
        '/api/bi-trade/enlaces',
        {'nombre': 'Gerencia Falabella', 'canal': 'falabella'},
        format='json',
    ).json()
    token, clave = creado['token'], creado['clave']

    puerta = APIClient().get(f'{RUTA_PUBLICA}/{token}').json()
    assert puerta == {'nombre': 'Gerencia Falabella', 'canal': 'falabella'}

    cliente = _con_acceso(token, clave)
    avance = cliente.get(f'{RUTA_PUBLICA}/{token}/avance-mensual?anio=2026&mes=3').json()
    assert avance['totales']['ventasDinero'] == 500_000
    assert cliente.get(f'{RUTA_PUBLICA}/{token}/tickets').status_code == 404
    productos = cliente.get(f'{RUTA_PUBLICA}/{token}/productos').json()
    assert {p['idProducto'] for p in productos} == {'FALP-001'}

    de_falabella = editor.get('/api/bi-trade/enlaces?canal=falabella').json()
    assert [e['nombre'] for e in de_falabella] == ['Gerencia Falabella']
    assert editor.get('/api/bi-trade/enlaces?canal=hc').json() == []


def test_la_semilla_de_falabella_carga_datos_y_no_duplica():
    call_command('seed_bi_trade_falabella_demo', stdout=StringIO())
    ventas = VentaFalabella.objects.count()
    metas = MetaComercialFalabella.objects.count()

    assert PuntoVentaFalabella.objects.count() == 5
    assert ProductoFalabella.objects.count() == 8
    assert InventarioFalabella.objects.count() == 5 * 8
    assert metas == 5 * 8 * 2  # el mes anterior y el actual
    assert ventas > 0

    call_command('seed_bi_trade_falabella_demo', stdout=StringIO())
    assert VentaFalabella.objects.count() == ventas
    assert MetaComercialFalabella.objects.count() == metas
    # Lo de Claro y Homecenter no se toca.
    assert not Venta.objects.exists()
    assert not VentaHc.objects.exists()


# ── Tmk Ecommerce Claro ──────────────────────────────────────────────────────────────

@pytest.fixture
def catalogo_tmk():
    """Una tienda y un producto de Tmk Ecommerce Claro."""
    tienda = PuntoVentaTmk.objects.create(
        id_punto_venta='TMK-301',
        nombre_pdv='Tmk Call Center Bogotá',
        regional=RegionalTmk.ZONA_NORTE,
    )
    audifonos = ProductoTmk.objects.create(
        id_producto='TMKP-001',
        nombre_producto='Audífonos JBL',
        marca='JBL',
        precio_venta_tmk=300_000,
        precio_venta_coltrade=250_000,
    )
    return tienda, audifonos


def _venta_tmk(punto, producto, unidades, fecha=date(2026, 3, 4)):
    return VentaTmk.objects.create(
        id_producto=producto, id_punto_venta=punto, fecha_venta=fecha, cantidad_vendida=unidades
    )


def test_las_tablas_y_llaves_de_tmk_terminan_en_tmk():
    modelos = (
        PuntoVentaTmk,
        ProductoTmk,
        VentaTmk,
        InventarioTmk,
        MetaComercialTmk,
    )
    for modelo in modelos:
        assert modelo._meta.db_table.endswith('_tmk'), modelo
        assert modelo._meta.pk.column.endswith('_tmk'), modelo
    for modelo in modelos[2:]:
        assert modelo._meta.get_field('id_producto').column == 'id_producto_tmk'
        assert modelo._meta.get_field('id_punto_venta').column == 'id_punto_venta_tmk'
    campo = ProductoTmk._meta.get_field('precio_venta_tmk')
    assert campo.column == 'precio_venta_tmk'


def test_tmk_no_se_mezcla_con_los_otros_canales(app_bi_trade, catalogo_hc, catalogo_tmk):
    tienda_hc, _centro, torre, _bombillo = catalogo_hc
    tienda, audifonos = catalogo_tmk
    _venta_hc(tienda_hc, torre, 3)
    _venta_tmk(tienda, audifonos, 2)

    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    mes = '?anio=2026&mes=3'
    claro = cliente.get(f'/api/bi-trade/avance-mensual{mes}').json()
    hc = cliente.get(f'/api/bi-trade/hc/avance-mensual{mes}').json()
    tmk = cliente.get(f'/api/bi-trade/tmk/avance-mensual{mes}').json()

    assert claro['totales']['ventasDinero'] == 0
    assert hc['totales']['ventasDinero'] == 3_000_000
    assert tmk['totales']['ventasDinero'] == 500_000
    assert cliente.get('/api/bi-trade/tmk/ventas').json()['total'] == 1
    tiendas = cliente.get('/api/bi-trade/tmk/puntos-venta').json()
    assert [p['idPuntoVenta'] for p in tiendas] == ['TMK-301']


def test_la_cuota_y_el_dia_de_tmk(app_bi_trade, catalogo_tmk):
    tienda, audifonos = catalogo_tmk
    # 25 × 250.000 en marzo de 2026, que tiene 25 días hábiles.
    MetaComercialTmk.objects.create(
        id_producto=audifonos, id_punto_venta=tienda, fecha_meta=date(2026, 3, 10),
        meta_cantidad=25,
    )
    _venta_tmk(tienda, audifonos, 2)

    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    mes = cliente.get('/api/bi-trade/tmk/avance-mensual?anio=2026&mes=3').json()
    assert mes['periodo']['diasHabiles'] == 25
    assert mes['totales']['metaDinero'] == 6_250_000
    assert mes['totales']['metaDiaria'] == 250_000

    dia = cliente.get('/api/bi-trade/tmk/cumplimiento-diario?fecha=2026-03-04').json()
    assert dia['totales']['realDinero'] == 500_000
    assert dia['totales']['metaDinero'] == 250_000
    assert dia['totales']['cumplimientoDinero'] == 200.0


def test_los_crud_de_tmk_escriben_en_sus_tablas(app_bi_trade):
    editor = _editor(app_bi_trade)
    assert editor.post(
        '/api/bi-trade/tmk/puntos-venta',
        {'idPuntoVenta': 'TMK-900', 'nombrePdv': 'Tmk Prueba', 'regional': 'Zona Sur'},
        format='json',
    ).status_code == 201
    assert editor.post(
        '/api/bi-trade/tmk/productos',
        {
            'idProducto': 'TMKP-900',
            'nombreProducto': 'Producto de prueba',
            'marca': 'JBL',
            'precioVentaClaro': 50_000,
            'precioVentaColtrade': 40_000,
        },
        format='json',
    ).status_code == 201

    assert ProductoTmk.objects.get(pk='TMKP-900').precio_venta_tmk == 50_000
    assert PuntoVentaTmk.objects.filter(pk='TMK-900').exists()
    assert not PuntoVenta.objects.filter(pk='TMK-900').exists()
    assert not PuntoVentaHc.objects.filter(pk='TMK-900').exists()
    assert not ProductoHc.objects.filter(pk='TMKP-900').exists()


def test_tmk_no_acepta_plaza_claro_ni_productos_de_otro_canal(
    app_bi_trade, catalogo_hc, catalogo_tmk
):
    editor = _editor(app_bi_trade)
    plaza = editor.post(
        '/api/bi-trade/tmk/puntos-venta',
        {'idPuntoVenta': 'TMK-901', 'nombrePdv': 'Tmk X', 'regional': 'Plaza Claro'},
        format='json',
    )
    assert plaza.status_code == 400

    de_hc = editor.post(
        '/api/bi-trade/tmk/ventas',
        {
            'idProducto': 'HCP-001',
            'idPuntoVenta': 'TMK-301',
            'fechaVenta': '2026-03-04',
            'cantidadVendida': 1,
        },
        format='json',
    )
    assert de_hc.status_code == 400
    assert not VentaTmk.objects.exists()


def test_las_plantillas_de_tmk_usan_encabezados_tmk(app_bi_trade):
    editor = _editor(app_bi_trade)
    producto, punto = 'id_producto_tmk', 'id_punto_venta_tmk'
    esperados = {
        'puntos-venta': [punto, 'nombre_pdv', 'regional', 'materiales'],
        'productos': [
            producto,
            'nombre_producto',
            'marca',
            'precio_venta_tmk',
            'precio_venta_coltrade',
        ],
        'ventas': [producto, punto, 'fecha_venta', 'cantidad_vendida'],
        'inventario': [producto, punto, 'cantidad_inventario'],
        'metas': [producto, punto, 'fecha_meta', 'meta_cantidad'],
    }
    for recurso, encabezados in esperados.items():
        respuesta = editor.get(f'/api/bi-trade/tmk/{recurso}/plantilla')
        assert respuesta.status_code == 200, recurso
        assert '-tmk' in respuesta['Content-Disposition'], recurso
        hoja = _hoja_de(respuesta)['Datos']
        assert [celda.value for celda in hoja[1]] == encabezados, recurso


def test_importar_en_tmk_con_encabezados_tmk(app_bi_trade, catalogo_tmk):
    editor = _editor(app_bi_trade)
    ruta = '/api/bi-trade/tmk/inventario/importar'
    encabezados = ['id_producto_tmk', 'id_punto_venta_tmk', 'cantidad_inventario']

    primera = editor.post(
        ruta, {'archivo': _xlsx(encabezados, [['TMKP-001', 'TMK-301', 7]])}, format='multipart'
    )
    assert primera.status_code == 200, primera.json()
    assert primera.json()['created'] == 1

    segunda = editor.post(
        ruta, {'archivo': _xlsx(encabezados, [['TMKP-001', 'TMK-301', 9]])}, format='multipart'
    )
    assert segunda.json()['updated'] == 1
    assert InventarioTmk.objects.get().cantidad_inventario == 9


def test_las_descargas_y_opciones_de_tmk(app_bi_trade, catalogo_tmk):
    tienda, audifonos = catalogo_tmk
    _venta_tmk(tienda, audifonos, 1)
    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))

    ventas = cliente.get('/api/bi-trade/tmk/ventas/exportar')
    assert 'filename="ventas-tmk-' in ventas['Content-Disposition']
    mes = cliente.get('/api/bi-trade/tmk/avance-mensual/exportar?anio=2026&mes=3')
    assert mes['Content-Disposition'] == 'attachment; filename="tmk-avance-2026-03.xlsx"'
    dia = cliente.get('/api/bi-trade/tmk/cumplimiento-diario/exportar?fecha=2026-03-04')
    assert dia['Content-Disposition'] == (
        'attachment; filename="tmk-cumplimiento-2026-03-04.xlsx"'
    )

    opciones = cliente.get('/api/bi-trade/tmk/opciones').json()
    regionales = [r['value'] for r in opciones['regionales']]
    assert 'Zona Centro' in regionales
    assert 'Plaza Claro' not in regionales
    assert opciones['marcas'] == ['JBL']


def test_escribir_en_tmk_pide_permiso(app_bi_trade):
    lector = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    assert lector.get('/api/bi-trade/tmk/productos').status_code == 200
    respuesta = lector.post('/api/bi-trade/tmk/productos', {'idProducto': 'X'}, format='json')
    assert respuesta.status_code == 403


def test_el_enlace_publico_de_tmk_muestra_tmk(
    app_bi_trade, catalogo_hc, catalogo_tmk
):
    cache.clear()
    tienda, audifonos = catalogo_tmk
    _venta_tmk(tienda, audifonos, 2)
    editor = _editor(app_bi_trade)
    creado = editor.post(
        '/api/bi-trade/enlaces',
        {'nombre': 'Gerencia Tmk Ecommerce Claro', 'canal': 'tmk'},
        format='json',
    ).json()
    token, clave = creado['token'], creado['clave']

    puerta = APIClient().get(f'{RUTA_PUBLICA}/{token}').json()
    assert puerta == {'nombre': 'Gerencia Tmk Ecommerce Claro', 'canal': 'tmk'}

    cliente = _con_acceso(token, clave)
    avance = cliente.get(f'{RUTA_PUBLICA}/{token}/avance-mensual?anio=2026&mes=3').json()
    assert avance['totales']['ventasDinero'] == 500_000
    assert cliente.get(f'{RUTA_PUBLICA}/{token}/tickets').status_code == 404
    productos = cliente.get(f'{RUTA_PUBLICA}/{token}/productos').json()
    assert {p['idProducto'] for p in productos} == {'TMKP-001'}

    de_tmk = editor.get('/api/bi-trade/enlaces?canal=tmk').json()
    assert [e['nombre'] for e in de_tmk] == ['Gerencia Tmk Ecommerce Claro']
    assert editor.get('/api/bi-trade/enlaces?canal=hc').json() == []


def test_la_semilla_de_tmk_carga_datos_y_no_duplica():
    call_command('seed_bi_trade_tmk_demo', stdout=StringIO())
    ventas = VentaTmk.objects.count()
    metas = MetaComercialTmk.objects.count()

    assert PuntoVentaTmk.objects.count() == 5
    assert ProductoTmk.objects.count() == 8
    assert InventarioTmk.objects.count() == 5 * 8
    assert metas == 5 * 8 * 2  # el mes anterior y el actual
    assert ventas > 0

    call_command('seed_bi_trade_tmk_demo', stdout=StringIO())
    assert VentaTmk.objects.count() == ventas
    assert MetaComercialTmk.objects.count() == metas
    # Lo de Claro y Homecenter no se toca.
    assert not Venta.objects.exists()
    assert not VentaHc.objects.exists()


# ── Tmk Ecommerce Claro: el informe del ERP ────────────────────────────────

def _importar_tmk(cliente, filas, anio=2026, mes=3, modo='completar'):
    return cliente.post(
        '/api/bi-trade/tmk/importar-informe',
        {'archivo': _informe(filas), 'anio': anio, 'mes': mes, 'modo': modo},
        format='multipart',
    )


def test_el_informe_del_erp_carga_en_las_tablas_de_tmk(editor, catalogo, catalogo_tmk):
    norte, _sur, caro, _barato = catalogo
    tienda, audifonos = catalogo_tmk
    Inventario.objects.create(id_producto=caro, id_punto_venta=norte, cantidad_inventario=5)
    InventarioTmk.objects.create(
        id_producto=audifonos, id_punto_venta=tienda, cantidad_inventario=99
    )

    respuesta = _importar_tmk(
        editor,
        [
            ('TMKP-001', 'TMK-301', '601', date(2026, 3, 4)),
            ('TMKP-001', 'TMK-301', '601', date(2026, 3, 4)),
            ('TMKP-001', 'TMK-301', '601A', date(2026, 3, 5)),
            ('TMKP-001', 'TMK-301', '1', None),
            # Existe en Claro, no en Tmk: no entra.
            ('SKU-1', 'PDV-001', '601', date(2026, 3, 4)),
        ],
    )

    assert respuesta.status_code == 200, respuesta.json()
    datos = respuesta.json()
    assert datos['ventas']['creadas'] == 1
    assert datos['ventas']['sinProducto'] == 1
    assert VentaTmk.objects.get().cantidad_vendida == 2
    # El inventario de Tmk se reemplaza; el de Claro no se toca.
    assert InventarioTmk.objects.get().cantidad_inventario == 1
    assert Inventario.objects.get().cantidad_inventario == 5
    assert not Venta.objects.exists()


def test_sobrescribir_en_tmk_solo_borra_las_ventas_de_tmk(editor, catalogo, catalogo_tmk):
    norte, _sur, caro, _barato = catalogo
    tienda, audifonos = catalogo_tmk
    Venta.objects.create(
        id_producto=caro, id_punto_venta=norte, fecha_venta=date(2026, 3, 4), cantidad_vendida=1
    )
    _venta_tmk(tienda, audifonos, 5)

    respuesta = _importar_tmk(
        editor, [('TMKP-001', 'TMK-301', '601', date(2026, 3, 6))], modo='sobrescribir'
    )

    assert respuesta.status_code == 200, respuesta.json()
    assert respuesta.json()['ventas']['eliminadas'] == 1
    assert list(VentaTmk.objects.values_list('fecha_venta', flat=True)) == [date(2026, 3, 6)]
    assert Venta.objects.count() == 1


def test_importar_el_informe_de_tmk_exige_permiso_de_edicion(app_bi_trade, catalogo_tmk):
    lector = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    filas = [('TMKP-001', 'TMK-301', '601', date(2026, 3, 4))]
    assert _importar_tmk(lector, filas).status_code == 403
    assert not VentaTmk.objects.exists()


def test_hc_y_falabella_no_tienen_importador_del_erp(editor):
    for canal in ('hc', 'falabella'):
        respuesta = editor.post(f'/api/bi-trade/{canal}/importar-informe', {}, format='multipart')
        assert respuesta.status_code == 404, canal


# ── Los canales aparte no miden en puntos ─────────────────────────────────

def test_solo_claro_exporta_puntos(app_bi_trade):
    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    dia = '/cumplimiento-diario/exportar?fecha=2026-03-04&corte=regional'

    def encabezados(ruta):
        respuesta = cliente.get(ruta)
        assert respuesta.status_code == 200, ruta
        return [celda.value for celda in _hoja_de(respuesta).active[1]]

    assert 'Vendido (pts)' in encabezados(f'/api/bi-trade{dia}')
    assert 'Meta puntos' in encabezados('/api/bi-trade/metas/exportar')
    for canal in ('hc', 'falabella', 'tmk'):
        del_dia = encabezados(f'/api/bi-trade/{canal}{dia}')
        assert 'Vendido ($)' in del_dia, canal
        assert not [e for e in del_dia if 'pts' in e], canal
        metas = encabezados(f'/api/bi-trade/{canal}/metas/exportar')
        assert 'Meta dinero' in metas, canal
        assert 'Meta puntos' not in metas, canal


def _llaves(datos):
    if isinstance(datos, dict):
        for clave, valor in datos.items():
            yield clave
            yield from _llaves(valor)
    elif isinstance(datos, list):
        for valor in datos:
            yield from _llaves(valor)


def _de_puntos(datos):
    """Las llaves del JSON que hablan de puntos (no de puntos de venta)."""
    return {
        clave for clave in _llaves(datos)
        if clave in ('puntos', 'puntaje') or clave.endswith('Puntos')
    }


def test_los_canales_aparte_no_tienen_puntos_en_nada(
    app_bi_trade, catalogo_hc, catalogo_falabella, catalogo_tmk
):
    for modelo in (ProductoHc, ProductoFalabella, ProductoTmk):
        assert 'puntaje' not in {campo.name for campo in modelo._meta.get_fields()}, modelo
    for modelo in (MetaComercialHc, MetaComercialFalabella, MetaComercialTmk):
        assert not hasattr(modelo, 'meta_puntos'), modelo

    tienda, _centro, torre, _bombillo = catalogo_hc
    _venta_hc(tienda, torre, 2)
    _meta_hc(tienda, torre, 5)
    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))

    for canal in ('hc', 'falabella', 'tmk'):
        for ruta in (
            'productos',
            'ventas',
            'ventas/resumen',
            'metas',
            'metas/resumen',
            'avance-mensual?anio=2026&mes=3',
            'cumplimiento-diario?fecha=2026-03-04',
        ):
            respuesta = cliente.get(f'/api/bi-trade/{canal}/{ruta}')
            assert respuesta.status_code == 200, (canal, ruta)
            assert _de_puntos(respuesta.json()) == set(), (canal, ruta)

    # Claro sí los tiene.
    dia = cliente.get('/api/bi-trade/cumplimiento-diario?fecha=2026-03-04').json()
    assert 'realPuntos' in _de_puntos(dia)


# ── Plan Partners ──────────────────────────────────────────────────────────

@pytest.fixture
def catalogo_partners():
    norte = RegionalPartner.objects.create(nombre='Región Centro (Z. Norte)')
    sur = RegionalPartner.objects.create(nombre='Región Centro (Z. Sur)')
    punto_norte = PuntoVentaPartner.objects.create(
        id_punto_venta='C192',
        nombre_pdv='Cav Cucuta Centro Av Quinta',
        id_regional=norte,
    )
    punto_sur = PuntoVentaPartner.objects.create(
        id_punto_venta='C108',
        nombre_pdv='Cav Bogota Plaza Claro',
        id_regional=sur,
    )
    producto = ProductoPartner.objects.create(id_producto='7015490', nombre_producto='Estandar')
    return punto_norte, punto_sur, producto


def _recomendacion(punto, producto, serial='69410004009849'):
    return {
        'idRegional': punto.id_regional_id,
        'marca': MarcaPartner.MOTOROLA.value,
        'idPuntoVenta': punto.pk,
        'idProducto': producto.pk,
        'fechaRecomendacion': str(timezone.localdate()),
        'serial': serial,
        'documentoPromotor': '1092389375',
        'factura': '10500000653130057938',
    }


def test_el_formulario_guarda_el_codigo_del_punto_y_del_producto(
    app_bi_trade, catalogo_partners
):
    """En el formulario se ven pegados; en la tabla van por separado."""
    norte, _sur, producto = catalogo_partners
    promotor = crear_usuario('promotor@supli.tech', app_bi_trade)

    respuesta = cliente_de(promotor).post(
        '/api/bi-trade/partners/registros', _recomendacion(norte, producto), format='json'
    )

    assert respuesta.status_code == 201, respuesta.data
    cuerpo = respuesta.json()
    assert cuerpo['idPuntoVenta'] == 'C192'
    assert cuerpo['idProducto'] == '7015490'
    assert cuerpo['puntoVentaEtiqueta'] == 'Cav Cucuta Centro Av Quinta - C192'
    assert cuerpo['productoEtiqueta'] == 'Estandar - 7015490'

    registro = RegistroPartner.objects.get()
    assert (registro.id_punto_venta_id, registro.id_producto_id) == ('C192', '7015490')
    assert registro.registrado_por == promotor


def test_el_punto_de_venta_tiene_que_ser_de_la_regional_elegida(app_bi_trade, catalogo_partners):
    norte, sur, producto = catalogo_partners
    promotor = crear_usuario('promotor@supli.tech', app_bi_trade)

    cuerpo = _recomendacion(norte, producto)
    cuerpo['idRegional'] = sur.id_regional_id
    respuesta = cliente_de(promotor).post(
        '/api/bi-trade/partners/registros', cuerpo, format='json'
    )

    assert respuesta.status_code == 400
    assert RegistroPartner.objects.count() == 0


def test_avisa_si_el_serial_ya_estaba_registrado(app_bi_trade, catalogo_partners):
    """El repetido no se bloquea —un equipo puede volver— pero sí se avisa."""
    norte, _sur, producto = catalogo_partners
    cliente = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade))
    ruta = '/api/bi-trade/partners/registros'

    primera = cliente.post(ruta, _recomendacion(norte, producto), format='json')
    segunda = cliente.post(ruta, _recomendacion(norte, producto), format='json')

    assert 'ya tenía' not in primera.json()['message']
    assert segunda.status_code == 201
    assert 'ya tenía 1 registro(s)' in segunda.json()['message']
    assert RegistroPartner.objects.count() == 2


def test_el_documento_del_promotor_va_sin_puntos(app_bi_trade, catalogo_partners):
    norte, _sur, producto = catalogo_partners
    cuerpo = _recomendacion(norte, producto)
    cuerpo['documentoPromotor'] = '1.092.389.375'

    respuesta = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade)).post(
        '/api/bi-trade/partners/registros', cuerpo, format='json'
    )

    assert respuesta.status_code == 400
    assert 'documentoPromotor' in respuesta.json()['errors']


def test_registrar_no_pide_permiso_pero_corregir_si(app_bi_trade, catalogo_partners):
    norte, _sur, producto = catalogo_partners
    promotor = crear_usuario('promotor@supli.tech', app_bi_trade)
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    ruta = '/api/bi-trade/partners/registros'

    creada = cliente_de(promotor).post(ruta, _recomendacion(norte, producto), format='json')
    assert creada.status_code == 201

    detalle = f'{ruta}/{creada.json()["idRegistro"]}'
    correccion = {'factura': '123456'}
    assert cliente_de(promotor).patch(detalle, correccion, format='json').status_code == 403
    assert cliente_de(editor).patch(detalle, correccion, format='json').status_code == 200


def test_las_opciones_traen_cada_punto_con_su_regional(app_bi_trade, catalogo_partners):
    """El formulario filtra los puntos con esto: antes eran tres preguntas."""
    norte, _sur, _producto = catalogo_partners
    respuesta = cliente_de(crear_usuario('promotor@supli.tech', app_bi_trade)).get(
        '/api/bi-trade/partners/opciones'
    )

    assert respuesta.status_code == 200
    opciones = respuesta.json()
    assert len(opciones['regionales']) == 2
    # Atado al modelo: agregar una marca no debe romper esta prueba.
    assert len(opciones['marcas']) == len(MarcaPartner.choices)
    por_codigo = {punto['value']: punto for punto in opciones['puntosVenta']}
    assert por_codigo['C192']['idRegional'] == norte.id_regional_id
    assert por_codigo['C192']['label'] == 'Cav Cucuta Centro Av Quinta - C192'


def test_las_listas_del_formulario_se_administran_desde_ahi(app_bi_trade, catalogo_partners):
    """Regionales, puntos y productos se agregan sin pasar por un despliegue."""
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)

    regional = cliente.post(
        '/api/bi-trade/partners/regionales', {'nombre': 'Región Eje Cafetero'}, format='json'
    )
    assert regional.status_code == 201

    punto = cliente.post(
        '/api/bi-trade/partners/puntos-venta',
        {
            'idPuntoVenta': 'C900',
            'nombrePdv': 'Cav Pereira Victoria',
            'idRegional': regional.json()['idRegional'],
        },
        format='json',
    )
    assert punto.status_code == 201
    assert punto.json()['etiqueta'] == 'Cav Pereira Victoria - C900'

    producto = cliente.post(
        '/api/bi-trade/partners/productos',
        {'idProducto': '7020000', 'nombreProducto': 'Clear'},
        format='json',
    )
    assert producto.status_code == 201

    # Lo que se agregó ya aparece en el formulario.
    opciones = cliente.get('/api/bi-trade/partners/opciones').json()
    assert 'C900' in {p['value'] for p in opciones['puntosVenta']}
    assert '7020000' in {p['value'] for p in opciones['productos']}

    # Y se puede quitar mientras nadie lo haya usado.
    assert cliente.delete('/api/bi-trade/partners/productos/7020000').status_code == 204
    assert cliente.delete('/api/bi-trade/partners/puntos-venta/C900').status_code == 204


def test_administrar_las_listas_pide_permiso(app_bi_trade, catalogo_partners):
    promotor = crear_usuario('promotor@supli.tech', app_bi_trade)
    respuesta = cliente_de(promotor).post(
        '/api/bi-trade/partners/regionales', {'nombre': 'Región Eje Cafetero'}, format='json'
    )
    assert respuesta.status_code == 403


def test_no_se_borra_del_catalogo_lo_que_ya_tiene_registros(app_bi_trade, catalogo_partners):
    """Se desactiva, no se borra: si no, cambiarían registros ya cargados."""
    norte, _sur, producto = catalogo_partners
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)
    cliente.post('/api/bi-trade/partners/registros', _recomendacion(norte, producto), format='json')

    for ruta in (
        f'/api/bi-trade/partners/puntos-venta/{norte.pk}',
        f'/api/bi-trade/partners/productos/{producto.pk}',
        f'/api/bi-trade/partners/regionales/{norte.id_regional_id}',
    ):
        respuesta = cliente.delete(ruta)
        assert respuesta.status_code == 400, ruta
        assert respuesta.json()['code'] == 'protected'

    # Desactivarlo sí se puede, y deja de aparecer en el formulario.
    assert (
        cliente.patch(
            f'/api/bi-trade/partners/productos/{producto.pk}', {'activo': False}, format='json'
        ).status_code
        == 200
    )
    opciones = cliente.get('/api/bi-trade/partners/opciones').json()
    assert producto.pk not in {p['value'] for p in opciones['productos']}


# ── Plan Partners: el formulario abierto por enlace ────────────────────────
#
# Va por su propio prefijo y sin contraseña: se diligencia a diario y no
# muestra nada. Lo que hay que cuidar es que por esa puerta no se pueda hacer
# nada más que enviar.

RUTA_FORMULARIO = '/api/publico/formulario'


@pytest.fixture
def enlace_formulario(app_bi_trade):
    """Un enlace del formulario del plan: solo el token, no tiene contraseña."""
    editor = crear_usuario('editor-form@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    datos = (
        cliente_de(editor)
        .post(
            '/api/bi-trade/enlaces',
            {'nombre': 'Promotores Costa', 'canal': 'partners'},
            format='json',
        )
        .json()
    )
    assert datos['clave'] == ''
    assert EnlacePublico.objects.get(token=datos['token']).clave_hash == ''
    return datos['token']


def test_el_formulario_publico_se_diligencia_sin_cuenta_ni_clave(
    enlace_formulario, catalogo_partners
):
    norte, _sur, producto = catalogo_partners
    token = enlace_formulario
    anonimo = APIClient()

    opciones = anonimo.get(f'{RUTA_FORMULARIO}/{token}/opciones')
    assert opciones.status_code == 200
    assert len(opciones.json()['puntosVenta']) == 2

    respuesta = anonimo.post(
        f'{RUTA_FORMULARIO}/{token}/registros',
        _recomendacion(norte, producto),
        format='json',
    )

    assert respuesta.status_code == 201
    assert respuesta.json()['origen'] == 'Enlace · Promotores Costa'
    registro = RegistroPartner.objects.get()
    assert registro.registrado_por is None
    assert registro.enlace.nombre == 'Promotores Costa'
    # En un enlace de formulario los «accesos» son los envíos recibidos.
    assert EnlacePublico.objects.get(pk=registro.enlace_id).accesos == 1


def test_por_el_formulario_no_se_puede_hacer_nada_mas(enlace_formulario):
    """Dos rutas y nada más: ni tableros, ni lo ya cargado, ni catálogos."""
    token = enlace_formulario
    anonimo = APIClient()

    for ruta in (
        f'{RUTA_PUBLICA}/{token}',  # la puerta del tablero
        f'{RUTA_PUBLICA}/{token}/avance-mensual',
        f'{RUTA_PUBLICA}/{token}/cumplimiento-diario',
        f'{RUTA_PUBLICA}/{token}/tickets',
        f'{RUTA_PUBLICA}/{token}/productos',
        f'{RUTA_PUBLICA}/{token}/puntos-venta',
    ):
        assert anonimo.get(ruta).status_code == 404, ruta

    # Enviar es lo único: los registros cargados no se leen desde aquí.
    assert anonimo.get(f'{RUTA_FORMULARIO}/{token}/registros').status_code == 405
    # Y la API de la app le sigue cerrada.
    assert anonimo.get('/api/bi-trade/partners/registros').status_code in (401, 403)


def test_un_enlace_de_tablero_no_sirve_para_el_formulario(enlace, catalogo_partners):
    norte, _sur, producto = catalogo_partners
    token, _clave, _id = enlace
    anonimo = APIClient()

    assert anonimo.get(f'{RUTA_FORMULARIO}/{token}/opciones').status_code == 404
    respuesta = anonimo.post(
        f'{RUTA_FORMULARIO}/{token}/registros',
        _recomendacion(norte, producto),
        format='json',
    )
    assert respuesta.status_code == 404
    assert RegistroPartner.objects.count() == 0


def test_un_formulario_revocado_deja_de_recibir(
    app_bi_trade, enlace_formulario, catalogo_partners
):
    norte, _sur, producto = catalogo_partners
    token = enlace_formulario
    editor = crear_usuario('revoca@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    id_enlace = EnlacePublico.objects.get(token=token).pk
    cliente_de(editor).patch(
        f'/api/bi-trade/enlaces/{id_enlace}', {'activo': False}, format='json'
    )

    anonimo = APIClient()
    assert anonimo.get(f'{RUTA_FORMULARIO}/{token}/opciones').status_code == 404
    respuesta = anonimo.post(
        f'{RUTA_FORMULARIO}/{token}/registros',
        _recomendacion(norte, producto),
        format='json',
    )
    assert respuesta.status_code == 404
    assert RegistroPartner.objects.count() == 0


def test_el_formulario_no_tiene_contrasena_que_regenerar(app_bi_trade, enlace_formulario):
    editor = crear_usuario('regenera@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    id_enlace = EnlacePublico.objects.get(token=enlace_formulario).pk

    respuesta = cliente_de(editor).post(f'/api/bi-trade/enlaces/{id_enlace}/regenerar-clave')

    assert respuesta.status_code == 400
    assert respuesta.json()['code'] == 'sin_clave'


# ── Plan Partners: metas y tablero ─────────────────────────────────────────


def _excel_de_metas(filas: list[tuple]) -> SimpleUploadedFile:
    """Un archivo con los mismos encabezados del Excel mensual de Trade."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'META'
    hoja.append(['MES', 'AÑO', 'CENTRO DE COSTOS', 'META', 'MARCA', 'PUNTO DE VENTA'])
    for fila in filas:
        hoja.append(list(fila))
    buffer = BytesIO()
    libro.save(buffer)
    return SimpleUploadedFile(
        'metas.xlsx',
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def test_las_metas_se_suben_con_el_excel_de_trade(app_bi_trade, catalogo_partners):
    """El archivo mensual entra tal cual; lo que no está en el catálogo se lista."""
    editor = crear_usuario('metas@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    archivo = _excel_de_metas(
        [
            ('JULIO', 2026, 'C192', 180.83, 'SAMSUNG', 'CAV CUCUTA CENTRO AV QUINTA'),
            ('JULIO', 2026, 'C108', '60,5', 'MOTOROLA', 'CAV BOGOTA PLAZA CLARO'),
            ('JULIO', 2026, 'C999', 10, 'SAMSUNG', 'CAV QUE NO EXISTE'),
            ('JULIO', 2026, 'C192', 10, 'XIAOMI', 'CAV CUCUTA CENTRO AV QUINTA'),
        ]
    )

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/partners/metas/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 200, respuesta.data
    datos = respuesta.json()
    assert (datos['created'], datos['updated'], datos['skipped']) == (2, 0, 2)
    assert datos['puntosDesconocidos'] == ['C999']
    assert datos['marcasDesconocidas'] == ['XIAOMI']
    # Los decimales del archivo se respetan, con coma o con punto.
    assert str(MetaPartner.objects.get(id_punto_venta='C192').meta_unidades) == '180.83'
    assert str(MetaPartner.objects.get(id_punto_venta='C108').meta_unidades) == '60.50'


def test_subir_las_metas_pide_permiso(app_bi_trade, catalogo_partners):
    promotor = crear_usuario('promotor@supli.tech', app_bi_trade)
    archivo = _excel_de_metas([('JULIO', 2026, 'C192', 10, 'SAMSUNG', 'X')])

    respuesta = cliente_de(promotor).post(
        '/api/bi-trade/partners/metas/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 403
    assert MetaPartner.objects.count() == 0


def test_el_tablero_mide_lo_registrado_contra_la_meta(app_bi_trade, catalogo_partners):
    norte, _sur, producto = catalogo_partners
    MetaPartner.objects.create(
        anio=2026, mes=7, id_punto_venta=norte, marca=MarcaPartner.MOTOROLA, meta_unidades=10
    )
    for dia in (3, 3, 8):
        RegistroPartner.objects.create(
            id_regional=norte.id_regional,
            marca=MarcaPartner.MOTOROLA,
            id_punto_venta=norte,
            id_producto=producto,
            fecha_recomendacion=date(2026, 7, dia),
            serial=f'S{dia}{RegistroPartner.objects.count()}',
            documento_promotor='1092389375',
            factura='F1',
        )
    # De otro mes: no debe entrar en el corte de julio.
    RegistroPartner.objects.create(
        id_regional=norte.id_regional,
        marca=MarcaPartner.MOTOROLA,
        id_punto_venta=norte,
        id_producto=producto,
        fecha_recomendacion=date(2026, 8, 1),
        serial='OTRO-MES',
        documento_promotor='1092389375',
        factura='F2',
    )

    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))
    datos = cliente.get('/api/bi-trade/partners/dashboard?anio=2026&mes=7').json()

    assert datos['filtros']['periodo'] == 'Julio 2026'
    assert datos['totales']['unidades'] == 3
    assert datos['totales']['meta'] == 10.0
    assert datos['totales']['cumplimiento'] == 30.0
    assert datos['totales']['faltante'] == 7.0
    assert datos['totales']['promotores'] == 1
    # Dos registros el día 3 y uno el 8.
    assert [(f['fecha'], f['unidades']) for f in datos['porDia']] == [
        ('2026-07-03', 2),
        ('2026-07-08', 1),
    ]
    motorola = next(f for f in datos['porMarca'] if f['marca'] == MarcaPartner.MOTOROLA.value)
    assert (motorola['unidades'], motorola['meta']) == (3, 10.0)


def test_el_tablero_filtra_por_marca_y_por_punto(app_bi_trade, catalogo_partners):
    """Filtrar por promotor no toca la meta: no está repartida por persona."""
    norte, sur, producto = catalogo_partners
    MetaPartner.objects.create(
        anio=2026, mes=7, id_punto_venta=norte, marca=MarcaPartner.MOTOROLA, meta_unidades=10
    )
    MetaPartner.objects.create(
        anio=2026, mes=7, id_punto_venta=sur, marca=MarcaPartner.SAMSUNG, meta_unidades=40
    )
    RegistroPartner.objects.create(
        id_regional=norte.id_regional,
        marca=MarcaPartner.MOTOROLA,
        id_punto_venta=norte,
        id_producto=producto,
        fecha_recomendacion=date(2026, 7, 5),
        serial='UNO',
        documento_promotor='111',
        factura='F',
    )
    cliente = cliente_de(crear_usuario('lector@supli.tech', app_bi_trade))

    solo_norte = cliente.get(
        f'/api/bi-trade/partners/dashboard?anio=2026&mes=7&punto={norte.pk}'
    ).json()
    assert (solo_norte['totales']['unidades'], solo_norte['totales']['meta']) == (1, 10.0)

    por_promotor = cliente.get(
        '/api/bi-trade/partners/dashboard?anio=2026&mes=7&promotor=111'
    ).json()
    assert por_promotor['totales']['unidades'] == 1
    assert por_promotor['totales']['meta'] == 50.0

    ano_completo = cliente.get('/api/bi-trade/partners/dashboard?anio=2026&mes=0').json()
    assert ano_completo['filtros']['periodo'] == 'Año 2026'
    assert ano_completo['totales']['unidades'] == 1


# ── Importaciones en bloque: el número de consultas no crece con las filas ──
#
# Antes cada fila iba sola a la base (buscar si existía, validar el producto y
# el punto, revisar duplicados y guardar). Con el backend local contra Render
# eran ~0,6 s por fila y el navegador cortaba a los 30 s aunque el servidor
# siguiera guardando. Estas pruebas fijan que el costo sea el mismo con 3 filas
# que con 60.


def _consultas_de(cliente, ruta, archivo):
    # La primera petición de un usuario llena la caché de sus permisos: se
    # hace antes para medir solo la importación.
    cliente.get('/api/bi-trade/opciones')
    with CaptureQueriesContext(connection) as consultas:
        respuesta = cliente.post(ruta, {'archivo': archivo}, format='multipart')
    assert respuesta.status_code == 200, respuesta.data
    return len(consultas), respuesta


def _catalogo_grande(puntos=6, productos=10):
    PuntoVenta.objects.bulk_create(
        PuntoVenta(id_punto_venta=f'PDV-{n}', nombre_pdv=f'Punto {n}') for n in range(puntos)
    )
    Producto.objects.bulk_create(
        Producto(
            id_producto=f'SKU-{n}',
            nombre_producto=f'Producto {n}',
            marca='Samsung',
            precio_venta_claro=1000,
            precio_venta_coltrade=900,
            puntaje=10,
        )
        for n in range(productos)
    )


def _filas_de_metas(cantidad, unidades=5):
    return [(f'SKU-{n % 10}', f'PDV-{n // 10}', '2026-10-01', unidades) for n in range(cantidad)]


def test_importar_metas_no_hace_una_consulta_por_fila(app_bi_trade):
    _catalogo_grande()
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)
    encabezados = ('id_producto', 'id_punto_venta', 'fecha_meta', 'meta_cantidad')
    ruta = '/api/bi-trade/metas/importar'

    pocas, _ = _consultas_de(cliente, ruta, _archivo_xlsx(encabezados, _filas_de_metas(3)))
    MetaComercial.objects.all().delete()
    muchas, respuesta = _consultas_de(
        cliente, ruta, _archivo_xlsx(encabezados, _filas_de_metas(60))
    )

    assert respuesta.data['created'] == 60
    assert MetaComercial.objects.count() == 60
    assert muchas == pocas

    # Reimportar actualiza en bloque: tampoco depende de cuántas filas son.
    actualizar, respuesta = _consultas_de(
        cliente, ruta, _archivo_xlsx(encabezados, _filas_de_metas(60, unidades=9))
    )
    assert respuesta.data['updated'] == 60
    assert MetaComercial.objects.count() == 60
    assert set(MetaComercial.objects.values_list('meta_cantidad', flat=True)) == {9}
    assert actualizar <= muchas + 1


def test_importar_puntos_de_venta_no_hace_una_consulta_por_fila(app_bi_trade):
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)
    encabezados = ('id_punto_venta', 'nombre_pdv', 'regional', 'materiales')
    ruta = '/api/bi-trade/puntos-venta/importar'

    def filas(cantidad):
        return [(f'PDV-{n}', f'Punto {n}', 'Zona Norte', None) for n in range(cantidad)]

    pocas, _ = _consultas_de(cliente, ruta, _archivo_xlsx(encabezados, filas(3)))
    PuntoVenta.objects.all().delete()
    muchas, _ = _consultas_de(cliente, ruta, _archivo_xlsx(encabezados, filas(40)))
    assert muchas == pocas
    assert PuntoVenta.objects.count() == 40

    # La mitad existe y la otra mitad es nueva: una sola carga hace las dos cosas.
    mixtas = [(f'PDV-{n}', f'Nuevo {n}', 'Zona Sur', None) for n in range(20, 60)]
    _, respuesta = _consultas_de(cliente, ruta, _archivo_xlsx(encabezados, mixtas))
    assert (respuesta.data['created'], respuesta.data['updated']) == (20, 20)
    assert PuntoVenta.objects.count() == 60
    assert PuntoVenta.objects.get(pk='PDV-25').nombre_pdv == 'Nuevo 25'
    assert PuntoVenta.objects.get(pk='PDV-5').nombre_pdv == 'Punto 5'


def test_importar_ventas_no_hace_una_consulta_por_fila(app_bi_trade):
    _catalogo_grande()
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)
    encabezados = ('id_producto', 'id_punto_venta', 'fecha_venta', 'cantidad_vendida')
    ruta = '/api/bi-trade/ventas/importar'

    pocas, _ = _consultas_de(cliente, ruta, _archivo_xlsx(encabezados, _filas_de_metas(3)))
    muchas, _ = _consultas_de(cliente, ruta, _archivo_xlsx(encabezados, _filas_de_metas(60)))

    assert muchas == pocas
    assert Venta.objects.count() == 63


def test_una_clave_repetida_en_el_archivo_actualiza_la_primera(app_bi_trade, catalogo):
    """Igual que antes: la segunda fila con la misma clave pisa a la primera."""
    pdv, _sur, producto, _barato = catalogo
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    archivo = _archivo_xlsx(
        ('id_producto', 'id_punto_venta', 'fecha_meta', 'meta_cantidad'),
        [
            (producto.pk, pdv.pk, '2026-10-01', 5),
            (producto.pk, pdv.pk, '2026-10-01', 8),
        ],
    )

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/metas/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 200, respuesta.data
    assert (respuesta.data['created'], respuesta.data['updated']) == (1, 1)
    assert MetaComercial.objects.get().meta_cantidad == 8


def test_importar_con_un_codigo_que_no_existe_no_guarda_nada(app_bi_trade, catalogo):
    pdv, _sur, producto, _barato = catalogo
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    archivo = _archivo_xlsx(
        ('id_producto', 'id_punto_venta', 'fecha_meta', 'meta_cantidad'),
        [
            (producto.pk, pdv.pk, '2026-10-01', 5),
            ('SKU-NO-EXISTE', pdv.pk, '2026-10-01', 5),
        ],
    )

    respuesta = cliente_de(editor).post(
        '/api/bi-trade/metas/importar', {'archivo': archivo}, format='multipart'
    )

    assert respuesta.status_code == 400
    assert respuesta.data['filas'][0]['fila'] == 3
    assert 'id_producto' in respuesta.data['filas'][0]['errores'][0]
    assert MetaComercial.objects.count() == 0


def test_importar_en_un_canal_usa_sus_encabezados_y_actualiza(app_bi_trade, catalogo_hc):
    """Los canales heredan la importación: `_hc` en el archivo, la tabla de HC en la base."""
    editor = crear_usuario('editor@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)
    encabezados = ('id_producto_hc', 'id_punto_venta_hc', 'fecha_meta', 'meta_cantidad')
    ruta = '/api/bi-trade/hc/metas/importar'

    primera = cliente.post(
        ruta,
        {'archivo': _archivo_xlsx(encabezados, [('HCP-001', 'HC-101', '2026-10-01', 4)])},
        format='multipart',
    )
    segunda = cliente.post(
        ruta,
        {'archivo': _archivo_xlsx(encabezados, [('HCP-001', 'HC-101', '2026-10-01', 7)])},
        format='multipart',
    )

    assert primera.data['created'] == 1
    assert segunda.data['updated'] == 1
    assert MetaComercialHc.objects.get().meta_cantidad == 7
    assert MetaComercial.objects.count() == 0


def test_el_listado_de_puntos_cuenta_las_ventas_sin_traerlas(app_bi_trade, catalogo):
    pdv, sur, producto, _barato = catalogo
    for _ in range(3):
        Venta.objects.create(
            id_producto=producto,
            id_punto_venta=pdv,
            fecha_venta=date(2026, 3, 4),
            cantidad_vendida=1,
        )
    lector = crear_usuario('lector@supli.tech', app_bi_trade)

    datos = cliente_de(lector).get('/api/bi-trade/puntos-venta').json()

    conteos = {fila['idPuntoVenta']: fila['ventasCount'] for fila in datos}
    assert conteos == {pdv.pk: 3, sur.pk: 0}


def test_importar_metas_partners_no_hace_una_consulta_por_fila(app_bi_trade, catalogo_partners):
    editor = crear_usuario('metas@supli.tech', app_bi_trade, ['bi-trade:data:manage'])
    cliente = cliente_de(editor)
    ruta = '/api/bi-trade/partners/metas/importar'
    marcas = [marca.upper() for marca in MarcaPartner.values]

    def filas(cantidad, meta=10):
        return [
            ('JULIO', 2026, ('C192', 'C108')[n % 2], meta, marcas[n % len(marcas)], 'X')
            for n in range(cantidad)
        ]

    pocas, _ = _consultas_de(cliente, ruta, _excel_de_metas(filas(2)))
    MetaPartner.objects.all().delete()
    muchas, respuesta = _consultas_de(cliente, ruta, _excel_de_metas(filas(2 * len(marcas))))

    assert muchas == pocas
    assert respuesta.json()['created'] == 2 * len(marcas)

    _, respuesta = _consultas_de(cliente, ruta, _excel_de_metas(filas(2 * len(marcas), meta=3)))
    assert respuesta.json()['updated'] == 2 * len(marcas)
    metas = {str(valor) for valor in MetaPartner.objects.values_list('meta_unidades', flat=True)}
    assert metas == {'3.00'}
